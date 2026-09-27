# -*- coding: utf-8 -*-
"""
DỰNG BỘ HỒ SƠ CẦU THÔN KHAI HOANG 2, KM14+363.65 TỪ DỮ LIỆU THẬT

Nguồn: bang_so_lieu.json (bảng số liệu trích từ PDF BVTC) + THKL G9 - CHUẨN.xlsx (sheet "Cau 15m").
Đầu ra (thư mục HO_SO_THIET_LAP cạnh dữ liệu gốc):
  01_BBS_Thep_Cau_Khai_Hoang_2.xlsx      — thống kê cốt thép từ bản vẽ, đối chiếu THKL (đầu vào pha rebar)
  02_BOQ_Khoi_Luong_Cau_Khai_Hoang_2.xlsx — tiên lượng từ THKL, cột đơn giá để trống (đầu vào pha qs)
  03_Tien_Do_Cau_Khai_Hoang_2.xlsx / .xml — tiến độ WBS đề xuất có quan hệ logic (đầu vào pha schedule)
  04_KCS_Nghiem_Thu_Thi_Nghiem_Cau_Khai_Hoang_2.xlsx — danh mục nghiệm thu + tần suất thí nghiệm
Mọi số liệu lấy từ nguồn; số liệu đề xuất (thời gian công tác) được ghi rõ là đề xuất.

Chạy:  python examples/build_khai_hoang_2_dossier.py
"""

from __future__ import annotations

import os
import sys
import xml.etree.ElementTree as ET
from collections import OrderedDict, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from khai_hoang_2_source import (SOURCE_DIR, GRADE, UNIT_WEIGHT, check_bbs_weights,  # noqa: E402
                                 load_bbs, load_thkl)

OUT_DIR = os.path.join(SOURCE_DIR, "HO_SO_THIET_LAP")
PROJECT = "Đường từ trung tâm huyện Đồng Văn đi Mốc 450 (nay là Mốc 456), huyện Mèo Vạc, tỉnh Hà Giang"
PACKAGE = "Gói thầu số 9: Km12 - Km24+862,93"
BRIDGE = "CẦU THÔN KHAI HOANG 2, KM14+363,65 (bản vẽ ghi Km14+363,35)"

# ── Style ────────────────────────────────────────────────────────────────────
F_TITLE = Font(name="Times New Roman", size=13, bold=True, color="1F497D")
F_SUB = Font(name="Times New Roman", size=10, italic=True, color="595959")
F_HDR = Font(name="Times New Roman", size=10, bold=True, color="FFFFFF")
F_BOLD = Font(name="Times New Roman", size=10, bold=True)
F_REG = Font(name="Times New Roman", size=10)
FILL_HDR = PatternFill("solid", start_color="1F497D")
FILL_SEC = PatternFill("solid", start_color="DCE6F1")
FILL_INPUT = PatternFill("solid", start_color="FFF2CC")
FILL_WARN = PatternFill("solid", start_color="FCE4D6")
FILL_OK = PatternFill("solid", start_color="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")


def title_block(ws, title, subtitle, ncols):
    ws["A1"] = f"DỰ ÁN: {PROJECT.upper()} — {PACKAGE}"
    ws["A1"].font = F_SUB
    ws["A2"] = title
    ws["A2"].font = F_TITLE
    ws["A3"] = subtitle
    ws["A3"].font = F_SUB
    for r in (1, 2, 3):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)


def header_row(ws, row, headers, widths=None):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row, c, h)
        cell.font, cell.fill, cell.alignment, cell.border = F_HDR, FILL_HDR, CENTER, BORDER
    ws.row_dimensions[row].height = 32
    if widths:
        for c, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row + 1, 1)


def put(ws, row, col, value, fmt=None, font=F_REG, align=None, fill=None):
    cell = ws.cell(row, col, value)
    cell.font = font
    cell.border = BORDER
    cell.alignment = align or (RIGHT if isinstance(value, (int, float)) or str(value).startswith("=") else LEFT)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    return cell


# ═════════════════════════════════════════════════════════════════════════════
# 01 — BBS
# ═════════════════════════════════════════════════════════════════════════════

# Nhóm BBS ↔ dòng THKL để đối chiếu: (tên nhóm, [hạng mục BBS], {nhóm Ø: mã THKL})
THKL_STEEL_MAP = [
    ("Dầm T L=15m (4 dầm)", ["Dầm T L=15m - dầm biên (1 dầm)", "Dầm T L=15m - dầm giữa (1 dầm)"],
     {"D<=10": ["2.1-2"], "10<D<=18": ["2.1-3"], "D>18": ["2.1-4"]}),
    ("Lớp phủ mặt cầu (lưới D6)", ["Lớp phủ mặt cầu + gờ lan can (1 nhịp)#B"], {"D<=10": ["2.1-9"]}),
    ("Gờ lan can (trên nhịp + trên mố)", ["Lớp phủ mặt cầu + gờ lan can (1 nhịp)#L"], {"10<D<=18": ["3.2-2"]}),
    ("Bản quá độ (2 bản)", ["Bản quá độ (1 bản)"],
     {"D<=10": ["2.3-3"], "10<D<=18": ["2.3-4"], "D>18": ["2.3-5"]}),
    ("Khe co giãn (2 khe)", ["Khe co giãn (1 khe)"], {"10<D<=18": ["3.2-13"]}),
    ("Thép neo D32 vào đá", ["Thép neo mố vào đá (toàn cầu)"], {"D>18": ["3.3-3"]}),
    ("Mố M1+M2 (bệ, thân, bệ kê gối, 4 tường cánh)",
     ["Bệ mố - bảng 1 (tr.14)", "Bệ mố - bảng 2 (tr.15)", "Thân mố + tường đỉnh (1 mố)",
      "Bệ kê gối + u neo dầm (1 mố)", "Tường cánh (1 tường)"],
     {"D<=10": ["3.3-7"], "10<D<=18": ["3.3-8"], "D>18": ["3.3-9"]}),
    ("Tường chắn đầu cầu (4 đốt, móng + thân)",
     ["Tường chắn đốt L=5,20m", "Tường chắn đốt L=5,89m", "Tường chắn đốt L=10,60m", "Tường chắn đốt L=13,50m"],
     {"10<D<=18": ["2.4-2", "2.4-6"], "D>18": ["2.4-3", "2.4-7"]}),
]


def dia_group(d):
    return "D<=10" if d <= 10 else ("10<D<=18" if d <= 18 else "D>18")


def build_bbs(rows, log, thkl):
    wb = Workbook()
    ws = wb.active
    ws.title = "BBS"
    headers = ["TT", "Hạng mục kết cấu", "Ký hiệu thanh", "Đường kính Ø (mm)", "Mác thép",
               "Chiều dài 1 thanh (mm)", "Số thanh / cấu kiện", "Số cấu kiện", "Tổng thanh toàn cầu",
               "Trọng lượng đơn vị (kg/m)", "Khối lượng toàn cầu (kg)", "KL ghi trên bản vẽ (kg/1 CK)",
               "Lệch so với bản vẽ", "Bản vẽ", "Trang PDF", "Căn cứ số cấu kiện"]
    title_block(ws, f"BẢNG THỐNG KÊ CỐT THÉP (BBS) — {BRIDGE}",
                "Trích từ bảng thép trên bản vẽ BVTC; khối lượng tính lại = L × n × kg/m chuẩn TCVN 1651. "
                "Ô cam: lệch > 2% so với bản vẽ — xem sheet CANH_BAO.", len(headers))
    HR = 5
    header_row(ws, HR, headers, [5, 36, 9, 10, 10, 11, 10, 9, 11, 11, 13, 13, 10, 8, 7, 44])
    r = HR + 1
    for i, b in enumerate(rows, start=1):
        put(ws, r, 1, i, align=CENTER)
        put(ws, r, 2, b.component)
        put(ws, r, 3, b.mark, align=CENTER)
        put(ws, r, 4, b.dia, align=CENTER)
        put(ws, r, 5, GRADE[b.dia], align=CENTER)
        put(ws, r, 6, b.length_mm, "#,##0")
        put(ws, r, 7, b.qty_per_unit, "#,##0")
        put(ws, r, 8, b.units, "#,##0", align=CENTER)
        put(ws, r, 9, f"=G{r}*H{r}", "#,##0")
        put(ws, r, 10, UNIT_WEIGHT[b.dia], "0.000")
        put(ws, r, 11, f"=F{r}/1000*I{r}*J{r}", "#,##0.00")
        put(ws, r, 12, b.drawing_kg, "#,##0.00")
        put(ws, r, 13, f'=IF(L{r}>0,(F{r}/1000*G{r}*J{r})/L{r}-1,"")', "+0.0%;-0.0%;0.0%")
        diff = (b.kg_per_unit / b.drawing_kg - 1) if b.drawing_kg else 0
        if abs(diff) > 0.02:
            for c in (12, 13):
                ws.cell(r, c).fill = FILL_WARN
        put(ws, r, 14, b.drawing, align=CENTER)
        put(ws, r, 15, b.page, align=CENTER)
        put(ws, r, 16, b.note)
        r += 1
    last = r - 1
    put(ws, r, 2, "TỔNG CỘNG", font=F_BOLD)
    put(ws, r, 9, f"=SUM(I{HR+1}:I{last})", "#,##0", font=F_BOLD)
    put(ws, r, 11, f"=SUM(K{HR+1}:K{last})", "#,##0.00", font=F_BOLD)
    ws.auto_filter.ref = f"A{HR}:P{last}"
    rng = lambda col: f"BBS!${col}${HR+1}:${col}${last}"  # noqa: E731

    # ── Tổng hợp theo hạng mục × đường kính ──
    ws2 = wb.create_sheet("TONG_HOP_THEO_D")
    comps = list(OrderedDict.fromkeys(b.component for b in rows))
    dias = sorted({b.dia for b in rows})
    headers2 = ["Hạng mục kết cấu"] + [f"Ø{d} (kg)" for d in dias] + ["D≤10", "10<D≤18", "D>18", "Tổng (kg)"]
    title_block(ws2, "TỔNG HỢP CỐT THÉP THEO HẠNG MỤC VÀ ĐƯỜNG KÍNH (toàn cầu)",
                "SUMIFS trực tiếp từ sheet BBS", len(headers2))
    header_row(ws2, 5, headers2, [40] + [10] * (len(dias) + 4))
    r = 6
    nd = len(dias)
    for comp in comps:
        put(ws2, r, 1, comp)
        for j, d in enumerate(dias, start=2):
            put(ws2, r, j, f'=SUMIFS({rng("K")},{rng("B")},$A{r},{rng("D")},{d})', "#,##0.00")
        cols = {d: get_column_letter(j) for j, d in enumerate(dias, start=2)}
        g1 = "+".join(f"{cols[d]}{r}" for d in dias if d <= 10) or "0"
        g2 = "+".join(f"{cols[d]}{r}" for d in dias if 10 < d <= 18) or "0"
        g3 = "+".join(f"{cols[d]}{r}" for d in dias if d > 18) or "0"
        put(ws2, r, nd + 2, f"={g1}", "#,##0.00")
        put(ws2, r, nd + 3, f"={g2}", "#,##0.00")
        put(ws2, r, nd + 4, f"={g3}", "#,##0.00")
        put(ws2, r, nd + 5, f"=SUM(B{r}:{get_column_letter(nd+1)}{r})", "#,##0.00", font=F_BOLD)
        r += 1
    put(ws2, r, 1, "TỔNG CỘNG", font=F_BOLD)
    for c in range(2, nd + 6):
        L = get_column_letter(c)
        put(ws2, r, c, f"=SUM({L}6:{L}{r-1})", "#,##0.00", font=F_BOLD)

    # ── Đối chiếu THKL ──
    ws3 = wb.create_sheet("DOI_CHIEU_THKL")
    thkl_by_code = {it.code: it for it in thkl}
    headers3 = ["Nhóm kết cấu", "Nhóm Ø", "BBS bản vẽ (kg)", "Mã dòng THKL", "THKL 'Cau 15m' (kg)",
                "Chênh lệch (kg)", "Chênh lệch (%)", "Kết luận"]
    title_block(ws3, "ĐỐI CHIẾU KHỐI LƯỢNG THÉP: BBS BẢN VẼ ↔ BẢNG THKL G9 (sheet 'Cau 15m')",
                "Ngưỡng chấp nhận ±2%. Số THKL nhập nguyên giá trị từ file gốc.", len(headers3))
    header_row(ws3, 5, headers3, [40, 10, 14, 14, 16, 14, 11, 44])
    r = 6
    results = []
    for name, comp_keys, groups in THKL_STEEL_MAP:
        for grp, codes in groups.items():
            bbs_kg = 0.0
            formula_parts = []
            for ck in comp_keys:
                comp, _, prefix = ck.partition("#")
                sel = [b for b in rows if b.component == comp and dia_group(b.dia) == grp
                       and (not prefix or b.mark.startswith(prefix))]
                bbs_kg += sum(b.kg_total for b in sel)
                if prefix:   # cùng ký hiệu có thể xuất hiện ở nhiều dòng (gờ trên nhịp / trên mố) → mỗi cặp 1 lần
                    for b in {(b.mark, b.dia): b for b in sel}.values():
                        formula_parts.append(
                            f'SUMIFS({rng("K")},{rng("B")},"{comp}",{rng("C")},"{b.mark}",{rng("D")},{b.dia})')
                else:
                    ds = sorted({b.dia for b in sel})
                    formula_parts += [f'SUMIFS({rng("K")},{rng("B")},"{comp}",{rng("D")},{d})' for d in ds]
            thkl_kg = sum((thkl_by_code[c].qty or 0) for c in codes)
            put(ws3, r, 1, name)
            put(ws3, r, 2, grp, align=CENTER)
            put(ws3, r, 3, "=" + ("+".join(formula_parts) if formula_parts else "0"), "#,##0.00")
            put(ws3, r, 4, " + ".join(codes), align=CENTER)
            put(ws3, r, 5, round(thkl_kg, 3), "#,##0.00")
            put(ws3, r, 6, f"=C{r}-E{r}", "#,##0.00")
            put(ws3, r, 7, f'=IF(E{r}=0,"",F{r}/E{r})', "+0.0%;-0.0%;0.0%")
            ok = thkl_kg > 0 and abs(bbs_kg - thkl_kg) / thkl_kg <= 0.02
            put(ws3, r, 8, f'=IF(E{r}=0,"THKL không có số",IF(ABS(G{r})<=2%,"KHỚP","LỆCH — cần đối chiếu"))',
                fill=FILL_OK if ok else FILL_WARN)
            results.append((name, grp, bbs_kg, thkl_kg, ok))
            r += 1
    put(ws3, r + 1, 1, "Không có BBS chi tiết trên bản vẽ (chỉ có trong THKL):", font=F_BOLD)
    for k, code in enumerate(["2.2-2", "2.2-3", "3.4-3"], start=2):
        it = thkl_by_code[code]
        put(ws3, r + k, 1, f"{it.section} — {it.name}")
        put(ws3, r + k, 4, code, align=CENTER)
        put(ws3, r + k, 5, it.qty, "#,##0.00")
        put(ws3, r + k, 8, "Không có trong kế hoạch cắt thép — cần bản vẽ/BBS dầm ngang từ TVTK", fill=FILL_WARN)

    # ── Cảnh báo ──
    ws4 = wb.create_sheet("CANH_BAO")
    title_block(ws4, "NHẬT KÝ SỬA OCR VÀ CẢNH BÁO SỐ LIỆU BẢN VẼ", "Cần đối chiếu PDF gốc trước khi phát hành", 2)
    header_row(ws4, 5, ["TT", "Nội dung"], [5, 150])
    for i, line in enumerate(log, start=1):
        put(ws4, 5 + i, 1, i, align=CENTER)
        put(ws4, 5 + i, 2, line, fill=FILL_WARN if "LỆCH" in line else None)
    path = os.path.join(OUT_DIR, "01_BBS_Thep_Cau_Khai_Hoang_2.xlsx")
    wb.save(path)
    return path, results


# ═════════════════════════════════════════════════════════════════════════════
# 02 — BOQ
# ═════════════════════════════════════════════════════════════════════════════

# Ghi chú kiểm tra cho các dòng THKL (mã → ghi chú)
BOQ_NOTES = {
    "2.1-2.1": "Tiêu đề trong THKL chép từ mẫu cầu DƯL L=33m; cầu này là 4 dầm T BTCT thường L=15m",
    "2.1-1": "Khớp bản vẽ PT-01: 2×8,5 (dầm biên) + 2×10 (dầm giữa) = 37 m3",
    "2.3-2": "Bản vẽ PD-08 ghi BT 25MPa; thuyết minh mục 7.2 ghi bản quá độ fc'=30MPa — cần TVTK thống nhất",
    "2.1-8": "Bản vẽ PT-10 ghi 35MPa; thuyết minh mục 7.2 ghi bản mặt cầu fc'=30MPa — cần TVTK thống nhất",
    "3.3-4": "Bản vẽ PD-02 ghi BT đệm 12MPa: 4,67 + 5,24 = 9,91 m3 ≠ THKL 9,348 m3 (10MPa)",
    "3.3-5": "THKL thiếu ĐVT (bổ sung m3). Bản vẽ PD-02: 110,00 + 123,75 = 233,75 m3 ≠ THKL 220 m3",
    "3.5-1": "THKL ghi trùng tên với dòng dưới ('Đào đất hố móng đá C3') — dòng này có thể là đào đất C3",
    "3.4-3": "Giá trị 1 kg — có thể là số tạm, cần kiểm tra",
    "3.4-4": "THKL để trống khối lượng",
    "3.5-7": "THKL để trống khối lượng",
    "3.5-8": "ĐVT 'm' cho gỗ xẻ — cần kiểm tra (thường là m3)",
    "2.4-2": "Nhãn móng/thân tường chắn trong THKL bị đảo so với bảng thép bản vẽ TC (xem 01_BBS)",
    "2.4-6": "Nhãn móng/thân tường chắn trong THKL bị đảo so với bảng thép bản vẽ TC (xem 01_BBS)",
    "3.2-1": "Khớp bản vẽ PT-10: BT gờ lan can 25MPa 12,56 m3",
    "3.2-2": "3.853,34 kg = tổng cả bảng PT-10, gồm cả lưới D6 699,30 kg đã tính ở dòng 2.1-9 → tính trùng; "
             "thép gờ lan can (L1, L2, L3, L3') theo bản vẽ = 3.154,03 kg",
    "3.3-8": "BBS bản vẽ lớn hơn THKL ~15%: 6 ký hiệu tường cánh (K1, K1A, K8, K8A, K19, K19A) bản vẽ ghi KL "
             "bằng ½ L×n×kg/m — cần TVTK xác nhận",
    "3.3-9": "BBS bản vẽ lớn hơn THKL ~11% (cùng nguyên nhân dòng trên)",
    "2.4-7": "Bảng thép TC dùng 1,980 kg/m cho Φ20 thân tường (đúng là 2,466) → THKL thiếu ~345 kg Φ20",
    "3.2-12": "Khớp bản vẽ PT-09: 2 khe × 7,30 m",
    "3.3-3": "Khớp bản vẽ PD-01: 98 thanh N1 D32 L=1,0m",
}


def build_boq(thkl):
    wb = Workbook()
    ws = wb.active
    ws.title = "QS"
    headers = ["STT", "Mã hiệu", "Nội dung công tác", "ĐVT", "Khối lượng", "Đơn giá", "Thành tiền",
               "Mã dòng THKL", "Ghi chú kiểm tra"]
    title_block(ws, f"BẢNG TIÊN LƯỢNG (BOQ) — {BRIDGE}",
                "Khối lượng lấy nguyên từ THKL G9 - CHUẨN.xlsx / sheet 'Cau 15m'. Cột Mã hiệu (mã định mức TT12/2021) "
                "và Đơn giá (ô vàng) CHƯA CÓ dữ liệu — cần bộ đơn giá địa phương Hà Giang để tính G_XD.", len(headers))
    HR = 5
    header_row(ws, HR, headers, [8, 12, 55, 7, 12, 13, 15, 10, 70])
    r = HR + 1
    last_section = None
    stt = 0
    first_item_row = r
    for it in thkl:
        if it.code == "2.1-2.1":      # dòng "4 dầm" là số lượng cấu kiện, không phải công tác
            continue
        if it.section != last_section:
            put(ws, r, 3, it.section, font=F_BOLD, fill=FILL_SEC)
            for c in (1, 2, 4, 5, 6, 7, 8, 9):
                ws.cell(r, c).fill = FILL_SEC
                ws.cell(r, c).border = BORDER
            last_section = it.section
            r += 1
        stt += 1
        unit = it.unit or ("m3" if it.code == "3.3-5" else "")
        put(ws, r, 1, stt, align=CENTER)
        put(ws, r, 2, None, fill=FILL_INPUT)
        put(ws, r, 3, it.name)
        put(ws, r, 4, unit, align=CENTER)
        put(ws, r, 5, round(it.qty, 4) if it.qty is not None else None, "#,##0.000",
            fill=FILL_WARN if it.qty is None else None)
        put(ws, r, 6, None, "#,##0", fill=FILL_INPUT)
        put(ws, r, 7, f'=IF(F{r}="","",E{r}*F{r})', "#,##0")
        put(ws, r, 8, it.code, align=CENTER)
        note = BOQ_NOTES.get(it.code, "")
        put(ws, r, 9, note, fill=FILL_WARN if note and "Khớp" not in note else None)
        r += 1
    put(ws, r, 3, "CHI PHÍ TRỰC TIẾP T", font=F_BOLD)
    put(ws, r, 7, f"=SUM(G{first_item_row}:G{r-1})", "#,##0", font=F_BOLD)

    ws2 = wb.create_sheet("TONG_HOP_GXD")
    title_block(ws2, "TỔNG HỢP CHI PHÍ XÂY DỰNG G_XD (TT 11/2021/TT-BXD)",
                "Tỷ lệ (ô vàng) CHƯA NHẬP — theo quy định hệ thống, không dùng tỷ lệ mặc định", 4)
    header_row(ws2, 5, ["Khoản mục", "Tỷ lệ (%)", "Cách tính", "Giá trị (đồng)"], [60, 12, 30, 20])
    lines = [
        ("Chi phí trực tiếp T", None, "Σ KL × ĐG", f"=QS!G{r}"),
        ("Chi phí chung (% T)", "rate", "T × tỷ lệ", "=D6*B7/100"),
        ("Chi phí nhà tạm để ở và điều hành thi công (% T)", "rate", "T × tỷ lệ", "=D6*B8/100"),
        ("Chi phí một số công việc không xác định được KL từ thiết kế (% T)", "rate", "T × tỷ lệ", "=D6*B9/100"),
        ("Chi phí gián tiếp GT", None, "C + LT + TT", "=D7+D8+D9"),
        ("Thu nhập chịu thuế tính trước (% (T+GT))", "rate", "(T+GT) × tỷ lệ", "=(D6+D10)*B11/100"),
        ("Giá trị trước thuế G", None, "T + GT + TL", "=D6+D10+D11"),
        ("Thuế giá trị gia tăng (% G)", "rate", "G × thuế suất", "=D12*B13/100"),
        ("CHI PHÍ XÂY DỰNG SAU THUẾ G_XD", None, "G + VAT", "=D12+D13"),
    ]
    for i, (name, rate, how, f) in enumerate(lines, start=6):
        put(ws2, i, 1, name, font=F_BOLD if rate is None else F_REG)
        put(ws2, i, 2, None, "0.00", fill=FILL_INPUT if rate else None)
        put(ws2, i, 3, how)
        put(ws2, i, 4, f, "#,##0", font=F_BOLD if rate is None else F_REG)
    path = os.path.join(OUT_DIR, "02_BOQ_Khoi_Luong_Cau_Khai_Hoang_2.xlsx")
    wb.save(path)
    return path


# ═════════════════════════════════════════════════════════════════════════════
# 03 — TIẾN ĐỘ (đề xuất)
# ═════════════════════════════════════════════════════════════════════════════

# (mã, tên, thời gian ngày làm việc, quan hệ logic, khối lượng chính / căn cứ)
WBS = [
    ("1", "CÔNG TÁC CHUẨN BỊ", None, "", ""),
    ("1.1", "Nhận bàn giao mặt bằng, định vị tim mố A0/B0, lưới mốc VN2000", 5, "", "Bảng tọa độ tr.10"),
    ("1.2", "Rà phá bom mìn khu vực xây dựng cầu", 7, "1.1FS", "Thuyết minh mục 11"),
    ("1.3", "San ủi mặt bằng công trường (đào đất 903,6 m3; đào đá 964,4 m3)", 15, "1.2FS", "THKL 3.1-2, 3.1-3"),
    ("1.4", "Lán trại, kho, bãi gia công thép, hàng rào tôn 55m, biển báo", 10, "1.3SS+5d", "THKL 3.1-5, 3.1-6; bản vẽ tr.34"),
    ("1.5", "Làm bệ đúc dầm T15m (đá dăm 3 m3, BT đệm 4 m3, tà vẹt 50 thanh)", 5, "1.3FS", "THKL 3.4; bản vẽ tr.37"),
    ("2", "MỐ M1", None, "", ""),
    ("2.1", "Đào hố móng mố M1 (đất + đá C3, phá đá bằng máy)", 8, "1.3FS", "THKL 3.5-1..3 (½ khối lượng)"),
    ("2.2", "Khoan cấy thép neo N1 D32 vào đá mố M1 (49 lỗ D42 sâu 50cm)", 4, "2.1FS", "Bản vẽ PD-01"),
    ("2.3", "Nghiệm thu hố móng + đổ BT đệm móng mố M1", 2, "2.2FS", "THKL 3.3-4"),
    ("2.4", "Gia công lắp dựng cốt thép, ván khuôn bệ mố M1", 6, "2.3FS; 1.4FS", "Bản vẽ PD-02"),
    ("2.5", "Đổ bê tông bệ mố M1 30MPa", 1, "2.4FS", "Bản vẽ PD-02 (bảng bệ tr.14: 110 m3)"),
    ("2.6", "Bảo dưỡng bê tông bệ mố M1", 3, "2.5FS", ""),
    ("2.7", "Cốt thép, ván khuôn thân mố + tường cánh M1", 12, "2.6FS", "Bản vẽ PD-04, PD-07"),
    ("2.8", "Đổ bê tông thân mố + tường cánh M1 30MPa", 2, "2.7FS", "THKL 3.3-6 (½)"),
    ("2.9", "Bệ kê gối, u neo dầm mố M1", 4, "2.8FS+7d", "Bản vẽ PD-04"),
    ("3", "MỐ M2", None, "", ""),
    ("3.1", "Đào hố móng mố M2 (đất + đá C3, phá đá bằng máy)", 8, "2.1FS", "THKL 3.5-1..3 (½ khối lượng)"),
    ("3.2", "Khoan cấy thép neo N1 D32 vào đá mố M2 (49 lỗ D42 sâu 50cm)", 4, "3.1FS; 2.2FS", "Bản vẽ PD-01"),
    ("3.3", "Nghiệm thu hố móng + đổ BT đệm móng mố M2", 2, "3.2FS", "THKL 3.3-4"),
    ("3.4", "Gia công lắp dựng cốt thép, ván khuôn bệ mố M2", 6, "3.3FS; 2.4FS", "Bản vẽ PD-02"),
    ("3.5", "Đổ bê tông bệ mố M2 30MPa", 1, "3.4FS", "Bản vẽ PD-02 (bảng bệ tr.15: 123,75 m3)"),
    ("3.6", "Bảo dưỡng bê tông bệ mố M2", 3, "3.5FS", ""),
    ("3.7", "Cốt thép, ván khuôn thân mố + tường cánh M2", 12, "3.6FS; 2.7FS", "Bản vẽ PD-04, PD-07"),
    ("3.8", "Đổ bê tông thân mố + tường cánh M2 30MPa", 2, "3.7FS", "THKL 3.3-6 (½)"),
    ("3.9", "Bệ kê gối, u neo dầm mố M2", 4, "3.8FS+7d; 2.9FS", "Bản vẽ PD-04"),
    ("4", "KẾT CẤU NHỊP", None, "", ""),
    ("4.1", "Gia công cốt thép dầm T L=15m (4 dầm)", 10, "1.4FS", "Bản vẽ PT-05, PT-07"),
    ("4.2", "Đúc dầm T đợt 1 (2 dầm) BT 40MPa", 10, "1.5FS; 4.1FS", "Bản vẽ PT-02..07"),
    ("4.3", "Đúc dầm T đợt 2 (2 dầm) BT 40MPa", 10, "4.2FS+3d", "Bản vẽ PT-02..07"),
    ("4.4", "Bảo dưỡng dầm đợt 2 đến đủ cường độ lao lắp", 21, "4.3FS", "Kiểm tra R mẫu trước khi cẩu"),
    ("4.5", "Lắp đặt gối cao su bản thép 350×500×84 (8 gối)", 2, "2.9FS; 3.9FS", "THKL 3.2-17"),
    ("4.6", "Cẩu lắp 4 dầm T vào vị trí (giữ dầm thẳng đứng)", 4, "4.4FS; 4.5FS", "Bản vẽ TC-03; thuyết minh mục 11"),
    ("4.7", "Dầm ngang (3 dầm ngang) + mối nối dọc bản", 8, "4.6FS", "THKL 2.2"),
    ("4.8", "Tường đỉnh / tường lưng mố M1, M2", 6, "4.6FS", "Bản vẽ PD-04"),
    ("4.9", "Lớp phủ bản mặt cầu BT 35MPa + lưới D6", 8, "4.7FS", "THKL 2.1-8, 2.1-9"),
    ("4.10", "Lắp đặt khe co giãn răng lược (2 khe × 7,3m), vữa không co ngót", 5, "4.9FS; 4.8FS", "Bản vẽ PT-09"),
    ("4.11", "Gờ lan can BT 25MPa + ống thoát nước D150", 8, "4.9FS", "Bản vẽ PT-10"),
    ("4.12", "Lắp dựng lan can thép mạ kẽm (1.270 kg)", 5, "4.11FS", "Bản vẽ PT-11"),
    ("4.13", "Lớp phòng nước + thảm BTNC mặt cầu dày 7cm (100,1 m2)", 3, "4.10FS; 4.12FS", "THKL 2.1-7, 2.1-10"),
    ("5", "ĐƯỜNG ĐẦU CẦU VÀ CÔNG TRÌNH PHỤ TRỢ", None, "", ""),
    ("5.1", "Tường chắn đầu cầu 4 đốt (móng + thân BT 25MPa)", 25, "3.8FS", "Bản vẽ TC-01..10"),
    ("5.2", "Quét nhựa đường mặt tiếp giáp đất của mố (217,4 m2)", 2, "2.8FS; 3.8FS", "THKL 3.3-11"),
    ("5.3", "Đắp đất chọn lọc sau mố + tứ nón K≥0,95 (1.024,6 m3), lớp ≤20cm", 12, "5.2FS; 4.8FS", "THKL 3.5-4, 3.5-5"),
    ("5.4", "Bản quá độ 2 đầu cầu (L=3,0m, dày 0,3m)", 8, "5.3FS", "Bản vẽ PD-08"),
    ("5.5", "Ốp tứ nón BTXM 20MPa trên đá dăm đệm", 8, "5.3FS", "THKL 2.4-9, 2.4-10"),
    ("5.6", "Nền, mặt đường hai đầu cầu (đào nền, khuôn, đá dăm nước, láng nhựa)", 15, "5.4FS; 5.1FS", "THKL mục I"),
    ("5.7", "Rãnh BTXM, gia cố lề, tường hộ lan", 10, "5.6SS+5d", "THKL mục I"),
    ("5.8", "Sơn gờ, vạch sơn, biển tên cầu", 3, "4.13FS; 5.6FS; 5.7FS", "THKL I-15, I-22, 3.2-5"),
    ("6", "HOÀN THIỆN", None, "", ""),
    ("6.1", "Dọn dẹp, hoàn trả mặt bằng, vệ sinh môi trường", 5, "5.8FS; 5.5FS", ""),
    ("6.2", "Nghiệm thu hoàn thành công trình, bàn giao", 3, "6.1FS", ""),
]


def build_schedule():
    wb = Workbook()
    ws = wb.active
    ws.title = "TIEN_DO"
    headers = ["Mã WBS", "Danh mục công tác", "Thời gian (ngày)", "Quan hệ logic", "Căn cứ / khối lượng"]
    title_block(ws, f"TIẾN ĐỘ THI CÔNG ĐỀ XUẤT — {BRIDGE}",
                "Thời gian công tác là ĐỀ XUẤT của nhà thầu (ngày làm việc, nghỉ Chủ nhật), cần TVGS/CĐT chấp thuận. "
                "Giới hạn theo thuyết minh: 12 tháng.", len(headers))
    header_row(ws, 5, headers, [9, 70, 12, 22, 40])
    r = 6
    for code, name, dur, pred, basis in WBS:
        summary = dur is None
        put(ws, r, 1, code, align=CENTER, font=F_BOLD if summary else F_REG, fill=FILL_SEC if summary else None)
        put(ws, r, 2, name, font=F_BOLD if summary else F_REG, fill=FILL_SEC if summary else None)
        if not summary:
            put(ws, r, 3, dur, "0", align=CENTER, fill=FILL_INPUT)
            put(ws, r, 4, pred or "-", align=CENTER)
            put(ws, r, 5, basis)
        else:
            for c in (3, 4, 5):
                ws.cell(r, c).fill = FILL_SEC
        r += 1
    # Excel dùng cho pha schedule: chỉ công việc chi tiết (bộ đọc không có khái niệm công việc tổng hợp)
    xlsx_path = os.path.join(OUT_DIR, "03_Tien_Do_Cau_Khai_Hoang_2.xlsx")
    ws_in = wb.create_sheet("CPM_INPUT")
    header_row(ws_in, 1, ["Mã WBS", "Danh mục công tác", "Thời gian (ngày)", "Quan hệ logic"], [9, 70, 12, 22])
    rr = 2
    for code, name, dur, pred, _ in WBS:
        if dur is None:
            continue
        for c, v in enumerate([code, name, dur, pred or "-"], start=1):
            put(ws_in, rr, c, v)
        rr += 1
    wb.save(xlsx_path)

    # MS Project XML (không có ngày khởi công → MS Project dùng ngày mở file; nhập lại StartDate khi có)
    ns = "http://schemas.microsoft.com/project"
    ET.register_namespace("", ns)
    proj = ET.Element(f"{{{ns}}}Project")
    ET.SubElement(proj, f"{{{ns}}}Name").text = "Tien_Do_Cau_Khai_Hoang_2"
    ET.SubElement(proj, f"{{{ns}}}Title").text = BRIDGE
    ET.SubElement(proj, f"{{{ns}}}MinutesPerDay").text = "480"
    tasks = ET.SubElement(proj, f"{{{ns}}}Tasks")
    uid_of = {}
    for i, (code, *_rest) in enumerate(WBS, start=1):
        uid_of[code] = i
    import re as _re
    for code, name, dur, pred, basis in WBS:
        t = ET.SubElement(tasks, f"{{{ns}}}Task")
        ET.SubElement(t, f"{{{ns}}}UID").text = str(uid_of[code])
        ET.SubElement(t, f"{{{ns}}}ID").text = str(uid_of[code])
        ET.SubElement(t, f"{{{ns}}}Name").text = name
        ET.SubElement(t, f"{{{ns}}}WBS").text = code
        ET.SubElement(t, f"{{{ns}}}OutlineNumber").text = code
        ET.SubElement(t, f"{{{ns}}}OutlineLevel").text = str(code.count(".") + 1)
        ET.SubElement(t, f"{{{ns}}}Summary").text = "1" if dur is None else "0"
        if dur is not None:
            ET.SubElement(t, f"{{{ns}}}Duration").text = f"PT{dur*8}H0M0S"
            ET.SubElement(t, f"{{{ns}}}DurationFormat").text = "7"
            for p in filter(None, (x.strip() for x in pred.split(";"))):
                m = _re.fullmatch(r"([\d.]+)(FS|SS|FF|SF)([+-]\d+)?d?", p)
                link = ET.SubElement(t, f"{{{ns}}}PredecessorLink")
                ET.SubElement(link, f"{{{ns}}}PredecessorUID").text = str(uid_of[m.group(1)])
                ET.SubElement(link, f"{{{ns}}}Type").text = {"FF": "0", "FS": "1", "SF": "2", "SS": "3"}[m.group(2)]
                lag_days = int(m.group(3) or 0)
                ET.SubElement(link, f"{{{ns}}}LinkLag").text = str(lag_days * 4800)   # 1/10 phút
                ET.SubElement(link, f"{{{ns}}}LagFormat").text = "7"
            if basis:
                ET.SubElement(t, f"{{{ns}}}Notes").text = basis
    xml_path = os.path.join(OUT_DIR, "03_Tien_Do_Cau_Khai_Hoang_2.xml")
    ET.ElementTree(proj).write(xml_path, encoding="utf-8", xml_declaration=True)
    return xlsx_path, xml_path


# ═════════════════════════════════════════════════════════════════════════════
# 04 — KCS: danh mục nghiệm thu + tần suất thí nghiệm
# ═════════════════════════════════════════════════════════════════════════════

# (mã WBS liên kết, loại biên bản, nội dung nghiệm thu, căn cứ kỹ thuật)
ACCEPTANCES = [
    ("—", "Vật liệu đầu vào", "Thép thanh CB400-V các loại Ø6..Ø32 (chứng chỉ + thí nghiệm kéo, uốn)", "TCVN 1651-2:2018"),
    ("—", "Vật liệu đầu vào", "Xi măng, cát, đá dăm, nước, phụ gia cho bê tông", "TCVN 6260, TCVN 7570, TCVN 4506"),
    ("—", "Vật liệu đầu vào", "Thiết kế cấp phối bê tông 10/20/25/30/35/40MPa, vữa không co ngót 50MPa", "TCVN 10306:2014"),
    ("—", "Vật liệu đầu vào", "Gối cao su bản thép 350×500×84mm (8 cái)", "AASHTO M251-92"),
    ("—", "Vật liệu đầu vào", "Khe co giãn răng lược, thép tấm không gỉ, bu lông M12", "AASHTO M297-2006"),
    ("—", "Vật liệu đầu vào", "Lan can thép mạ kẽm nhúng nóng, ống thoát nước gang D150", "TCVN 5408:2007"),
    ("—", "Vật liệu đầu vào", "Đất đắp chọn lọc sau mố, đá dăm đệm", "QĐ 3095/QĐ-BGTVT"),
    ("1.1", "Công việc xây dựng", "Định vị tim mố A0, B0 và góc bệ (tọa độ VN2000)", "Bảng tọa độ tr.10"),
    ("2.3", "Công việc xây dựng", "Hố móng mố M1: cao độ đáy móng, ngàm vào đá gốc ≥ 50cm", "Thuyết minh 7.1; TC-02"),
    ("2.2", "Công việc xây dựng", "Khoan cấy thép neo D32 mố M1 (vị trí, chiều sâu lỗ, thí nghiệm kéo nhổ)", "Bản vẽ PD-01"),
    ("2.4", "Công việc xây dựng", "Cốt thép + ván khuôn bệ mố M1 (Hold point trước đổ BT)", "TCVN 4453:1995"),
    ("2.5", "Công việc xây dựng", "Bê tông bệ mố M1", "TCVN 4453:1995"),
    ("2.7", "Công việc xây dựng", "Cốt thép + ván khuôn thân mố, tường cánh M1", "TCVN 4453:1995"),
    ("2.8", "Công việc xây dựng", "Bê tông thân mố, tường cánh M1", "TCVN 4453:1995"),
    ("2.9", "Công việc xây dựng", "Bệ kê gối mố M1 (cao độ, vị trí tấm đệm)", "Bản vẽ PD-04"),
    ("3.3", "Công việc xây dựng", "Hố móng mố M2: cao độ đáy móng, ngàm vào đá gốc ≥ 50cm", "Thuyết minh 7.1; TC-02"),
    ("3.2", "Công việc xây dựng", "Khoan cấy thép neo D32 mố M2 (vị trí, chiều sâu lỗ, thí nghiệm kéo nhổ)", "Bản vẽ PD-01"),
    ("3.4", "Công việc xây dựng", "Cốt thép + ván khuôn bệ mố M2 (Hold point trước đổ BT)", "TCVN 4453:1995"),
    ("3.5", "Công việc xây dựng", "Bê tông bệ mố M2", "TCVN 4453:1995"),
    ("3.7", "Công việc xây dựng", "Cốt thép + ván khuôn thân mố, tường cánh M2", "TCVN 4453:1995"),
    ("3.8", "Công việc xây dựng", "Bê tông thân mố, tường cánh M2", "TCVN 4453:1995"),
    ("3.9", "Công việc xây dựng", "Bệ kê gối mố M2 (cao độ, vị trí tấm đệm)", "Bản vẽ PD-04"),
    ("3.9", "Giai đoạn", "Hoàn thành kết cấu phần dưới (mố M1, M2)", "NĐ 06/2021/NĐ-CP"),
    ("1.5", "Công việc xây dựng", "Bệ đúc dầm T15m", "Bản vẽ tr.37"),
    ("4.2", "Công việc xây dựng", "Cốt thép + ván khuôn + bê tông dầm T đợt 1 (2 dầm)", "TCVN 4453:1995; PT-02..07"),
    ("4.3", "Công việc xây dựng", "Cốt thép + ván khuôn + bê tông dầm T đợt 2 (2 dầm)", "TCVN 4453:1995; PT-02..07"),
    ("4.5", "Công việc xây dựng", "Lắp đặt gối cầu", "AASHTO M251-92"),
    ("4.6", "Công việc xây dựng", "Lao lắp dầm T vào vị trí", "TCVN 11823:2017; TC-03"),
    ("4.7", "Công việc xây dựng", "Dầm ngang + mối nối dọc bản", "TCVN 4453:1995"),
    ("4.9", "Công việc xây dựng", "Lớp phủ bản mặt cầu BT 35MPa + lưới thép D6", "TCVN 4453:1995; PT-10"),
    ("4.10", "Công việc xây dựng", "Khe co giãn", "Bản vẽ PT-09"),
    ("4.11", "Công việc xây dựng", "Gờ lan can + ống thoát nước", "Bản vẽ PT-10"),
    ("4.12", "Công việc xây dựng", "Lan can thép", "Bản vẽ PT-11"),
    ("4.13", "Công việc xây dựng", "Lớp phòng nước + BTNC mặt cầu", "TCVN 8819:2011"),
    ("4.13", "Giai đoạn", "Hoàn thành kết cấu phần trên", "NĐ 06/2021/NĐ-CP"),
    ("5.1", "Công việc xây dựng", "Tường chắn đầu cầu (móng, thân)", "Bản vẽ TC-01..10"),
    ("5.3", "Công việc xây dựng", "Đắp đất sau mố, tứ nón (từng lớp ≤20cm, K≥0,95; lòng mố K≥0,98)", "TCVN 9436:2012; QĐ 3095"),
    ("5.4", "Công việc xây dựng", "Bản quá độ", "Bản vẽ PD-08"),
    ("5.5", "Công việc xây dựng", "Ốp tứ nón BTXM 20MPa", "Thuyết minh 7.3"),
    ("5.6", "Công việc xây dựng", "Nền, móng, mặt đường hai đầu cầu", "TCVN 9436:2012; TCVN 8863:2011"),
    ("5.7", "Công việc xây dựng", "Rãnh, gia cố lề, tường hộ lan", "THKL mục I"),
    ("5.8", "Công việc xây dựng", "Sơn vạch, biển tên cầu", "QCVN 41:2019/BGTVT"),
    ("6.2", "Hoàn thành", "Nghiệm thu hoàn thành công trình đưa vào sử dụng", "NĐ 06/2021/NĐ-CP"),
]

# Bê tông để tính số tổ mẫu: (cấu kiện, cấp BT, V m3, nguồn, số lần đổ tối thiểu, V một tổ mẫu, căn cứ tần suất)
CONCRETE = [
    ("Bệ mố M1", "30MPa", 110.00, "Bản vẽ PD-02", 1, 100, "Móng lớn: 1 tổ/100 m3"),
    ("Bệ mố M2", "30MPa", 123.75, "Bản vẽ PD-02", 1, 100, "Móng lớn: 1 tổ/100 m3"),
    ("Thân mố + tường cánh M1+M2", "30MPa", 169.29, "THKL 3.3-6", 2, 20, "Kết cấu thành mỏng: 1 tổ/20 m3"),
    ("Dầm T L=15m (4 dầm)", "40MPa", 37.00, "THKL 2.1-1", 4, 20, "Mỗi dầm ≥ 1 tổ"),
    ("Dầm ngang", "30MPa", 2.86, "THKL 2.2-1", 1, 20, "Mỗi đợt đổ ≥ 1 tổ"),
    ("Lớp phủ bản mặt cầu", "35MPa", 14.44, "THKL 2.1-8", 1, 20, "1 tổ/20 m3"),
    ("Gờ lan can", "25MPa", 12.56, "THKL 3.2-1", 2, 20, "Mỗi đợt đổ ≥ 1 tổ"),
    ("Bản quá độ (2 bản)", "25MPa", 12.53, "THKL 2.3-2", 2, 20, "Mỗi bản ≥ 1 tổ"),
    ("Móng tường chắn", "25MPa", 111.69, "THKL 2.4-1", 4, 100, "Mỗi đốt ≥ 1 tổ"),
    ("Thân tường chắn", "25MPa", 67.37, "THKL 2.4-5", 4, 20, "Mỗi đốt ≥ 1 tổ"),
    ("Ốp tứ nón", "20MPa", 29.91, "THKL 2.4-9", 1, 100, "1 tổ/100 m3"),
    ("Vữa không co ngót khe co giãn", "50MPa", 2.20, "THKL 3.2-14", 2, 20, "Mỗi khe ≥ 1 tổ"),
]


def build_kcs(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "DANH_MUC_NGHIEM_THU"
    headers = ["STT", "Mã WBS (tiến độ)", "Loại nghiệm thu", "Nội dung nghiệm thu", "Căn cứ kỹ thuật",
               "Ngày dự kiến", "Số biên bản", "Ghi chú"]
    title_block(ws, f"DANH MỤC BIÊN BẢN NGHIỆM THU — {BRIDGE}",
                "Ngày dự kiến điền sau khi có ngày khởi công và chạy CPM (lấy ngày hoàn thành EF của mã WBS). "
                "Biểu mẫu theo NĐ 06/2021/NĐ-CP và quy định hiện hành của CĐT.", len(headers))
    header_row(ws, 5, headers, [5, 10, 18, 70, 28, 13, 13, 25])
    for i, (code, kind, content, basis) in enumerate(ACCEPTANCES, start=1):
        r = 5 + i
        put(ws, r, 1, i, align=CENTER)
        put(ws, r, 2, code, align=CENTER)
        put(ws, r, 3, kind, fill=FILL_SEC if kind in ("Giai đoạn", "Hoàn thành") else None)
        put(ws, r, 4, content)
        put(ws, r, 5, basis)
        put(ws, r, 6, None, fill=FILL_INPUT)
        put(ws, r, 7, None, fill=FILL_INPUT)
        put(ws, r, 8, None)

    ws2 = wb.create_sheet("TAN_SUAT_THI_NGHIEM")
    headers2 = ["STT", "Đối tượng thí nghiệm", "Chỉ tiêu", "Khối lượng", "ĐVT", "Nguồn khối lượng",
                "Tần suất (1 mẫu / ...)", "Số lần đổ / lô tối thiểu", "Số mẫu (tổ)", "Căn cứ tần suất"]
    title_block(ws2, "KẾ HOẠCH TẦN SUẤT THÍ NGHIỆM KIỂM TRA CHẤT LƯỢNG",
                "Số mẫu = MAX(ROUNDUP(khối lượng / tần suất), số lần đổ tối thiểu). Tần suất cần TVGS chấp thuận.",
                len(headers2))
    header_row(ws2, 5, headers2, [5, 34, 26, 12, 7, 18, 12, 12, 11, 40])
    r = 6
    n = 0
    put(ws2, r, 2, "BÊ TÔNG — nén mẫu R7/R28 (tổ 3 mẫu trụ 15×30cm)", font=F_BOLD, fill=FILL_SEC)
    r += 1
    first_c = r
    for comp, grade, vol, src, pours, per, basis in CONCRETE:
        n += 1
        put(ws2, r, 1, n, align=CENTER)
        put(ws2, r, 2, f"{comp} ({grade})")
        put(ws2, r, 3, "Cường độ nén R7, R28; độ sụt mỗi xe")
        put(ws2, r, 4, vol, "#,##0.00")
        put(ws2, r, 5, "m3", align=CENTER)
        put(ws2, r, 6, src)
        put(ws2, r, 7, per, "#,##0", fill=FILL_INPUT)
        put(ws2, r, 8, pours, "#,##0", fill=FILL_INPUT)
        put(ws2, r, 9, f"=MAX(ROUNDUP(D{r}/G{r},0),H{r})", "#,##0", font=F_BOLD)
        put(ws2, r, 10, f"TCVN 4453:1995 mục 7.1.7 — {basis}")
        r += 1
    put(ws2, r, 2, "Cộng số tổ mẫu bê tông", font=F_BOLD)
    put(ws2, r, 9, f"=SUM(I{first_c}:I{r-1})", "#,##0", font=F_BOLD)
    r += 2
    put(ws2, r, 2, "CỐT THÉP — kéo, uốn (tổ 3 mẫu kéo + 3 mẫu uốn)", font=F_BOLD, fill=FILL_SEC)
    r += 1
    first_s = r
    by_d = defaultdict(float)
    for b in rows:
        by_d[b.dia] += b.kg_total
    for d in sorted(by_d):
        n += 1
        put(ws2, r, 1, n, align=CENTER)
        put(ws2, r, 2, f"Thép Ø{d} CB400-V")
        put(ws2, r, 3, "Giới hạn chảy, bền, độ giãn dài, uốn")
        put(ws2, r, 4, round(by_d[d] / 1000, 3), "#,##0.000")
        put(ws2, r, 5, "tấn", align=CENTER)
        put(ws2, r, 6, "01_BBS (toàn cầu)")
        put(ws2, r, 7, 50, "#,##0", fill=FILL_INPUT)
        put(ws2, r, 8, 1, "#,##0", fill=FILL_INPUT)
        put(ws2, r, 9, f"=MAX(ROUNDUP(D{r}/G{r},0),H{r})", "#,##0", font=F_BOLD)
        put(ws2, r, 10, "TCVN 1651-2:2018 — mỗi lô ≤ 50 tấn cùng Ø, cùng mác, cùng nhà sản xuất")
        r += 1
    put(ws2, r, 2, "Cộng số tổ mẫu thép", font=F_BOLD)
    put(ws2, r, 9, f"=SUM(I{first_s}:I{r-1})", "#,##0", font=F_BOLD)
    r += 2
    put(ws2, r, 2, "ĐẤT ĐẮP, NEO, THIẾT BỊ", font=F_BOLD, fill=FILL_SEC)
    r += 1
    others = [
        ("Đất đắp sau mố + tứ nón", "Độ chặt K (mỗi lớp)", 1024.63, "m3", "THKL 3.5-4 + 3.5-5", 200, 10,
         "TCVN 9436:2012 — mỗi lớp ≥ 1 điểm; tối thiểu theo số lớp đầm (đề xuất 10 lớp)"),
        ("Đất đắp chọn lọc", "Thành phần hạt, đầm nén tiêu chuẩn", 1024.63, "m3", "THKL 3.5-4 + 3.5-5", 1000, 1,
         "QĐ 3095/QĐ-BGTVT — mỗi nguồn vật liệu"),
        ("Thép neo D32 khoan cấy vào đá", "Kéo nhổ thử", 98, "thanh", "Bản vẽ PD-01", 50, 3,
         "Đề xuất 2% số neo, không ít hơn 3 neo — TVTK/TVGS quyết định"),
        ("Gối cao su bản thép", "Chứng chỉ + kiểm tra kích thước, độ cứng", 8, "cái", "THKL 3.2-17", 8, 1,
         "AASHTO M251 — mỗi lô"),
        ("Vữa không co ngót 50MPa", "Cường độ nén mẫu lập phương", 2.2, "m3", "THKL 3.2-14", 1, 2,
         "Mỗi khe co giãn ≥ 1 tổ"),
    ]
    for name, what, q, u, src, per, mn, basis in others:
        n += 1
        put(ws2, r, 1, n, align=CENTER)
        put(ws2, r, 2, name)
        put(ws2, r, 3, what)
        put(ws2, r, 4, q, "#,##0.00")
        put(ws2, r, 5, u, align=CENTER)
        put(ws2, r, 6, src)
        put(ws2, r, 7, per, "#,##0", fill=FILL_INPUT)
        put(ws2, r, 8, mn, "#,##0", fill=FILL_INPUT)
        put(ws2, r, 9, f"=MAX(ROUNDUP(D{r}/G{r},0),H{r})", "#,##0", font=F_BOLD)
        put(ws2, r, 10, basis)
        r += 1
    path = os.path.join(OUT_DIR, "04_KCS_Nghiem_Thu_Thi_Nghiem_Cau_Khai_Hoang_2.xlsx")
    wb.save(path)
    return path


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    os.makedirs(OUT_DIR, exist_ok=True)
    rows, log = load_bbs()
    log = log + check_bbs_weights(rows)
    thkl = load_thkl()
    bbs_path, steel_check = build_bbs(rows, log, thkl)
    boq_path = build_boq(thkl)
    sch_xlsx, sch_xml = build_schedule()
    kcs_path = build_kcs(rows)
    print("[V] Đã tạo:")
    for p in (bbs_path, boq_path, sch_xlsx, sch_xml, kcs_path):
        print("   ", p)
    print(f"\nBBS: {len(rows)} dòng, {sum(b.total_qty for b in rows):,} thanh, "
          f"{sum(b.kg_total for b in rows)/1000:,.3f} tấn")
    print("\nĐối chiếu thép BBS ↔ THKL:")
    for name, grp, bbs_kg, thkl_kg, ok in steel_check:
        pct = (bbs_kg - thkl_kg) / thkl_kg if thkl_kg else float("nan")
        print(f"  {'KHỚP' if ok else 'LỆCH'}  {name:48s} {grp:9s} BBS {bbs_kg:10.2f}  THKL {thkl_kg:10.2f}  {pct:+.1%}")


if __name__ == "__main__":
    main()

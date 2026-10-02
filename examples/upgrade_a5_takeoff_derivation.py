# -*- coding: utf-8 -*-
"""
Nâng cấp bảng diễn giải khối lượng cống hộp A5 trong mọi workbook ví dụ: đưa kích thước lòng cống,
số khoang, cạnh vút và quy ước ván khuôn thành ô đầu vào có nhãn (xem tools/takeoff_rules.py).

Chạy:  python examples/upgrade_a5_takeoff_derivation.py [thư_mục_gốc=examples]
Idempotent: workbook đã nâng cấp (K8 có giá trị) sẽ được bỏ qua.
"""
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import openpyxl

from tools.takeoff_rules import apply_culvert_derivation

MARKER = "Bê tông thân B20 cống đơn 2.0x2.0m"


def main(base: str) -> int:
    changed = 0
    for dp, _, files in os.walk(base):
        for fn in files:
            if not fn.lower().endswith(".xlsx") or fn.startswith("~$"):
                continue
            path = os.path.join(dp, fn)
            wb = openpyxl.load_workbook(path)
            touched = False
            for ws in wb.worksheets:
                if str(ws["B8"].value or "").startswith(MARKER) and ws["K8"].value is None:
                    apply_culvert_derivation(ws)
                    touched = True
                    print(f"  nâng cấp: {os.path.relpath(path, base)} [{ws.title}]")
            if touched:
                wb.save(path)
                changed += 1
    print(f"Đã nâng cấp {changed} workbook.")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "examples")))

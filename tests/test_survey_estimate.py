# -*- coding: utf-8 -*-
"""
Golden test dự toán khảo sát xây dựng — đối chiếu với một dự toán THẬT đã thẩm định.

Fixture tests/fixtures/du_toan_khao_sat_trich.xls: trích 3 sheet (bảng đo bóc khối lượng, bảng
tổng hợp dự toán hạng mục, sheet Hệ số) của một file dự toán khảo sát đã thẩm định, giữ nguyên số
liệu, bỏ tên công trình và người lập. Các con số kỳ vọng dưới đây lấy từ file gốc, không tính từ
chính chương trình.

Chạy thêm với file gốc đầy đủ:  AEC_TEST_SURVEY_XLS=/đường/dẫn/du_toan.xls
"""

import argparse
import contextlib
import io
import os
import tempfile
import unittest

from tools.money import round_vnd
from tools.qs_loader import QSLoadError, _find_header, _parse_items, load_qs
from tools.survey_estimate import SurveyEstimate, _read_summary, load_survey_estimate

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIXTURE = os.path.join(ROOT, "tests", "fixtures", "du_toan_khao_sat_trich.xls")

# Số ghi trong file đã thẩm định (sheet "TH dự toán hạng mục")
APPROVED = {
    "VL": 59_915_021, "NC": 1_041_693_266, "M": 70_584_811, "T": 1_172_193_098,
    "C": 677_100_622.9, "TL": 110_957_623.254, "Gks": 1_960_251_344.154,
    "Glpa": 29_403_770.1623, "Glbc": 49_006_283.6039, "Ghmc": 107_813_823.9285,
    "G": 2_146_475_221.8486, "GTGT": 214_647_522.1849, "Gxd": 2_361_122_744.0335,
    "Gdp": 118_056_137.2017, "Tong": 2_479_178_881.2352,
}
APPROVED_ROUNDED = 2_479_179_000
APPROVED_RATES = {"chung_nc": 0.65, "tl": 0.06, "lpa": 0.015, "lbc": 0.025, "co": 0.025,
                  "dc": 0.02, "atgt": 0.01, "bh": 0.0, "vat": 0.1, "dp": 0.05}


def _assert_matches_approved(test, est):
    res = est.compute()
    for key, expected in APPROVED.items():
        test.assertLess(abs(float(res[key]) - expected), 0.01, key)
    test.assertEqual(est.total_rounded(1000), APPROVED_ROUNDED)
    test.assertEqual(est.errors, [])
    # chỉ còn cảnh báo 8 công tác khối lượng 0 (khoan dưới nước, thí nghiệm cầu, SPT)
    test.assertEqual(len(est.warnings), 8, est.warnings)
    test.assertTrue(all("khối lượng = 0" in w for w in est.warnings))


class SurveyGoldenTest(unittest.TestCase):
    def test_matches_approved_estimate(self):
        est = load_survey_estimate(FIXTURE)
        self.assertEqual(len(est.qs.items), 35)
        self.assertEqual(est.rates, APPROVED_RATES)
        self.assertTrue(all(src.endswith("#TH dự toán hạng mục") for src in est.rate_sources.values()))
        self.assertEqual(est.file_total_rounded, APPROVED_ROUNDED)
        _assert_matches_approved(self, est)

    def test_every_line_matches_file_amounts(self):
        est = load_qs(FIXTURE)
        self.assertEqual(est.errors, [])
        self.assertFalse([w for w in est.warnings if "thành tiền trong file" in w])
        first = est.items[1]                       # CF.11220: 6 điểm
        self.assertEqual((first.code, first.quantity), ("CF.11220", 6))
        self.assertEqual((first.price_vl, first.price_nc, first.price_m), (176192, 7488784, 2155258))
        self.assertEqual(first.file_amount, 6 * (176192 + 7488784 + 2155258))

    def test_summary_values_cross_checked(self):
        est = load_survey_estimate(FIXTURE)
        est.file_values["G"] += 5
        est.compute()
        self.assertTrue(any(w.startswith("G ghi trong file") for w in est.warnings))

    def test_construction_formula_not_applied_to_survey(self):
        """load_qs (G_XD xây lắp) không được đọc 'NC x 65%' thành chi phí chung 65% của T."""
        est = load_qs(FIXTURE)
        self.assertNotIn("chung", est.rates)
        self.assertTrue(any("dự toán khảo sát" in w for w in est.warnings))
        with self.assertRaises(QSLoadError):
            est.compute()

    def test_cli(self):
        import run_state_graph
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            code = run_state_graph.run_survey(argparse.Namespace(survey=FIXTURE, survey_sheet=None))
        self.assertEqual(code, 0)
        self.assertIn("2,479,179,000", out.getvalue())


class SurveyRulesTest(unittest.TestCase):
    def _est(self, rows):
        est = SurveyEstimate(qs=None)
        _read_summary(est, "test", rows)
        return est

    def test_rate_with_unexpected_base_is_an_error(self):
        est = self._est([["II", "CHI PHÍ CHUNG", "T x 5%", 100.0, "C"],
                         ["", "Chi phí lập phương án", "Gks x 1,5%", 10.0, "Glpa"]])
        self.assertNotIn("chung_nc", est.rates)
        self.assertEqual(est.rates["lpa"], 0.015)
        self.assertTrue(any("cần cơ sở NC" in e for e in est.errors))

    def test_missing_rate_is_reported_not_defaulted(self):
        est = load_survey_estimate(FIXTURE)
        del est.rates["dp"]
        with self.assertRaises(QSLoadError) as cm:
            est.compute()
        self.assertIn("dự phòng", str(cm.exception))

    def test_hs_sheet_mismatch_warns(self):
        import openpyxl
        import xlrd
        book = xlrd.open_workbook(FIXTURE)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for sh in book.sheets():
            ws = wb.create_sheet(sh.name)
            for r in range(sh.nrows):
                for c in range(sh.ncols):
                    v = sh.cell_value(r, c)
                    if v != "":
                        ws.cell(r + 1, c + 1, v)
        for row in wb["Hệ số"].iter_rows():
            if row[1].value == "hsTL":
                row[2].value = 0.055
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "ks.xlsx")
            wb.save(path)
            est = load_survey_estimate(path)
        self.assertEqual(est.rates["tl"], 0.06)            # bảng tổng hợp là nguồn chính
        self.assertTrue(any("sheet Hệ số 5.50%" in w for w in est.warnings), est.warnings)


class TwoRowHeaderTest(unittest.TestCase):
    ROWS = [
        ["BẢNG TÍNH TOÁN, ĐO BÓC KHỐI LƯỢNG"],
        ["STT", "Mã hiệu công tác", "Danh mục công tác đo bóc", "Đơn vị", "Khối lượng cấu kiện",
         "Khối lượng", "Đơn giá", "", "", "Thành tiền", "", "", "Đơn giá"],
        ["", "", "", "", "", "", "Vật liệu", "Nhân công", "Máy thi công", "Vật liệu", "Nhân công",
         "Máy thi công"],
        ["", "I", "Khảo sát tuyến", "", "", "", "", "", "", "", "", "", ""],
        ["1", "CH.11130", "Đo vẽ mặt cắt dọc", "100m", "", "2", "100", "1000", "10", "200", "2000", "20",
         "BO_DON_GIA_2018"],
    ]

    def test_merged_header_columns(self):
        last, cols = _find_header(self.ROWS)
        self.assertEqual(last, 2)
        self.assertEqual(cols["quantity"], 5)            # "Khối lượng", không phải "Khối lượng cấu kiện"
        self.assertEqual((cols["price_vl"], cols["price_nc"], cols["price_m"]), (6, 7, 8))
        self.assertEqual((cols["amount_vl"], cols["amount_nc"], cols["amount_m"]), (9, 10, 11))

    def test_item_uses_component_prices_when_price_column_is_text(self):
        est = _parse_items([list(r) for r in self.ROWS], source="test")
        self.assertEqual(est.errors, [])
        item = est.items[0]
        self.assertEqual((item.quantity, item.unit_price, item.amount), (2, 1110, 2220))
        self.assertEqual((item.price_vl, item.price_nc, item.price_m), (100, 1000, 10))
        self.assertEqual(item.file_amount, 2220)
        self.assertEqual(item.section, "Khảo sát tuyến")

    def test_title_row_is_not_a_header(self):
        rows = [["Tên bảng: Khối lượng = Dài x Rộng"],
                ["TT", "Nội dung công tác", "ĐVT", "Khối lượng", "Đơn giá", "Thành tiền"],
                ["1", "Đào đất", "m3", "10", "5", "50"]]
        last, cols = _find_header(rows)
        self.assertEqual((last, cols["quantity"], cols["price"]), (1, 3, 4))


@unittest.skipUnless(os.environ.get("AEC_TEST_SURVEY_XLS"), "đặt AEC_TEST_SURVEY_XLS để chạy với file gốc")
class FullSurveyFileTest(unittest.TestCase):
    def test_full_file_matches(self):
        est = load_survey_estimate(os.environ["AEC_TEST_SURVEY_XLS"])
        _assert_matches_approved(self, est)
        self.assertEqual(round_vnd(est.results["G"]), 2_146_475_222)


if __name__ == "__main__":
    unittest.main()

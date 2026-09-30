# -*- coding: utf-8 -*-
"""
UNIT TESTS — DYNAMIC CPM SCHEDULE BUILDER
Kiểm thử bộ công cụ sinh tiến độ thi công 100% công thức động & Native Excel Charts.
"""

import os
import unittest
import tempfile
import openpyxl

from tools.dynamic_schedule_builder import DynamicScheduleBuilder


class TestDynamicCPMBuilder(unittest.TestCase):

    def setUp(self):
        self.builder = DynamicScheduleBuilder()
        self.temp_dir = tempfile.mkdtemp()
        self.output_file = os.path.join(self.temp_dir, "test_schedule.xlsx")

    def tearDown(self):
        if os.path.exists(self.output_file):
            os.remove(self.output_file)
        if os.path.exists(self.temp_dir):
            os.rmdir(self.temp_dir)

    def test_build_cum_b9_workbook_sheets_and_charts(self):
        """Kiểm tra tạo đầy đủ 6 Sheet và 4 biểu đồ Native Excel."""
        self.builder.build_cum_b9_workbook(self.output_file)
        self.assertTrue(os.path.exists(self.output_file))

        wb = openpyxl.load_workbook(self.output_file)
        expected_sheets = [
            "01_THONG_SO_DU_AN",
            "02_DINH_MUC_CA_MAY_VA_DAU",
            "03_TIEN_DO_GANTT_CPM",
            "04_TONG_HOP_CA_MAY_VA_DAU",
            "05_NHU_CAU_VAT_TU_CHINH",
            "06_SO_SANH_DINH_MUC_VS_THUC_TE",
        ]
        self.assertEqual(wb.sheetnames, expected_sheets)

        # Kiểm tra biểu đồ trong từng sheet
        self.assertEqual(len(wb["03_TIEN_DO_GANTT_CPM"]._charts), 1)
        self.assertEqual(len(wb["04_TONG_HOP_CA_MAY_VA_DAU"]._charts), 1)
        self.assertEqual(len(wb["05_NHU_CAU_VAT_TU_CHINH"]._charts), 1)
        self.assertEqual(len(wb["06_SO_SANH_DINH_MUC_VS_THUC_TE"]._charts), 1)

    def test_no_corrupting_formulas_in_text_cells(self):
        """Kiểm tra không có ô chú thích văn bản nào bắt đầu bằng '=' gây lỗi cho Excel."""
        self.builder.build_cum_b9_workbook(self.output_file)
        wb = openpyxl.load_workbook(self.output_file)

        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for r in range(1, ws.max_row + 1):
                for c in range(1, ws.max_column + 1):
                    val = ws.cell(r, c).value
                    if isinstance(val, str) and val.startswith("="):
                        # Đảm bảo không chứa các từ ghi chú tiếng Việt hoặc gạch ngang
                        self.assertNotIn(" — ", val, f"Ô rác công thức tại {sheet} R{r}C{c}: {val}")
                        self.assertNotIn("Tự động", val, f"Ô rác công thức tại {sheet} R{r}C{c}: {val}")

    def test_sheet06_comparison_formulas(self):
        """Kiểm tra Sheet 06 liên kết động sang Sheet 04."""
        self.builder.build_cum_b9_workbook(self.output_file)
        wb = openpyxl.load_workbook(self.output_file, data_only=False)
        ws6 = wb["06_SO_SANH_DINH_MUC_VS_THUC_TE"]

        # Ô E6 phải liên kết sang Sheet 04 ô G6
        self.assertEqual(ws6["E6"].value, "='04_TONG_HOP_CA_MAY_VA_DAU'!$G$6")
        # Ô G6 (chênh lệch) = F6 - E6
        self.assertEqual(ws6["G6"].value, "=F6-E6")
        # Ô H6 (tỷ lệ %) = IF(E6>0, (F6-E6)/E6, 0)
        self.assertEqual(ws6["H6"].value, "=IF(E6>0, (F6-E6)/E6, 0)")


if __name__ == "__main__":
    unittest.main()

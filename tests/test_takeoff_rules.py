# -*- coding: utf-8 -*-
"""Kiểm thử diễn giải khối lượng cống hộp: hình học tính tay, và đối chiếu công thức trong sheet với module."""

import math
import os
import sys
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import openpyxl

from tools.excel_eval import WorkbookEvaluator
from tools.takeoff_rules import BoxCulvert, CULVERT_ROWS

A5 = os.path.join(ROOT, "examples", "HO_SO_CONG_HOP_TUYEN_A5", "BO_HO_SO_02_VI_MO_CHUYEN_SAU_14_BO",
                  "03_QS_Dien_Giai_Chi_Tiet_Takeoff_Cong_A5.xlsx")
CASES = {  # (cống, thể tích BT trên 1 m tính tay)
    8: (BoxCulvert(2.5, 2.5, 2.0, 2.0, 1, 0.20), 2.33),        # 6.25 − 4 + 4·½·0.04          = 2.330
    9: (BoxCulvert(3.6, 3.6, 3.0, 3.0, 1, 0.25), 4.085),       # 12.96 − 9 + 4·½·0.0625       = 4.085
    10: (BoxCulvert(6.9, 3.6, 3.0, 3.0, 2, 0.25), 7.09),       # 24.84 − 18 + 8·½·0.0625      = 7.090
}


class BoxCulvertGeometryTest(unittest.TestCase):

    def test_concrete_area_by_hand(self):
        for row, (bc, expected) in CASES.items():
            self.assertAlmostEqual(bc.concrete_area(), expected, places=9, msg=f"dòng {row}")

    def test_inner_perimeter_haunch_replaces_two_legs_by_hypotenuse(self):
        # lòng 2×2, vút 0.2: 8 − 4·0.2·(2 − √2) = 8 − 0.8·0.585786... = 7.531371
        self.assertAlmostEqual(BoxCulvert(2.5, 2.5, 2, 2, 1, 0.2).inner_perimeter(), 7.5313708499, places=9)
        self.assertEqual(BoxCulvert(2.5, 2.5, 2, 2, 1, 0.0).inner_perimeter(), 8.0)   # không vút = chu vi chữ nhật

    def test_formwork_faces_by_hand(self):
        bc = BoxCulvert(2.5, 2.5, 2, 2, 1, 0.2)
        fw = bc.formwork_per_m(11.3)                     # mặc định: trong + hông ngoài + đầu đốt
        self.assertAlmostEqual(fw["trong"], 7.5313708499, places=9)
        self.assertEqual(fw["hong_ngoai"], 5.0)           # 2 × 2.5
        self.assertAlmostEqual(fw["dau_dot"], 2 * 2.33 / 11.3, places=12)
        self.assertEqual(fw["mat_tren"] + fw["mat_day"], 0.0)
        full = bc.formwork_per_m(11.3, top=True, bottom=True)
        self.assertEqual(full["mat_tren"], 2.5)
        self.assertAlmostEqual(full["tong"] - fw["tong"], 5.0, places=12)   # thêm nắp + đáy = 2 × 2.5

    def test_double_cell_has_two_inner_perimeters(self):
        bc = BoxCulvert(6.9, 3.6, 3, 3, 2, 0.25)
        self.assertAlmostEqual(bc.formwork_per_m(11.3)["trong"], 2 * (12 - 4 * 0.25 * (2 - math.sqrt(2))), places=9)


@unittest.skipUnless(os.path.exists(A5), "thiếu workbook ví dụ A5")
class A5SheetDerivationTest(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.ev = WorkbookEvaluator(A5)
        cls.ws = openpyxl.load_workbook(A5)["QS_TAKEOFF"]

    def test_concrete_quantities_unchanged_and_match_hand_values(self):
        # số đốt × dài × diện tích mặt cắt: 34·11.3·2.33, 50·11.3·4.085, 108·11.3·7.09
        for row, expected in ((8, 895.186), (9, 2308.025), (10, 8652.636)):
            self.assertAlmostEqual(self.ev.value("QS_TAKEOFF", row, 9), expected, places=3)

    def test_no_hardcoded_dimensions_in_concrete_formulas(self):
        for row in CULVERT_ROWS:
            f = self.ws[f"I{row}"].value
            self.assertNotIn("2.0*2.0", f)
            self.assertNotIn("3.0*3.0", f)
            self.assertNotIn("0.25", f)
            self.assertNotIn("0.2*", f)
            for col in "KLMN":
                self.assertIsNotNone(self.ws[f"{col}{row}"].value, f"thiếu đầu vào {col}{row}")

    def test_sheet_formwork_formula_matches_python_module(self):
        for crow, spec in CULVERT_ROWS.items():
            bc, _ = CASES[crow]
            expected = bc.formwork_per_m(11.3)["tong"]
            self.assertAlmostEqual(self.ev.value("QS_TAKEOFF", spec["formwork_row"], 19), expected, places=9)

    def test_legacy_formwork_numbers_are_kept_not_silently_overwritten(self):
        # Số ván khuôn đang dùng cho hợp đồng được giữ nguyên; chênh lệch hiển thị ở cột T để QS quyết định.
        self.assertEqual([self.ws[f"F{r}"].value for r in (11, 12, 13)], [11.14, 16.48, 25.66])


if __name__ == "__main__":
    unittest.main()

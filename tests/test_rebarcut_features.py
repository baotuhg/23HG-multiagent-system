# -*- coding: utf-8 -*-
"""
Kiểm thử các tính năng học từ RebarCut Pro Excel: giới hạn tổ cắt (PA3), cắt đầu cây,
phân loại đầu thừa 20D/100D, phương án nối thép (PA4 an toàn hơn), xuất bố cục RebarCut.
Chạy: python -m unittest discover tests
"""

import contextlib
import io
import os
import sys
import tempfile
import unittest
from collections import Counter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from core.agents.rebar_agent import RebarAgent
from core.state.shared_state import ProjectPhase
from core.supervisor.supervisor_agent import AECSupervisor
from tests.test_cutting_stock_solver import random_demands
from tools.bbs_loader import load_bbs
from tools.cutting_stock_solver import CuttingStockSolver, CutDemand
from tools.rebarcut_export import write_rebarcut_workbook

# BBS ví dụ trong RebarCut_Pro_Excel_V5.xlsm (RebarCut: PA1 388 cây, PA4 300 cây / 193 mối nối)
REBARCUT_SAMPLE = [("F1", 7262, 57), ("F1A", 8262, 56), ("F1B", 6025, 56), ("F2", 7262, 57),
                   ("F2A", 8262, 56), ("W1", 8467, 54), ("W1A", 9467, 52)]


def rebarcut_demands(zones=((0.0, 1.0),), allowed=True):
    return [CutDemand(length, qty, 25, mark, splice_allowed=allowed, lap_xd=40, max_splice_ratio=0.5,
                      splice_zones=list(zones) if zones else None)
            for mark, length, qty in REBARCUT_SAMPLE]


class SplicingTest(unittest.TestCase):

    def assert_valid_spliced_plan(self, sol, demands, zones, usable=11_700, kerf=0):
        for bar in sol.assignment:
            self.assertLessEqual(sum(bar["cuts_mm"]) + kerf * (len(bar["cuts_mm"]) - 1), usable)
        whole = Counter()
        for bar in sol.assignment:
            for length, mark in zip(bar["cuts_mm"], bar["marks"]):
                if "(nối-" not in mark:
                    whole[mark] += 1
        splices = Counter(s["mark"] for s in sol.splices)
        segments = Counter(m for bar in sol.assignment for m in bar["marks"] if "(nối-" in m)
        for d in demands:
            self.assertEqual(whole[d.mark] + splices[d.mark], d.quantity, d.mark)
            self.assertLessEqual(splices[d.mark], int(d.quantity * 0.5), d.mark)
            self.assertEqual(segments[f"{d.mark}(nối-A)"], splices[d.mark])
            self.assertEqual(segments[f"{d.mark}(nối-B)"], splices[d.mark])
        for s in sol.splices:
            self.assertEqual(s["a_mm"] + s["b_mm"] - s["lap_mm"], s["bar_length_mm"])
            self.assertEqual(s["lap_mm"], 1000)                        # 40 × Ø25
            self.assertGreaterEqual(min(s["a_mm"], s["b_mm"]), 1000)   # ≥ max(L nối, 20D)
            self.assertTrue(any(z0 * s["bar_length_mm"] - 1e-6 <= s["lap_from_mm"]
                                and s["lap_to_mm"] <= z1 * s["bar_length_mm"] + 1e-6 for z0, z1 in zones))
            self.assertIn(f"{s['mark']}(nối-A)", sol.assignment[s["bar_a"] - 1]["marks"])
            self.assertIn(f"{s['mark']}(nối-B)", sol.assignment[s["bar_b"] - 1]["marks"])

    def test_beats_rebarcut_pa4_on_its_own_sample(self):
        demands = rebarcut_demands()
        sol = CuttingStockSolver(kerf_mm=0).solve(demands, allow_splicing=True)
        self.assertEqual(CuttingStockSolver(kerf_mm=0).solve(demands).total_bars_needed, 388)
        self.assertLess(sol.total_bars_needed, 300)           # RebarCut PA4: 300 cây
        self.assertLessEqual(sol.total_splices, 193)          # RebarCut PA4: 193 mối nối
        self.assert_valid_spliced_plan(sol, demands, [(0.0, 1.0)])

    def test_restricted_zones_are_respected(self):
        zones = [(0.0, 0.3), (0.7, 1.0)]
        demands = rebarcut_demands(zones)
        sol = CuttingStockSolver(kerf_mm=0).solve(demands, allow_splicing=True)
        self.assertGreater(sol.total_splices, 0)
        self.assert_valid_spliced_plan(sol, demands, zones)

    def test_no_zone_means_no_splice(self):
        sol = CuttingStockSolver(kerf_mm=0).solve(rebarcut_demands(zones=None), allow_splicing=True)
        self.assertEqual(sol.total_splices, 0)
        self.assertEqual(sol.total_bars_needed, 388)
        self.assertTrue(any("chưa có vùng cho phép nối" in w for w in sol.warnings))

    def test_not_allowed_means_no_splice(self):
        sol = CuttingStockSolver(kerf_mm=0).solve(rebarcut_demands(allowed=False), allow_splicing=True)
        self.assertEqual(sol.total_splices, 0)


class WorkshopConstraintTest(unittest.TestCase):

    def test_max_pieces_and_marks_per_bar(self):
        demands = random_demands(3, diameters=(16,))
        sol = CuttingStockSolver(max_pieces_per_bar=3, max_marks_per_bar=2, time_limit_s=5).solve(demands)
        self.assertIn(sol.status, ("OPTIMAL", "FEASIBLE"))
        for bar in sol.assignment:
            self.assertLessEqual(len(bar["cuts_mm"]), 3)
            self.assertLessEqual(len(set(bar["marks"])), 2)
        cut = Counter(m for bar in sol.assignment for m in bar["marks"])
        self.assertEqual(cut, Counter({d.mark: d.quantity for d in demands}))

    def test_end_trim_reduces_usable_length(self):
        # 2 × 5800 = 11600 vừa cây 11.7m, nhưng không vừa khi cắt 2 × 100mm đầu cây
        self.assertEqual(CuttingStockSolver(kerf_mm=0).solve([CutDemand(5800, 2, 16)]).total_bars_needed, 1)
        sol = CuttingStockSolver(kerf_mm=0, end_trim_mm=100).solve([CutDemand(5800, 2, 16)])
        self.assertEqual(sol.total_bars_needed, 2)
        self.assertEqual(CuttingStockSolver(end_trim_mm=100).solve([CutDemand(11_600, 1, 16)]).status, "INFEASIBLE")

    def test_offcut_classification(self):
        s = CuttingStockSolver()
        self.assertEqual(s.classify_offcut(0, 25), "Hết cây")
        self.assertEqual(s.classify_offcut(2500, 25), "Tái sử dụng")    # ≥ 100D
        self.assertEqual(s.classify_offcut(600, 25), "Đầu thừa ngắn")   # ≥ 20D
        self.assertEqual(s.classify_offcut(400, 25), "Phế")


class RebarCutIOTest(unittest.TestCase):

    def test_loader_reads_splice_columns(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "bbs.csv")
            with open(path, "w", encoding="utf-8") as f:
                f.write("Bar Mark;Ø (mm);Số lượng;Chiều dài (m);Cho nối?;Kiểu nối;L nối xD;Max nối %;Vùng cho phép nối\n"
                        "A;25;10;9,0;Có;Nối chồng;40;50;0-25%; 75-100%\n".replace("0-25%; 75-100%", "\"0-25%; 75-100%\"")
                        + "B;25;10;9,0;Không;;;;\n"
                        + "C;25;10;9,0;Có;Coupler;;;0-1\n"
                        + "D;25;10;9,0;Có;Nối chồng;40;;0.5-0.2\n"
                        + "E;25;10;9,0;Có;Nối chồng;;;Theo thiết kế\n")
            r = load_bbs(path)
        by = {d.mark: d for d in r.demands}
        self.assertEqual((by["A"].splice_allowed, by["A"].lap_xd, by["A"].max_splice_ratio, by["A"].splice_zones),
                         (True, 40, 0.5, [(0.0, 0.25), (0.75, 1.0)]))
        self.assertFalse(by["B"].splice_allowed)
        self.assertTrue(by["E"].splice_allowed)
        self.assertIsNone(by["E"].splice_zones)          # chữ, không phải vùng số → không nối
        self.assertEqual(len(r.errors), 2)               # Coupler thiếu L nối; vùng 0.5-0.2 sai

    def test_export_roundtrip(self):
        demands = rebarcut_demands()
        solver = CuttingStockSolver(kerf_mm=0)
        base, spliced = solver.solve(demands), solver.solve(demands, allow_splicing=True)
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "rc.xlsx")
            write_rebarcut_workbook(path, demands, base, spliced)
            back = load_bbs(path)
            import openpyxl
            wb = openpyxl.load_workbook(path)
            self.assertEqual(wb.sheetnames, ["INPUT", "SO_SANH", "PA_TOI_UU", "PA_NOI", "MOI_NOI", "REMAIN", "CHI_TIET"])
            self.assertEqual(wb["SO_SANH"]["B7"].value, spliced.total_bars_needed)
            self.assertEqual(wb["MOI_NOI"].max_row - 4, spliced.total_splices)
        self.assertEqual([(d.mark, d.length_mm, d.quantity, d.splice_zones) for d in back.demands],
                         [(d.mark, d.length_mm, d.quantity, d.splice_zones) for d in demands])

    def test_agent_splice_plan_in_real_mode(self):
        with tempfile.TemporaryDirectory() as tmp:
            bbs = os.path.join(tmp, "bbs.csv")
            with open(bbs, "w", encoding="utf-8") as f:
                f.write("mark,diameter_mm,length_mm,quantity,Cho nối?,Vùng cho phép nối\n")
                for mark, length, qty in REBARCUT_SAMPLE:
                    f.write(f"{mark},25,{length},{qty},Có,0-1\n")
            out = os.path.join(tmp, "rc.xlsx")
            sup = AECSupervisor(project_root=ROOT, human_gate_mode="auto", persist_path=os.path.join(tmp, "s.json"))
            with contextlib.redirect_stdout(io.StringIO()):
                sup.register_agent(RebarAgent(bbs_path=bbs, kerf_mm=0, splice=True, rebarcut_out=out))
                ok = sup.run(phases=[ProjectPhase.REBAR_CUT])
            self.assertTrue(ok)
            plan = sup.bus.get_cutting_dict()["splice_plan"]
            self.assertEqual(sup.bus.get_cutting_dict()["total_bars_needed"], 388)   # phiếu chính: không nối
            self.assertGreater(plan["bars_saved"], 88)
            self.assertTrue(os.path.exists(out))


if __name__ == "__main__":
    unittest.main()

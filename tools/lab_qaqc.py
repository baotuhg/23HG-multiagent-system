# -*- coding: utf-8 -*-
"""
LAB QA/QC — Đánh giá phiếu thí nghiệm vật liệu & giải tỏa điểm dừng kỹ thuật (Hold Point)

Pure Python, Zero LLM. Mỗi dòng của bảng phiếu thí nghiệm là một phiếu / tổ mẫu:
  - Nén bê tông : cường độ từng viên (Mẫu 1..n) hoặc kết quả trung bình, ngày đúc, ngày nén,
                  yêu cầu (số MPa, "C30/37", "C30", "M300"), loại mẫu (lập phương / trụ, kích thước)
  - Kéo thép    : mác thép (CB240-T, CB300-V, CB400-V, CB500-V), giới hạn chảy, giới hạn bền,
                  độ giãn dài
  - Chỉ tiêu khác (PDA, siêu âm cọc, độ sụt...): Kết quả so với Yêu cầu dạng "≥ 7800",
                  "≤ 3", "18±2", "16-20" hoặc chữ ("Loại 1", "Đạt")

TIÊU CHÍ MẶC ĐỊNH (chỉnh được qua LabCriteria, phải đối chiếu chỉ dẫn kỹ thuật dự án):
  Bê tông tuổi ≥ 28 ngày: cường độ tổ mẫu ≥ 100% yêu cầu VÀ không viên nào < 85% yêu cầu
    (tham khảo nguyên tắc nghiệm thu của TCVN 4453:1995). Cường độ tổ mẫu 3 viên: trung bình;
    nếu viên lớn nhất / nhỏ nhất lệch > 15% so với viên giữa thì lấy viên giữa
    (tham khảo TCVN 3118). Quy đổi về mẫu lập phương 150mm: 100mm × 0.91, 200mm × 1.05.
  Bê tông tuổi < 28 ngày: CHỜ R28; cảnh báo sớm nếu R7 < 65% yêu cầu.
  Thép: giới hạn chảy, giới hạn bền, độ giãn dài ≥ giá trị tối thiểu theo mác
    (bảng STEEL_GRADES tham khảo TCVN 1651-1/-2:2018 — kiểm tra lại với chứng chỉ / hồ sơ).

Kết quả gom theo BBNT liên kết: có phiếu KHÔNG ĐẠT → BBNT bị CHẶN; còn phiếu chờ → CHỜ;
tất cả đạt → GIẢI TỎA.
"""

from __future__ import annotations
import csv
import json
import os
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Dict, List, Optional, Sequence, Tuple

from tools.bbs_loader import _FORMULA, _norm, _to_number
from tools.excel_eval import serial_to_date

HEADER_SCAN_ROWS = 40

# Mác thép: (giới hạn chảy min MPa, giới hạn bền min MPa, độ giãn dài min %) — tham khảo, cần đối chiếu
STEEL_GRADES: Dict[str, Tuple[float, float, float]] = {
    "CB240T": (240, 380, 20),
    "CB300T": (300, 440, 16),
    "CB300V": (300, 450, 19),
    "CB400V": (400, 570, 14),
    "CB500V": (500, 650, 14),
}

# Cấp bê tông EN 206: C(cường độ mẫu trụ)/(cường độ mẫu lập phương), MPa
CONCRETE_CLASSES: Dict[int, int] = {12: 15, 16: 20, 20: 25, 25: 30, 30: 37, 35: 45, 40: 50,
                                    45: 55, 50: 60, 55: 67, 60: 75}

CUBE_SIZE_FACTOR = {100: 0.91, 150: 1.0, 200: 1.05}   # quy đổi về lập phương 150mm


class LabLoadError(ValueError):
    """Không đọc được bảng phiếu thí nghiệm."""


@dataclass
class LabCriteria:
    group_min_ratio: float = 1.0        # cường độ tổ mẫu / yêu cầu
    specimen_min_ratio: float = 0.85    # viên nhỏ nhất / yêu cầu
    acceptance_age_days: int = 28
    early_warning_ratio: float = 0.65   # R7 / yêu cầu dưới mức này → cảnh báo sớm
    outlier_pct: float = 15.0           # lệch so với viên giữa (tổ 3 viên)


@dataclass
class LabRecord:
    row: int
    test_id: str
    kind: str                   # CONCRETE / STEEL / OTHER
    component: str
    bbnt: str
    status: str = "PASS"        # PASS / FAIL / PENDING
    value: Optional[float] = None
    required: Optional[float] = None
    ratio: Optional[float] = None
    age_days: Optional[int] = None
    detail: str = ""


@dataclass
class HoldPoint:
    bbnt: str
    status: str                 # GIẢI TỎA / CHỜ / CHẶN
    tests: List[str] = field(default_factory=list)
    reasons: List[str] = field(default_factory=list)


@dataclass
class LabEvaluation:
    source: str
    records: List[LabRecord] = field(default_factory=list)
    hold_points: List[HoldPoint] = field(default_factory=list)
    errors: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)

    def count(self, status: str) -> int:
        return sum(1 for r in self.records if r.status == status)


# ─────────────────────────────────────────────────────────────────────────────
# ĐỌC FILE
# ─────────────────────────────────────────────────────────────────────────────

def _detect_columns(headers: Sequence[Any]) -> Dict[str, Any]:
    cols: Dict[str, Any] = {"specimens": []}
    for idx, raw in enumerate(headers):
        if raw is None or raw is _FORMULA:
            continue
        h = _norm(raw)
        if not h:
            continue
        if re.fullmatch(r"(mau|vien|r|mauthu|vienmau)\d{1,2}(mpa)?", h):
            cols["specimens"].append(idx)
        elif "maphieu" in h or "sophieu" in h or h in ("phieu", "testid", "id"):
            cols.setdefault("test_id", idx)
        elif "loaimau" in h or "hinhdangmau" in h or "kichthuocmau" in h or h in ("specimen", "specimentype"):
            cols.setdefault("specimen_type", idx)
        elif "loaithinghiem" in h or h in ("loai", "testtype", "type"):
            cols.setdefault("kind", idx)
        elif "bbnt" in h or "bienban" in h or h in ("holdpoint",):
            cols.setdefault("bbnt", idx)
        elif "caukien" in h or "hangmuc" in h or "vitri" in h or h in ("component",):
            cols.setdefault("component", idx)
        elif "ngayduc" in h or h in ("castingdate", "cast"):
            cols.setdefault("cast_date", idx)
        elif "ngaythinghiem" in h or "ngaynen" in h or "ngaythi" in h or h in ("testdate",):
            cols.setdefault("test_date", idx)
        elif "tuoi" in h or h in ("age", "agedays"):
            cols.setdefault("age", idx)
        elif "macthep" in h or h in ("steelgrade",):
            cols.setdefault("steel_grade", idx)
        elif "gioihanchay" in h or h in ("fy", "reh", "yield"):
            cols.setdefault("fy", idx)
        elif "gioihanben" in h or h in ("fu", "rm", "tensile"):
            cols.setdefault("fu", idx)
        elif "giandai" in h or h in ("elongation", "a5", "a"):
            cols.setdefault("elong", idx)
        elif "yeucau" in h or h in ("required", "requirement", "macyc", "capyc"):
            cols.setdefault("required", idx)
        elif "chitieu" in h or h in ("indicator",):
            cols.setdefault("indicator", idx)
        elif "ketqua" in h or h in ("result", "value", "giatri"):
            cols.setdefault("result", idx)
    return cols


def _find_header(rows):
    for r, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        cols = _detect_columns(row)
        if "test_id" in cols and ("result" in cols or cols["specimens"] or "fy" in cols):
            return r, cols
    return None


def load_lab_results(path: str, sheet: Optional[str] = None) -> List[Tuple[int, Dict[str, Any]]]:
    """Trả về [(số dòng, {trường: giá trị})]."""
    if not os.path.exists(path):
        raise LabLoadError(f"Không tìm thấy file phiếu thí nghiệm: {path}")
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        from tools.excel_eval import WorkbookEvaluator
        ev = WorkbookEvaluator(path)
        if sheet and sheet not in ev.sheetnames:
            raise LabLoadError(f"Không có sheet '{sheet}' trong {path}")
        for name in ([sheet] if sheet else ev.sheetnames):
            rows = ev.rows(name, missing=_FORMULA)
            if _find_header(rows):
                return _to_records(rows, f"{path}#{name}")
        raise LabLoadError(f"Không tìm thấy bảng phiếu thí nghiệm trong {path} — cần cột Mã phiếu và "
                           f"Kết quả / Mẫu 1..n / Giới hạn chảy")
    if ext in (".csv", ".txt"):
        with open(path, "r", encoding="utf-8-sig", newline="") as f:
            text = f.read().replace("\r\n", "\n")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return _to_records([r for r in csv.reader(text.splitlines(), dialect)], path)
    if ext == ".json":
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        items = data.get("items", []) if isinstance(data, dict) else data
        keys: List[str] = []
        for i in items:
            keys.extend(k for k in i if k not in keys)
        return _to_records([keys] + [[i.get(k) for k in keys] for i in items], path)
    raise LabLoadError(f"Định dạng chưa hỗ trợ: {ext} (dùng .xlsx, .csv hoặc .json)")


def _to_records(rows, source: str):
    found = _find_header(rows)
    if found is None:
        raise LabLoadError(f"{source}: cần cột Mã phiếu và Kết quả / Mẫu 1..n / Giới hạn chảy")
    header_idx, cols = found
    out = []
    for offset, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        def cell(i):
            if i is None or i >= len(row):
                return None
            v = row[i]
            return None if isinstance(v, str) and not v.strip() else v
        rec = {k: cell(v) for k, v in cols.items() if k != "specimens"}
        rec["specimens"] = [cell(i) for i in cols["specimens"]]
        if all(v is None for k, v in rec.items() if k != "specimens") and not any(
                v is not None for v in rec["specimens"]):
            continue
        rec["_source"] = source
        out.append((offset, rec))
    return out


# ─────────────────────────────────────────────────────────────────────────────
# ĐÁNH GIÁ
# ─────────────────────────────────────────────────────────────────────────────

def _as_date(v: Any) -> Optional[date]:
    if v is None or v is _FORMULA:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    d = serial_to_date(v)                    # ngày do công thức Excel tính ra là số seri
    if d is not None:
        return d
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _kind(rec: Dict[str, Any]) -> str:
    text = _norm(f"{rec.get('kind') or ''} {rec.get('indicator') or ''}")
    if rec.get("fy") is not None or "thep" in text or "keo" in text:
        return "STEEL"
    if rec.get("specimens") and any(v is not None for v in rec["specimens"]):
        return "CONCRETE"
    if "betong" in text and ("nen" in text or "cuongdo" in text):
        return "CONCRETE"
    return "OTHER"


def _specimen_factor(text: Any) -> Tuple[float, bool, str]:
    """(hệ số quy đổi về lập phương 150, là mẫu trụ?, mô tả)."""
    t = _norm(text or "")
    size = re.search(r"(100|150|200)", t)
    if "tru" in t or "cylinder" in t:
        return 1.0, True, "trụ"
    s = int(size.group(1)) if size else 150
    return CUBE_SIZE_FACTOR[s], False, f"lập phương {s}mm"


def concrete_requirement(required: Any, cylinder: bool) -> float:
    """Cường độ yêu cầu (MPa) theo cùng loại mẫu: số MPa, C30/37, C30, M300."""
    n = _to_number(required)
    if n is not None:
        return n
    t = str(required or "").strip().upper().replace(" ", "")
    m = re.fullmatch(r"C(\d+)/(\d+)", t)
    if m:
        return float(m.group(1) if cylinder else m.group(2))
    m = re.fullmatch(r"C(\d+)", t)
    if m and int(m.group(1)) in CONCRETE_CLASSES:
        c = int(m.group(1))
        return float(c if cylinder else CONCRETE_CLASSES[c])
    m = re.fullmatch(r"M(\d+)", t)
    if m and not cylinder:
        return round(int(m.group(1)) * 0.0980665, 2)   # kG/cm² → MPa
    if re.fullmatch(r"B\d+(\.\d+)?", t):
        raise ValueError(f"cấp độ bền '{required}' là cường độ đặc trưng — ghi cường độ yêu cầu bằng số MPa")
    raise ValueError(f"không hiểu yêu cầu cường độ '{required}' (dùng số MPa, C30/37, C30 hoặc M300)")


def group_strength(values: List[float], outlier_pct: float) -> Tuple[float, str]:
    if len(values) == 3:
        lo, mid, hi = sorted(values)
        if mid > 0 and (abs(hi - mid) / mid * 100 > outlier_pct or abs(mid - lo) / mid * 100 > outlier_pct):
            return mid, f"viên lệch > {outlier_pct:g}% so với viên giữa → lấy viên giữa"
    return sum(values) / len(values), ""


def _parse_requirement(text: Any) -> Tuple[str, Any]:
    """'≥ 7800' → ('>=', 7800); '18±2' → ('range', (16, 20)); 'Loại 1' → ('eq', 'loai1')."""
    s = str(text or "").strip().replace(" ", "").replace(",", ".")
    if not s:
        raise ValueError("thiếu yêu cầu")
    m = re.fullmatch(r"(≥|>=|≤|<=|>|<)(-?\d+(?:\.\d+)?)", s)
    if m:
        op = {"≥": ">=", "≤": "<="}.get(m.group(1), m.group(1))
        return op, float(m.group(2))
    m = re.fullmatch(r"(-?\d+(?:\.\d+)?)±(\d+(?:\.\d+)?)", s)
    if m:
        c, d = float(m.group(1)), float(m.group(2))
        return "range", (c - d, c + d)
    m = re.fullmatch(r"(-?\d+(?:\.\d+)?)(?:-|–|÷)(-?\d+(?:\.\d+)?)", s)
    if m:
        return "range", (float(m.group(1)), float(m.group(2)))
    if re.fullmatch(r"-?\d+(?:\.\d+)?", s):
        return ">=", float(s)
    return "eq", _norm(text)


def evaluate(rows: List[Tuple[int, Dict[str, Any]]], source: str = "",
             criteria: Optional[LabCriteria] = None) -> LabEvaluation:
    c = criteria or LabCriteria()
    ev = LabEvaluation(source=source)
    seen_ids: Dict[str, int] = {}

    for row, rec in rows:
        test_id = str(rec.get("test_id") or "").strip()
        where = f"Dòng {row} ({test_id or '?'})"
        if not test_id:
            ev.errors.append(f"{where}: thiếu mã phiếu thí nghiệm")
            continue
        if test_id in seen_ids:
            ev.errors.append(f"{where}: trùng mã phiếu với dòng {seen_ids[test_id]}")
            continue
        seen_ids[test_id] = row
        if any(v is _FORMULA for v in list(rec.values()) + rec["specimens"]):
            ev.errors.append(f"{where}: có công thức không tính được")
            continue
        kind = _kind(rec)
        r = LabRecord(row=row, test_id=test_id, kind=kind,
                      component=str(rec.get("component") or "").strip(),
                      bbnt=str(rec.get("bbnt") or "").strip())
        try:
            if kind == "CONCRETE":
                _evaluate_concrete(rec, r, c)
            elif kind == "STEEL":
                _evaluate_steel(rec, r)
            else:
                _evaluate_other(rec, r)
        except ValueError as e:
            ev.errors.append(f"{where}: {e}")
            continue
        if r.status == "PENDING" and r.ratio is not None and r.ratio < c.early_warning_ratio:
            ev.warnings.append(f"{test_id} ({r.component}): R{r.age_days} = {r.value:.1f} MPa chỉ đạt "
                               f"{r.ratio:.0%} yêu cầu — cảnh báo sớm, theo dõi R28")
        ev.records.append(r)

    groups: Dict[str, List[LabRecord]] = {}
    for r in ev.records:
        if r.bbnt:
            groups.setdefault(r.bbnt, []).append(r)
        if r.status == "FAIL":
            ev.warnings.append(f"{r.test_id} ({r.component}) KHÔNG ĐẠT: {r.detail}")
    for bbnt, recs in groups.items():
        if any(x.status == "FAIL" for x in recs):
            status = "CHẶN"
        elif any(x.status == "PENDING" for x in recs):
            status = "CHỜ"
        else:
            status = "GIẢI TỎA"
        ev.hold_points.append(HoldPoint(
            bbnt=bbnt, status=status, tests=[x.test_id for x in recs],
            reasons=[f"{x.test_id}: {x.detail}" for x in recs if x.status != "PASS"],
        ))
    unlinked = [r.test_id for r in ev.records if not r.bbnt]
    if unlinked:
        ev.warnings.append(f"{len(unlinked)} phiếu chưa liên kết BBNT: {unlinked[:5]}")
    return ev


def _evaluate_concrete(rec, r: LabRecord, c: LabCriteria) -> None:
    factor, cylinder, shape = _specimen_factor(rec.get("specimen_type"))
    required = concrete_requirement(rec.get("required"), cylinder)
    raw = [v for v in rec["specimens"] if v is not None]
    if not raw and rec.get("result") is not None:
        raw = [rec["result"]]
    values = [_to_number(v) for v in raw]
    if not values or any(v is None or v <= 0 for v in values):
        raise ValueError(f"cường độ các viên mẫu không hợp lệ: {raw}")
    values = [v * factor for v in values]

    age = _to_number(rec.get("age"))
    cast, tested = _as_date(rec.get("cast_date")), _as_date(rec.get("test_date"))
    if cast and tested:
        if tested < cast:
            raise ValueError(f"ngày thí nghiệm {tested} trước ngày đúc {cast}")
        from_dates = (tested - cast).days
        if age is not None and int(age) != from_dates:
            raise ValueError(f"tuổi mẫu ghi {int(age)} ngày ≠ tính theo ngày đúc/ngày nén {from_dates} ngày")
        age = from_dates
    if age is None:
        raise ValueError("thiếu tuổi mẫu (ghi Tuổi hoặc Ngày đúc + Ngày thí nghiệm)")
    r.age_days = int(age)

    strength, note = group_strength(values, c.outlier_pct)
    r.value, r.required = round(strength, 2), required
    r.ratio = strength / required
    lowest = min(values)
    base = f"{shape}, R{r.age_days} = {strength:.1f} MPa / yêu cầu {required:g} MPa ({r.ratio:.0%})"
    if note:
        base += f"; {note}"
    if factor != 1.0:
        base += f"; đã quy đổi × {factor:g}"
    if r.age_days < c.acceptance_age_days:
        r.status = "PENDING"
        r.detail = f"{base} — chờ R{c.acceptance_age_days}"
        return
    problems = []
    if strength < required * c.group_min_ratio - 1e-9:
        problems.append(f"tổ mẫu < {c.group_min_ratio:.0%} yêu cầu")
    if lowest < required * c.specimen_min_ratio - 1e-9:
        problems.append(f"viên thấp nhất {lowest:.1f} MPa < {c.specimen_min_ratio:.0%} yêu cầu")
    r.status = "FAIL" if problems else "PASS"
    r.detail = base + ("; " + "; ".join(problems) if problems else "")


def _evaluate_steel(rec, r: LabRecord) -> None:
    grade_raw = rec.get("steel_grade") or rec.get("required")
    grade = _norm(grade_raw or "").upper()
    if grade not in STEEL_GRADES:
        raise ValueError(f"mác thép '{grade_raw}' chưa có trong bảng {sorted(STEEL_GRADES)}")
    fy_min, fu_min, el_min = STEEL_GRADES[grade]
    fy, fu, el = (_to_number(rec.get(k)) for k in ("fy", "fu", "elong"))
    if fy is None or fu is None or el is None:
        raise ValueError("thiếu giới hạn chảy / giới hạn bền / độ giãn dài")
    problems = []
    if fy < fy_min:
        problems.append(f"giới hạn chảy {fy:g} < {fy_min:g} MPa")
    if fu < fu_min:
        problems.append(f"giới hạn bền {fu:g} < {fu_min:g} MPa")
    if el < el_min:
        problems.append(f"độ giãn dài {el:g}% < {el_min:g}%")
    r.value, r.required, r.ratio = fy, fy_min, fy / fy_min
    r.status = "FAIL" if problems else "PASS"
    r.detail = (f"{grade_raw}: fy {fy:g} / fu {fu:g} MPa, giãn dài {el:g}%"
                + ("; " + "; ".join(problems) if problems else ""))


def _evaluate_other(rec, r: LabRecord) -> None:
    indicator = str(rec.get("indicator") or rec.get("kind") or "").strip()
    result = rec.get("result")
    if result is None:
        raise ValueError("thiếu kết quả")
    op, target = _parse_requirement(rec.get("required"))
    if op == "eq":
        ok = _norm(result) == target
    else:
        v = _to_number(result)
        if v is None:
            raise ValueError(f"kết quả '{result}' không phải số trong khi yêu cầu '{rec.get('required')}' là số")
        r.value = v
        if op == "range":
            ok = target[0] <= v <= target[1]
        else:
            r.required = target
            ok = {">=": v >= target, ">": v > target, "<=": v <= target, "<": v < target}[op]
    r.status = "PASS" if ok else "FAIL"
    r.detail = f"{indicator}: {result} / yêu cầu {rec.get('required')}"


# ─────────────────────────────────────────────────────────────────────────────
# XUẤT BÁO CÁO
# ─────────────────────────────────────────────────────────────────────────────

def write_lab_report(path: str, ev: LabEvaluation, criteria: Optional[LabCriteria] = None) -> None:
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    c = criteria or LabCriteria()
    fills = {"PASS": "C6EFCE", "FAIL": "FFC7CE", "PENDING": "FFEB9C",
             "GIẢI TỎA": "C6EFCE", "CHẶN": "FFC7CE", "CHỜ": "FFEB9C"}
    label = {"PASS": "ĐẠT", "FAIL": "KHÔNG ĐẠT", "PENDING": "CHỜ"}

    def header(ws, row, values):
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3A5E")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HOLD_POINTS"
    ws["A1"] = "TỔNG HỢP ĐIỂM DỪNG KỸ THUẬT THEO BIÊN BẢN NGHIỆM THU"
    ws["A2"] = f"Nguồn: {ev.source}"
    header(ws, 4, ["BBNT", "Trạng thái", "Phiếu thí nghiệm", "Lý do"])
    for hp in ev.hold_points:
        ws.append([hp.bbnt, hp.status, ", ".join(hp.tests), "; ".join(hp.reasons)])
        ws.cell(row=ws.max_row, column=2).fill = PatternFill("solid", fgColor=fills[hp.status])
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 90

    ws = wb.create_sheet("KET_QUA")
    header(ws, 1, ["Dòng file", "Mã phiếu", "Loại", "Cấu kiện", "BBNT", "Tuổi (ngày)", "Giá trị", "Yêu cầu",
                   "Tỷ lệ", "Kết luận", "Chi tiết"])
    for r in ev.records:
        ws.append([r.row, r.test_id, r.kind, r.component, r.bbnt, r.age_days, r.value, r.required,
                   None if r.ratio is None else round(r.ratio, 3), label[r.status], r.detail])
        ws.cell(row=ws.max_row, column=10).fill = PatternFill("solid", fgColor=fills[r.status])
    ws.column_dimensions["K"].width = 90

    ws = wb.create_sheet("TIEU_CHI")
    ws.append(["Tiêu chí đã dùng (đối chiếu chỉ dẫn kỹ thuật dự án)"])
    ws.append(["Bê tông: tổ mẫu ≥ yêu cầu ×", c.group_min_ratio])
    ws.append(["Bê tông: viên thấp nhất ≥ yêu cầu ×", c.specimen_min_ratio])
    ws.append(["Tuổi nghiệm thu (ngày)", c.acceptance_age_days])
    ws.append(["Cảnh báo sớm khi R < yêu cầu ×", c.early_warning_ratio])
    ws.append(["Tổ 3 viên: lấy viên giữa khi lệch > %", c.outlier_pct])
    for grade, (fy, fu, el) in STEEL_GRADES.items():
        ws.append([f"Thép {grade}: fy ≥ {fy:g} MPa, fu ≥ {fu:g} MPa, giãn dài ≥ {el:g}%"])
    ws.column_dimensions["A"].width = 60
    if ev.warnings:
        ws = wb.create_sheet("CANH_BAO")
        for w in ev.warnings:
            ws.append([w])
        ws.column_dimensions["A"].width = 120
    wb.save(path)

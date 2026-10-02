# -*- coding: utf-8 -*-
"""
TAKEOFF LOADER — đọc bảng cấu kiện (CSV / Excel / JSON) → đo bóc bằng tools/takeoff_rules → xuất Bảng 6.2 / 6.1.

Mỗi dòng là một cấu kiện / công tác. Cột nhận diện theo tiêu đề (không phân biệt hoa thường, có dấu hay không):

  Loại (bắt buộc)      : be_tong | van_khuon | coc_khoan_nhoi | khoan | cot_tron | dao_hao | dao_ho | mat_cat
                         | ong | dan_giao_trong | dan_giao_cot
  Tên / Danh mục công tác (bắt buộc), Ký hiệu bản vẽ, Mã hiệu, Ghi chú
  Số lượng (số bộ phận giống nhau, mặc định 1), Dài (m), Rộng (m), Cao (m)
  Kiểu (van_khuon: mong | cot | dam | san | tuong), Dày sàn (m), Có cạnh (1/0), Có đầu (1/0)
  Đường kính (m), Đập đầu cọc (m), Sâu khoan (m), Mái taluy (m = ngang/đứng)
  Lỗ rỗng (m3 với bê tông, m2 với ván khuôn; nhiều lỗ ngăn bằng ';'), Thép (m3, chỉ dùng để xét có trừ hay không)
  Lý trình (m), Diện tích (m2)  — loại mat_cat: mỗi dòng một mặt cắt, các dòng cùng Tên gộp thành một khối lượng
  Trừ hố ga (m), Thoát nước (1/0) — loại ong;  Diện tích bằng (m2), Chu vi (m) — dàn giáo

Dòng sai dữ liệu → dừng và liệt kê từng dòng (không đoán).
Pure Python, Zero LLM.
"""

from __future__ import annotations
import json
import os
from typing import Any, Dict, List, Optional, Sequence, Tuple

from tools.bbs_loader import _FORMULA, _norm, _read_csv, _to_number
from tools.takeoff_rules import (
    DEFAULT_PROFILE, MeasurementProfile, PROFILE_TT13_2021_PL_VI, Quantity,
    average_end_volume, bang_6_1, bang_6_2, bored_length, circular_column, load_profile, pile, pipe_length,
    pit_excavation, rect_concrete, rect_formwork, scaffold_column, scaffold_inner, trench_excavation,
)

HEADER_SCAN_ROWS = 30

# khóa chuẩn → các tiêu đề chấp nhận (đã chuẩn hóa bằng _norm)
_ALIASES = {
    "loai": ["loai", "loaicongtac", "type", "kind"],
    "ten": ["ten", "tencongtac", "danhmuccongtac", "noidungcongtac", "hangmuc", "name"],
    "ban_ve": ["kyhieubanve", "banve", "drawing"],
    "ma_hieu": ["mahieu", "mahieucongtac", "madinhmuc", "code"],
    "ghi_chu": ["ghichu", "note", "notes"],
    "so_luong": ["soluong", "sobophangiongnhau", "sobophan", "n", "count"],
    "dai": ["dai", "daim", "chieudai", "chieudaim", "length"],
    "rong": ["rong", "rongm", "chieurong", "chieurongm", "width"],
    "cao": ["cao", "caom", "chieucao", "chieucaom", "caoday", "caodaym", "height"],
    "kieu": ["kieu", "kieuvankhuon", "caukien", "formworkkind"],
    "day_san": ["daysan", "daysanm", "chieudaysan", "slabthickness"],
    "co_canh": ["cocanh", "edges"],
    "co_dau": ["codau", "ends"],
    "duong_kinh": ["duongkinh", "duongkinhm", "d", "diameter"],
    "dap_dau": ["dapdaucoc", "dapdaucocm", "dapdau", "cutoff"],
    "sau_khoan": ["saukhoan", "saukhoanm", "chieusaukhoan", "drilldepth"],
    "mai": ["maitaluy", "mai", "slope"],
    "lo_rong": ["lorong", "lorongm3", "lorongm2", "openings"],
    "thep_m3": ["thep", "thepm3", "thetichthep", "rebarvolume"],
    "ly_trinh": ["lytrinh", "lytrinhm", "station"],
    "dien_tich": ["dientich", "dientichm2", "area"],
    "tru_ho_ga": ["truhoga", "truhogam", "chambers"],
    "thoat_nuoc": ["thoatnuoc", "drainage"],
    "dien_tich_bang": ["dientichbang", "dientichbangm2", "planarea"],
    "chu_vi": ["chuvi", "chuvim", "perimeter"],
}
_LOOKUP = {alias: key for key, aliases in _ALIASES.items() for alias in aliases}

KINDS = ("be_tong", "van_khuon", "coc_khoan_nhoi", "khoan", "cot_tron", "dao_hao", "dao_ho", "mat_cat", "ong",
         "dan_giao_trong", "dan_giao_cot")
_REQUIRED = {
    "be_tong": ("dai", "rong", "cao"),
    "van_khuon": ("kieu", "dai", "rong", "cao"),
    "coc_khoan_nhoi": ("duong_kinh", "dai"),
    "khoan": ("sau_khoan",),
    "cot_tron": ("duong_kinh", "cao"),
    "dao_hao": ("dai", "rong", "cao"),
    "dao_ho": ("dai", "rong", "cao"),
    "mat_cat": ("ly_trinh", "dien_tich"),
    "ong": ("dai",),
    "dan_giao_trong": ("dien_tich_bang", "cao"),
    "dan_giao_cot": ("chu_vi", "cao"),
}
_NUMERIC = ("so_luong", "dai", "rong", "cao", "day_san", "duong_kinh", "dap_dau", "sau_khoan", "mai", "thep_m3",
            "ly_trinh", "dien_tich", "tru_ho_ga", "dien_tich_bang", "chu_vi")
PROFILES = {"mac-dinh": DEFAULT_PROFILE, "tt13-2021": PROFILE_TT13_2021_PL_VI}


class TakeoffLoadError(ValueError):
    """Không đọc được bảng cấu kiện hoặc có dòng sai dữ liệu."""


def resolve_profile(name_or_path: Optional[str]) -> MeasurementProfile:
    if not name_or_path:
        return DEFAULT_PROFILE
    if name_or_path in PROFILES:
        return PROFILES[name_or_path]
    if os.path.exists(name_or_path):
        return load_profile(name_or_path)
    raise TakeoffLoadError(f"Hồ sơ quy tắc '{name_or_path}' không có: dùng {', '.join(PROFILES)} hoặc đường dẫn JSON")


def _read_rows(path: str, sheet: Optional[str]) -> List[List[Any]]:
    if not os.path.exists(path):
        raise TakeoffLoadError(f"Không tìm thấy file: {path}")
    ext = os.path.splitext(path)[1].lower()
    if ext in (".csv", ".txt"):
        return _read_csv(path)
    if ext == ".json":
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        items = data.get("items", []) if isinstance(data, dict) else data
        if not isinstance(items, list) or not all(isinstance(i, dict) for i in items):
            raise TakeoffLoadError("File JSON phải là danh sách đối tượng hoặc {\"items\": [...]}")
        keys: List[str] = []
        for item in items:
            keys += [k for k in item if k not in keys]
        return [keys] + [[item.get(k) for k in keys] for item in items]
    if ext in (".xlsx", ".xlsm"):
        from tools.excel_eval import WorkbookEvaluator
        ev = WorkbookEvaluator(path)
        names = [sheet] if sheet else ev.sheetnames
        for name in names:
            if name not in ev.sheetnames:
                raise TakeoffLoadError(f"Không có sheet '{name}' trong {path}. Các sheet: {ev.sheetnames}")
            rows = ev.rows(name, missing=_FORMULA)
            if _find_header(rows):
                return rows
        raise TakeoffLoadError(f"Không tìm thấy bảng cấu kiện (cần cột 'Loại' và 'Tên') trong {path}")
    raise TakeoffLoadError(f"Định dạng chưa hỗ trợ: {ext} (dùng .csv, .xlsx hoặc .json)")


def _find_header(rows: Sequence[Sequence[Any]]) -> Optional[Tuple[int, Dict[str, int]]]:
    for r, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        cols: Dict[str, int] = {}
        for i, h in enumerate(row):
            key = _LOOKUP.get(_norm(h)) if h not in (None, "", _FORMULA) else None
            if key and key not in cols:
                cols[key] = i
        if "loai" in cols and "ten" in cols:
            return r, cols
    return None


def _openings(v: Any) -> Tuple[List[float], bool]:
    if v in (None, "", _FORMULA):
        return [], True
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return [float(v)], True
    parts = [p for p in str(v).replace("|", ";").split(";") if p.strip()]
    nums = [_to_number(p) for p in parts]
    return [n for n in nums if n is not None], all(n is not None and n >= 0 for n in nums)


def _flag(v: Any, default: bool) -> bool:
    if v in (None, "", _FORMULA):
        return default
    return _norm(v) in ("1", "co", "yes", "true", "x")


def load_takeoff(path: str, sheet: Optional[str] = None,
                 profile: MeasurementProfile = DEFAULT_PROFILE) -> List[Quantity]:
    """Đọc bảng cấu kiện và trả các khối lượng đo bóc. Có dòng sai → TakeoffLoadError liệt kê toàn bộ."""
    rows = _read_rows(path, sheet)
    found = _find_header(rows)
    if not found:
        raise TakeoffLoadError(f"{path}: không tìm thấy dòng tiêu đề (cần cột 'Loại' và 'Tên công tác')")
    hdr, cols = found
    errors: List[str] = []
    out: List[Quantity] = []
    sections: Dict[str, Dict[str, Any]] = {}       # mat_cat gộp theo tên

    for offset, row in enumerate(rows[hdr + 1:], start=hdr + 2):
        get = lambda k: row[cols[k]] if k in cols and cols[k] < len(row) else None
        if all(v in (None, "", _FORMULA) for v in row):
            continue
        kind = {k.replace("_", ""): k for k in KINDS}.get(_norm(get("loai") or ""))
        name = str(get("ten") or "").strip()
        where = f"Dòng {offset}" + (f" ({name[:40]})" if name else "")
        if not kind:
            errors.append(f"{where}: loại '{get('loai')}' không hợp lệ (dùng: {', '.join(KINDS)})")
            continue
        if not name:
            errors.append(f"{where}: thiếu tên công tác")
            continue
        vals: Dict[str, Optional[float]] = {}
        bad = False
        for k in _NUMERIC:
            raw = get(k)
            if raw in (None, ""):
                vals[k] = None
                continue
            if raw is _FORMULA:
                errors.append(f"{where}: cột {k} có công thức không tính được")
                bad = True
                continue
            num = _to_number(raw)
            if num is None or num < 0:
                errors.append(f"{where}: cột {k} = '{raw}' không phải số ≥ 0")
                bad = True
            vals[k] = num
        missing = [k for k in _REQUIRED[kind] if (k == "kieu" and not get("kieu")) or (k != "kieu" and vals.get(k) is None)]
        if missing:
            errors.append(f"{where}: loại {kind} thiếu cột {', '.join(missing)}")
            bad = True
        holes, holes_ok = _openings(get("lo_rong"))
        if not holes_ok:
            errors.append(f"{where}: cột lỗ rỗng '{get('lo_rong')}' sai (dùng số ≥ 0, ngăn bằng ';')")
            bad = True
        if bad:
            continue
        n = vals["so_luong"] if vals["so_luong"] is not None else 1.0
        refs = dict(drawing=str(get("ban_ve") or "").strip(), code=str(get("ma_hieu") or "").strip())
        try:
            if kind == "be_tong":
                qs = [rect_concrete(name, n, vals["dai"], vals["rong"], vals["cao"], openings_m3=holes,
                                    rebar_volume_m3=vals["thep_m3"] or 0.0, profile=profile)]
            elif kind == "van_khuon":
                fk = _norm(get("kieu"))
                qs = [rect_formwork(name, fk, n, vals["dai"], vals["rong"], vals["cao"],
                                    slab_thickness=vals["day_san"] or 0.0, edges=_flag(get("co_canh"), False),
                                    ends=_flag(get("co_dau"), False), openings_m2=holes, profile=profile)]
            elif kind == "coc_khoan_nhoi":
                qs = [q for q in pile(name, n, vals["duong_kinh"], vals["dai"], vals["dap_dau"] or 0.0, profile=profile) if q]
            elif kind == "khoan":
                qs = [bored_length(name, n, vals["sau_khoan"], profile=profile)]
            elif kind == "cot_tron":
                qs = list(circular_column(name, n, vals["duong_kinh"], vals["cao"], profile=profile))
            elif kind == "dao_hao":
                qs = [trench_excavation(name, vals["dai"], vals["rong"], vals["cao"], vals["mai"] or 0.0, profile=profile)]
            elif kind == "dao_ho":
                qs = [pit_excavation(name, vals["dai"], vals["rong"], vals["cao"], vals["mai"] or 0.0, profile=profile)]
            elif kind == "ong":
                qs = [pipe_length(name, vals["dai"], vals["tru_ho_ga"] or 0.0, _flag(get("thoat_nuoc"), True), profile=profile)]
            elif kind == "dan_giao_trong":
                qs = [scaffold_inner(name, vals["dien_tich_bang"], vals["cao"], profile=profile)]
            elif kind == "dan_giao_cot":
                qs = [scaffold_column(name, vals["chu_vi"], vals["cao"], profile=profile)]
            else:   # mat_cat: gộp sau
                sec = sections.setdefault(name, {"st": [], "ar": [], "refs": refs, "first": offset})
                sec["st"].append(vals["ly_trinh"])
                sec["ar"].append(vals["dien_tich"])
                continue
        except ValueError as e:
            errors.append(f"{where}: {e}")
            continue
        for q in qs:
            sub = ("dap_dau_coc" if q.name.endswith("đập đầu cọc") else
                   "van_khuon" if kind == "cot_tron" and q.unit == "m2" else kind)
            out.append(q.with_ref(kind=sub, **refs))

    for name, sec in sections.items():
        order = sorted(zip(sec["st"], sec["ar"]))
        try:
            q = average_end_volume(name, [s for s, _ in order], [a for _, a in order], profile=profile)
            out.append(q.with_ref(kind="mat_cat", **sec["refs"]))
        except ValueError as e:
            errors.append(f"Mặt cắt '{name}' (từ dòng {sec['first']}): {e}")

    if errors:
        raise TakeoffLoadError(f"{path}: {len(errors)} dòng sai dữ liệu:\n  - " + "\n  - ".join(errors))
    if not out:
        raise TakeoffLoadError(f"{path}: không có cấu kiện nào")
    return out


def write_takeoff_workbook(path: str, quantities: Sequence[Quantity], profile: MeasurementProfile,
                           project_name: str = "") -> None:
    """Xuất Excel: BANG_6_2_CHI_TIET, BANG_6_1_TONG_HOP, QUY_TAC_DO_BOC."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    hdr_font, hdr_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")

    def table(ws, title, rows, widths):
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append([f"Dự án: {project_name}" if project_name else ""])
        ws.append([])
        for r in rows:
            ws.append(r)
        for c in ws[4]:
            c.font, c.fill = hdr_font, hdr_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for i, w in enumerate(widths):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w

    ws = wb.active
    ws.title = "BANG_6_2_CHI_TIET"
    table(ws, "BẢNG CHI TIẾT KHỐI LƯỢNG CÔNG TÁC XÂY DỰNG (mẫu 6.2)", bang_6_2(quantities),
          [6, 14, 14, 40, 9, 12, 60, 16, 18, 50])
    for row in ws.iter_rows(min_row=5):
        row[7].number_format = row[8].number_format = "#,##0.000"
    ws2 = wb.create_sheet("BANG_6_1_TONG_HOP")
    table(ws2, "BẢNG TỔNG HỢP KHỐI LƯỢNG XÂY DỰNG (mẫu 6.1)", bang_6_1(quantities), [6, 14, 40, 9, 40, 18, 30])
    for row in ws2.iter_rows(min_row=5):
        row[5].number_format = "#,##0.000"
    ws3 = wb.create_sheet("QUY_TAC_DO_BOC")
    ws3.append(["Hồ sơ quy tắc", profile.name])
    ws3.append(["Nguồn", profile.source or "(chưa ghi)"])
    ws3.append(["Đã đối chiếu bản gốc", "CÓ" if profile.verified else "CHƯA"])
    ws3.append(["Số chữ số thập phân", profile.decimals])
    for k, v in sorted(profile.no_deduct_below.items()):
        ws3.append([f"Không trừ lỗ rỗng nhỏ hơn ({k})", v])
    if profile.rebar_no_deduct_below_ratio is not None:
        ws3.append(["Không trừ cốt thép khi hàm lượng nhỏ hơn", profile.rebar_no_deduct_below_ratio])
    for w in profile.warnings():
        ws3.append(["CẢNH BÁO", w])
    ws3.column_dimensions["A"].width, ws3.column_dimensions["B"].width = 42, 110
    wb.save(path)

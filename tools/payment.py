# -*- coding: utf-8 -*-
"""
PAYMENT 03A — Bảng xác định giá trị khối lượng công việc hoàn thành (Mẫu số 03a, NĐ 99/2021/NĐ-CP)

Pure Python, Zero LLM.
  Hợp đồng  : công tác, khối lượng, đơn giá từ bảng QS (tools/qs_loader)
  Tiến độ   : khối lượng lũy kế đến hết kỳ trước + khối lượng thực hiện kỳ này (file --progress)
  Kết quả   : giá trị hoàn thành kỳ này, lũy kế; VAT; thu hồi tạm ứng; giữ lại; đề nghị thanh toán

Nguyên tắc (không âm thầm cho qua):
  - Khối lượng lũy kế vượt khối lượng hợp đồng: phần vượt KHÔNG thanh toán theo Mẫu 03a,
    được liệt kê riêng (cần phụ lục hợp đồng / khối lượng phát sinh).
  - Công việc trong file tiến độ không khớp công tác hợp đồng: liệt kê riêng, không thanh toán.
  - Đơn giá: 'direct'  → đơn giá bảng QS là chi phí trực tiếp, nhân hệ số G/T (GT + TL) để ra
                         đơn giá trước thuế;
             'contract'→ đơn giá bảng QS đã là đơn giá hợp đồng trước thuế.
  - Tỷ lệ thu hồi tạm ứng và giữ lại phải khai báo (có thể 0); không có mặc định.

File tiến độ (Excel/CSV/JSON) — nhận diện cột theo tiêu đề:
  Mã hiệu / STT / Nội dung công tác  (khóa ghép với bảng QS, ưu tiên theo thứ tự này)
  Khối lượng lũy kế kỳ trước (tùy chọn, mặc định 0)
  Khối lượng thực hiện kỳ này
"""

from __future__ import annotations
import csv
import json
import os
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Sequence

from tools.bbs_loader import _FORMULA, _norm, _to_number
from tools.qs_loader import QSEstimate, QSItem

HEADER_SCAN_ROWS = 40


class PaymentError(ValueError):
    """Không lập được bảng 03a (thiếu cột, dữ liệu sai...)."""


@dataclass
class ProgressRow:
    row: int
    code: str
    stt: str
    description: str
    previous_qty: float
    this_qty: float


@dataclass
class PaymentLine:
    item: QSItem
    unit_price: int                # đơn giá hợp đồng trước thuế (đồng)
    contract_value: int
    previous_qty: float
    this_qty: float                # khối lượng kỳ này ĐƯỢC thanh toán
    cumulative_qty: float
    cumulative_value: int
    this_value: int
    overrun_qty: float = 0.0       # khối lượng kỳ này vượt hợp đồng (không thanh toán)
    note: str = ""


@dataclass
class PaymentResult:
    period: str
    price_basis: str
    price_factor: float
    lines: List[PaymentLine] = field(default_factory=list)
    unmatched: List[ProgressRow] = field(default_factory=list)
    errors: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    contract_value: int = 0
    cumulative_value: int = 0
    this_value: int = 0            # giá trị hoàn thành kỳ này, trước thuế
    vat_rate: float = 0.0
    this_vat: int = 0
    this_total: int = 0            # gồm VAT
    advance_recovery_pct: float = 0.0
    advance_recovery: int = 0
    retention_pct: float = 0.0
    retention: int = 0
    payable: int = 0               # đề nghị thanh toán kỳ này
    overrun_value: int = 0         # giá trị khối lượng vượt hợp đồng (chưa thanh toán)


# ─────────────────────────────────────────────────────────────────────────────
# ĐỌC FILE TIẾN ĐỘ THANH TOÁN
# ─────────────────────────────────────────────────────────────────────────────

def _detect_columns(headers: Sequence[Any]) -> Dict[str, int]:
    cols: Dict[str, int] = {}
    for idx, raw in enumerate(headers):
        if raw is None or raw is _FORMULA:
            continue
        h = _norm(raw)
        if not h:
            continue
        if "kytruoc" in h or h in ("previous", "prevqty", "cumulativeprevious"):
            cols.setdefault("previous", idx)
        elif ("kynay" in h and "luyke" not in h) or h in ("thisperiod", "thisqty", "qty", "quantity"):
            cols.setdefault("this", idx)
        elif "mahieu" in h or "madinhmuc" in h or h in ("ma", "code"):
            cols.setdefault("code", idx)
        elif h in ("tt", "stt", "no"):
            cols.setdefault("stt", idx)
        elif "noidung" in h or "tencongviec" in h or "tencongtac" in h or h in ("description", "congviec"):
            cols.setdefault("description", idx)
    return cols


def _find_header(rows):
    for r, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        cols = _detect_columns(row)
        if "this" in cols and any(k in cols for k in ("code", "stt", "description")):
            return r, cols
    return None


def load_progress(path: str, sheet: Optional[str] = None) -> List[ProgressRow]:
    if not os.path.exists(path):
        raise PaymentError(f"Không tìm thấy file khối lượng thực hiện: {path}")
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        from tools.excel_eval import WorkbookEvaluator
        ev = WorkbookEvaluator(path)
        if sheet and sheet not in ev.sheetnames:
            raise PaymentError(f"Không có sheet '{sheet}' trong {path}")
        for name in ([sheet] if sheet else ev.sheetnames):
            rows = ev.rows(name, missing=_FORMULA)
            if _find_header(rows):
                return _parse_progress(rows, f"{path}#{name}")
        raise PaymentError(f"Không tìm thấy bảng khối lượng thực hiện trong {path} — "
                           f"cần cột 'Khối lượng thực hiện kỳ này' và Mã hiệu / STT / Nội dung")
    if ext in (".csv", ".txt"):
        with open(path, "r", encoding="utf-8-sig", newline="") as f:
            text = f.read()
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return _parse_progress([r for r in csv.reader(text.splitlines(), dialect)], path)
    if ext == ".json":
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        items = data.get("items", []) if isinstance(data, dict) else data
        keys: List[str] = []
        for i in items:
            keys.extend(k for k in i if k not in keys)
        return _parse_progress([keys] + [[i.get(k) for k in keys] for i in items], path)
    raise PaymentError(f"Định dạng chưa hỗ trợ: {ext} (dùng .xlsx, .csv hoặc .json)")


def _parse_progress(rows, source: str) -> List[ProgressRow]:
    found = _find_header(rows)
    if found is None:
        raise PaymentError(f"{source}: cần cột 'Khối lượng thực hiện kỳ này' và Mã hiệu / STT / Nội dung")
    header_idx, cols = found
    out: List[ProgressRow] = []
    errors: List[str] = []

    def cell(row, key):
        idx = cols.get(key)
        if idx is None or idx >= len(row):
            return None
        v = row[idx]
        return None if isinstance(v, str) and not v.strip() else v

    for offset, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        raw_this, raw_prev = cell(row, "this"), cell(row, "previous")
        if raw_this is None and raw_prev is None:
            continue
        code = str(cell(row, "code") or "").strip()
        stt = str(cell(row, "stt") or "").strip()
        desc = str(cell(row, "description") or "").strip()
        where = f"Dòng {offset} ({code or stt} {desc[:40]})".strip()
        if _FORMULA in (raw_this, raw_prev):
            errors.append(f"{where}: công thức không tính được")
            continue
        this_qty = _to_number(raw_this) if raw_this is not None else 0.0
        prev_qty = _to_number(raw_prev) if raw_prev is not None else 0.0
        if this_qty is None or prev_qty is None or this_qty < 0 or prev_qty < 0:
            errors.append(f"{where}: khối lượng không hợp lệ (kỳ trước '{raw_prev}', kỳ này '{raw_this}')")
            continue
        if not (code or stt or desc):
            errors.append(f"{where}: thiếu Mã hiệu / STT / Nội dung để ghép với hợp đồng")
            continue
        out.append(ProgressRow(offset, code, stt, desc, prev_qty, this_qty))
    if errors:
        raise PaymentError(f"{source}: {len(errors)} dòng sai dữ liệu:\n    " + "\n    ".join(errors[:10]))
    return out


# ─────────────────────────────────────────────────────────────────────────────
# TÍNH 03A
# ─────────────────────────────────────────────────────────────────────────────

def _match(items: List[QSItem], p: ProgressRow) -> Optional[QSItem]:
    """Ghép theo Mã hiệu (nếu duy nhất), rồi STT, rồi nội dung. Mơ hồ → PaymentError."""
    for attr, value in (("code", p.code), ("stt", p.stt), ("description", p.description)):
        if not value:
            continue
        key = _norm(value)
        hits = [i for i in items if _norm(getattr(i, attr)) == key]
        if len(hits) == 1:
            return hits[0]
        if len(hits) > 1 and attr != "code":
            raise PaymentError(f"Dòng {p.row}: '{value}' khớp {len(hits)} công tác hợp đồng — ghi rõ STT")
    return None


def compute_payment(
    estimate: QSEstimate,
    progress: List[ProgressRow],
    price_basis: str,
    advance_recovery_pct: float,
    retention_pct: float,
    vat_rate: Optional[float] = None,
    advance_outstanding: Optional[float] = None,
    period: str = "",
) -> PaymentResult:
    if price_basis not in ("direct", "contract"):
        raise PaymentError("price_basis phải là 'direct' hoặc 'contract'")
    if price_basis == "direct":
        if not estimate.T:
            raise PaymentError("Cần tính G_XD (tỷ lệ GT, TL) trước khi quy đổi đơn giá trực tiếp")
        factor = estimate.G / estimate.T
    else:
        factor = 1.0
    vat = estimate.rates.get("vat") if vat_rate is None else vat_rate
    if vat is None:
        raise PaymentError("Thiếu thuế suất VAT")
    for label, pct in (("thu hồi tạm ứng", advance_recovery_pct), ("giữ lại", retention_pct)):
        if not 0 <= pct <= 100:
            raise PaymentError(f"Tỷ lệ {label} {pct}% không hợp lệ")

    res = PaymentResult(period=period, price_basis=price_basis, price_factor=factor, vat_rate=vat,
                        advance_recovery_pct=advance_recovery_pct, retention_pct=retention_pct)
    done: Dict[int, ProgressRow] = {}
    for p in progress:
        try:
            item = _match(estimate.items, p)
        except PaymentError as e:
            res.errors.append(str(e))
            continue
        if item is None:
            res.unmatched.append(p)
            continue
        if id(item) in done:
            res.errors.append(f"Dòng {p.row}: công tác '{item.code} {item.description[:40]}' đã có ở dòng "
                              f"{done[id(item)].row}")
            continue
        done[id(item)] = p

    for item in estimate.items:
        p = done.get(id(item))
        prev = p.previous_qty if p else 0.0
        this = p.this_qty if p else 0.0
        price = round(item.unit_price * factor)
        q = item.quantity
        payable = max(0.0, min(this, q - prev))
        overrun = this - payable
        note = ""
        if prev > q + 1e-9:
            res.warnings.append(f"{item.code} {item.description[:40]}: lũy kế kỳ trước {prev:,.3f} đã vượt "
                                f"khối lượng hợp đồng {q:,.3f}")
        if overrun > 1e-9:
            note = f"Vượt HĐ {overrun:,.3f} {item.unit} — chưa thanh toán"
            res.warnings.append(f"{item.code} {item.description[:40]}: khối lượng kỳ này vượt hợp đồng "
                                f"{overrun:,.3f} {item.unit} — cần phụ lục hợp đồng / phát sinh")
        cumulative = min(prev + payable, max(q, prev))
        line = PaymentLine(
            item=item, unit_price=price, contract_value=round(q * price), previous_qty=prev,
            this_qty=payable, cumulative_qty=cumulative, cumulative_value=round(min(cumulative, q) * price),
            this_value=round(payable * price), overrun_qty=overrun, note=note,
        )
        res.lines.append(line)
        res.overrun_value += round(overrun * price)

    for p in res.unmatched:
        res.warnings.append(f"Dòng {p.row} ({p.code or p.stt} {p.description[:40]}): không có trong hợp đồng — "
                            f"không thanh toán theo 03a")

    res.contract_value = sum(l.contract_value for l in res.lines)
    res.cumulative_value = sum(l.cumulative_value for l in res.lines)
    res.this_value = sum(l.this_value for l in res.lines)
    res.this_vat = round(res.this_value * vat)
    res.this_total = res.this_value + res.this_vat
    recovery = round(res.this_total * advance_recovery_pct / 100)
    if advance_outstanding is not None:
        recovery = min(recovery, round(advance_outstanding))
    res.advance_recovery = recovery
    res.retention = round(res.this_total * retention_pct / 100)
    res.payable = res.this_total - res.advance_recovery - res.retention
    if res.payable < 0:
        res.errors.append(f"Số đề nghị thanh toán âm ({res.payable:,}) — kiểm tra tỷ lệ khấu trừ")
    return res


# ─────────────────────────────────────────────────────────────────────────────
# XUẤT MẪU 03A
# ─────────────────────────────────────────────────────────────────────────────

def write_payment_workbook(path: str, res: PaymentResult, project_name: str = "") -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PHU_LUC_03A"
    ws["A1"] = "BẢNG XÁC ĐỊNH GIÁ TRỊ KHỐI LƯỢNG CÔNG VIỆC HOÀN THÀNH ĐỀ NGHỊ THANH TOÁN"
    ws["A2"] = f"Mẫu số 03a — Nghị định 99/2021/NĐ-CP. Kỳ thanh toán: {res.period or '...'}. {project_name}"
    ws["A3"] = ("Đơn giá hợp đồng trước thuế = đơn giá bảng QS × " + f"{res.price_factor:.6f} (hệ số G/T)"
                if res.price_basis == "direct" else "Đơn giá hợp đồng trước thuế = đơn giá bảng QS")
    headers = ["STT", "Mã hiệu", "Nội dung công việc theo hợp đồng", "Đơn vị tính",
               "Khối lượng theo HĐ (1)", "Đơn giá theo HĐ (2)", "Thành tiền theo HĐ (3=1x2)",
               "Khối lượng lũy kế đến hết kỳ trước (4)", "Khối lượng thực hiện kỳ này (5)",
               "Khối lượng lũy kế đến hết kỳ này (6=4+5)", "Giá trị lũy kế đến hết kỳ này (7=6x2)",
               "Giá trị đề nghị thanh toán kỳ này (8=5x2)", "Ghi chú"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5E")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    section = None
    for line in res.lines:
        i = line.item
        if i.section != section:
            section = i.section
            ws.append([section])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.append([i.stt, i.code, i.description, i.unit, i.quantity, line.unit_price, line.contract_value,
                   line.previous_qty, line.this_qty, line.cumulative_qty, line.cumulative_value,
                   line.this_value, line.note])
    ws.append(["", "", "TỔNG CỘNG (trước thuế)", "", "", "", res.contract_value, "", "", "",
               res.cumulative_value, res.this_value, ""])
    ws.cell(row=ws.max_row, column=3).font = Font(bold=True)
    for row in ws.iter_rows(min_row=6, max_row=ws.max_row):
        for c in (5, 8, 9, 10):
            row[c - 1].number_format = "#,##0.000"
        for c in (6, 7, 11, 12):
            row[c - 1].number_format = "#,##0"

    r = ws.max_row + 2
    summary = [
        ("1", "Giá trị khối lượng hoàn thành kỳ này (trước thuế)", res.this_value),
        ("2", f"Thuế GTGT ({res.vat_rate * 100:g}%)", res.this_vat),
        ("3", "Giá trị hoàn thành kỳ này (gồm thuế) (3 = 1 + 2)", res.this_total),
        ("4", f"Thu hồi tạm ứng ({res.advance_recovery_pct:g}% giá trị kỳ này)", res.advance_recovery),
        ("5", f"Giữ lại ({res.retention_pct:g}% giá trị kỳ này)", res.retention),
        ("6", "SỐ TIỀN ĐỀ NGHỊ THANH TOÁN KỲ NÀY (6 = 3 − 4 − 5)", res.payable),
        ("", "Giá trị khối lượng vượt hợp đồng (chưa thanh toán, cần phụ lục / phát sinh)", res.overrun_value),
    ]
    ws.cell(row=r, column=2, value="TỔNG HỢP GIÁ TRỊ ĐỀ NGHỊ THANH TOÁN").font = Font(bold=True)
    for k, (no, label, value) in enumerate(summary, start=r + 1):
        ws.cell(row=k, column=2, value=no)
        ws.cell(row=k, column=3, value=label)
        ws.cell(row=k, column=12, value=value).number_format = "#,##0"
    ws.column_dimensions["C"].width = 60
    for col in "EFGHIJKL":
        ws.column_dimensions[col].width = 16

    if res.unmatched or any(l.overrun_qty > 1e-9 for l in res.lines) or res.warnings:
        ws2 = wb.create_sheet("PHAT_SINH_CANH_BAO")
        ws2["A1"] = "KHỐI LƯỢNG VƯỢT HỢP ĐỒNG / NGOÀI HỢP ĐỒNG — KHÔNG THANH TOÁN THEO MẪU 03A"
        ws2.append([])
        ws2.append(["Loại", "Mã hiệu / STT", "Nội dung", "Khối lượng", "Đơn giá HĐ", "Giá trị"])
        for l in res.lines:
            if l.overrun_qty > 1e-9:
                ws2.append(["Vượt HĐ", l.item.code or l.item.stt, l.item.description, l.overrun_qty,
                            l.unit_price, round(l.overrun_qty * l.unit_price)])
        for p in res.unmatched:
            ws2.append(["Ngoài HĐ", p.code or p.stt, p.description, p.this_qty, None, None])
        ws2.append([])
        for w in res.warnings:
            ws2.append(["Cảnh báo", "", w])
        ws2.column_dimensions["C"].width = 70
    wb.save(path)

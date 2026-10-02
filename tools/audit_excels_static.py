# -*- coding: utf-8 -*-
"""
AUDIT EXCEL TĨNH (không cần Microsoft Excel / win32com — chạy được trên Linux, macOS, CI)

Quét mọi .xlsx dưới một thư mục và báo cáo từng file:
  - Lỗi đã lưu sẵn trong ô (#REF!, #VALUE!, #DIV/0!, #N/A, #NAME?, #NUM!, #NULL!)
  - Công thức có chứa #REF! ngay trong chuỗi công thức, hoặc trỏ tới sheet không tồn tại
  - File 'vỏ rỗng' (không công thức, ≤ 5 ô dữ liệu)
  - Defined Names gãy (#REF!) và tổng số Defined Names (phát hiện "name rác")
  - Công thức chưa có giá trị lưu sẵn (file chưa từng được Excel tính) — thử tự tính bằng
    tools/excel_eval.py; hàm chưa hỗ trợ được tính riêng, không coi là lỗi
  - Liên kết ngoài (external links) và tham chiếu tới file khác
  - Số công thức / số ô hằng số (tham khảo, không phải lỗi)
  - Trùng nội dung giữa các file (MD5)

Giới hạn: đây là kiểm tra tĩnh + bộ tính tự viết, KHÔNG thay cho Excel thật. Ô công thức dùng hàm
ngoài danh sách của excel_eval chỉ được kiểm tra khi file đã có giá trị Excel lưu sẵn.

Dùng:
    python -m tools.audit_excels_static examples
    python -m tools.audit_excels_static examples --json bao_cao.json --no-eval
Mã thoát 1 nếu có lỗi (dùng được trong CI); --strict coi cả ô chưa tính được là lỗi.
"""

from __future__ import annotations
import argparse
import hashlib
import json
import os
import re
import sys
from typing import Any, Dict, List

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

import openpyxl

_SHEET_REF = re.compile(r"(?:'([^']+)'|([A-Za-z0-9_\u00C0-\u1EF9.]+))!")
ERROR_LITERALS = ("#REF!", "#VALUE!", "#DIV/0!", "#N/A", "#NAME?", "#NUM!", "#NULL!")
MAX_DETAIL = 10          # số ví dụ lỗi lưu cho mỗi loại / file
NAME_JUNK_THRESHOLD = 500


def _md5(path: str) -> str:
    h = hashlib.md5()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def find_xlsx(root: str) -> List[str]:
    out = []
    for dp, _, files in os.walk(root):
        for fn in files:
            if fn.lower().endswith((".xlsx", ".xlsm")) and not fn.startswith("~$"):
                out.append(os.path.join(dp, fn))
    return sorted(out)


def audit_file(path: str, try_eval: bool = True) -> Dict[str, Any]:
    res: Dict[str, Any] = {
        "file": path, "size_kb": round(os.path.getsize(path) / 1024, 1), "md5": _md5(path),
        "sheets": 0, "formulas": 0, "constants": 0,
        "cached_errors": 0, "formula_ref_errors": 0, "missing_sheet_refs": 0, "empty_shell": False, "uncached_formulas": 0,
        "eval_ok": 0, "eval_unsupported": 0, "eval_error_results": 0,
        "defined_names": 0, "broken_names": 0, "external_links": 0,
        "details": [], "fatal": None,
    }

    def note(kind: str, where: str, text: str = ""):
        if sum(1 for d in res["details"] if d["kind"] == kind) < MAX_DETAIL:
            res["details"].append({"kind": kind, "where": where, "text": text[:160]})

    try:
        fwb = openpyxl.load_workbook(path, data_only=False)
        vwb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:  # file hỏng / không mở được là một kết quả audit
        res["fatal"] = f"{type(e).__name__}: {e}"
        return res

    res["sheets"] = len(fwb.worksheets)
    res["external_links"] = len(getattr(fwb, "_external_links", []) or [])
    try:
        names = list(fwb.defined_names.items())
    except AttributeError:
        names = []
    res["defined_names"] = len(names)
    for nm, dn in names:
        if "#REF!" in str(getattr(dn, "attr_text", "")):
            res["broken_names"] += 1
            note("broken_name", nm, str(dn.attr_text))

    uncached: List[tuple] = []
    for ws in fwb.worksheets:
        vws = vwb[ws.title]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    continue
                if isinstance(v, str) and v.startswith("="):
                    res["formulas"] += 1
                    if "#REF!" in v:
                        res["formula_ref_errors"] += 1
                        note("formula_ref", f"{ws.title}!{c.coordinate}", v)
                    for m in _SHEET_REF.finditer(v):
                        ref = m.group(1) or m.group(2)
                        if ref not in fwb.sheetnames:
                            res["missing_sheet_refs"] += 1
                            note("missing_sheet", f"{ws.title}!{c.coordinate}", f"'{ref}' <- {v}")
                            break
                    cv = vws[c.coordinate].value
                    if isinstance(cv, str) and cv in ERROR_LITERALS:
                        res["cached_errors"] += 1
                        note("cached_error", f"{ws.title}!{c.coordinate}", f"{cv} <- {v}")
                    elif cv is None:
                        res["uncached_formulas"] += 1
                        uncached.append((ws.title, c.row, c.column, c.coordinate, v))
                else:
                    res["constants"] += 1
                    if isinstance(v, str) and v in ERROR_LITERALS:
                        res["cached_errors"] += 1
                        note("cached_error", f"{ws.title}!{c.coordinate}", v)

    res["empty_shell"] = (res["formulas"] == 0 and res["constants"] <= 5)
    if try_eval and uncached:
        from tools.excel_eval import WorkbookEvaluator, FormulaError, ExcelErrorResult
        try:
            ev = WorkbookEvaluator(path)
        except Exception as e:
            ev = None
            note("eval_unavailable", "", f"{type(e).__name__}: {e}")
        if ev is not None:
            for sheet, r, c, coord, f in uncached:
                try:
                    val = ev.value(sheet, r, c)
                    if isinstance(val, str) and val in ERROR_LITERALS:
                        res["eval_error_results"] += 1
                        note("eval_error", f"{sheet}!{coord}", f"{val} <- {f}")
                    else:
                        res["eval_ok"] += 1
                except ExcelErrorResult as e:
                    res["eval_error_results"] += 1
                    note("eval_error", f"{sheet}!{coord}", f"{e.code} <- {f}")
                except FormulaError as e:
                    res["eval_unsupported"] += 1
                    note("eval_unsupported", f"{sheet}!{coord}", str(e))
                except ZeroDivisionError:
                    res["eval_error_results"] += 1
                    note("eval_error", f"{sheet}!{coord}", f"#DIV/0! <- {f}")
                except Exception as e:
                    res["eval_unsupported"] += 1
                    note("eval_unsupported", f"{sheet}!{coord}", f"{type(e).__name__}: {e}")
    return res


def hard_errors(r: Dict[str, Any], strict: bool = False) -> int:
    n = (r["cached_errors"] + r["formula_ref_errors"] + r["broken_names"] + r["missing_sheet_refs"]
         + r["eval_error_results"] + (1 if r["fatal"] else 0))
    if strict:
        n += r["eval_unsupported"] + (r["uncached_formulas"] - r["eval_ok"]
                                      - r["eval_unsupported"] - r["eval_error_results"])
    return n


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Audit tĩnh toàn bộ file Excel (không cần Excel COM)")
    ap.add_argument("root", nargs="?", default="examples", help="thư mục cần quét (mặc định: examples)")
    ap.add_argument("--json", help="ghi báo cáo chi tiết ra file JSON")
    ap.add_argument("--no-eval", action="store_true", help="không tự tính ô chưa có giá trị lưu sẵn")
    ap.add_argument("--strict", action="store_true", help="coi ô không tính được / chưa tính là lỗi")
    args = ap.parse_args(argv)

    files = find_xlsx(args.root)
    print("=" * 78)
    print(f"AUDIT EXCEL TĨNH — {len(files)} file .xlsx dưới '{args.root}'")
    print("=" * 78)

    results = []
    for i, p in enumerate(files, 1):
        r = audit_file(p, try_eval=not args.no_eval)
        results.append(r)
        bad = hard_errors(r, args.strict)
        flag = "OK " if bad == 0 else "LỖI"
        rel = os.path.relpath(p, args.root)
        extra = []
        if r["uncached_formulas"]:
            extra.append(f"chưa-tính={r['uncached_formulas']} (tự tính được {r['eval_ok']}, "
                         f"hàm chưa hỗ trợ {r['eval_unsupported']})")
        if r["defined_names"] > NAME_JUNK_THRESHOLD:
            extra.append(f"names={r['defined_names']} (nhiều bất thường)")
        if r["external_links"]:
            extra.append(f"liên-kết-ngoài={r['external_links']}")
        if r["empty_shell"]:
            extra.append("RỖNG: gần như không có dữ liệu (vỏ file)")
        if r["fatal"]:
            extra.append(f"không mở được: {r['fatal']}")
        print(f"[{i:>2}/{len(files)}] {flag} {rel}  | công thức={r['formulas']} hằng={r['constants']} "
              f"lỗi={bad}" + ("  | " + "; ".join(extra) if extra else ""))
        for d in r["details"]:
            if d["kind"] in ("cached_error", "formula_ref", "broken_name", "eval_error", "missing_sheet"):
                print(f"        - {d['kind']}: {d['where']} {d['text']}")

    by_md5: Dict[str, List[str]] = {}
    for r in results:
        by_md5.setdefault(r["md5"], []).append(r["file"])
    dups = {k: v for k, v in by_md5.items() if len(v) > 1}

    total_bad = sum(hard_errors(r, args.strict) for r in results)
    bad_files = sum(1 for r in results if hard_errors(r, args.strict))
    print("-" * 78)
    print(f"Tổng: {len(files)} file | {len(files) - bad_files} sạch | {bad_files} có lỗi | "
          f"{total_bad} lỗi")
    print(f"File rỗng (vỏ, không dữ liệu): {sum(1 for r in results if r['empty_shell'])}")
    print(f"Công thức: {sum(r['formulas'] for r in results)} | hằng số: {sum(r['constants'] for r in results)}")
    print(f"Ô công thức chưa có giá trị lưu sẵn: {sum(r['uncached_formulas'] for r in results)} "
          f"(tự tính được {sum(r['eval_ok'] for r in results)}, "
          f"hàm chưa hỗ trợ {sum(r['eval_unsupported'] for r in results)})")
    print(f"File trùng nội dung (cùng MD5): {sum(len(v) for v in dups.values())} file "
          f"trong {len(dups)} nhóm")
    if args.json:
        with open(args.json, "w", encoding="utf-8") as f:
            json.dump({"results": results, "duplicate_groups": list(dups.values())},
                      f, ensure_ascii=False, indent=2)
        print(f"Đã ghi báo cáo: {args.json}")
    return 1 if total_bad else 0


if __name__ == "__main__":
    sys.exit(main())

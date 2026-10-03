"""
Cập nhật tiến độ toàn tuyến Cống hộp A5 sang mốc: 05/09/2026 đến 12/11/2026 (69 ngày).
Giữ nguyên vẹn 100% bố cục 3 tầng sống động (Tasks + Mobilization + Diesel) và công thức sống.
"""
import os
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_FILES = [
    os.path.join(ROOT, "examples/HO_SO_CONG_HOP_TUYEN_A5/HO_SO_THUC_CHIEN_HUB_AND_SPOKE_CONG_A5/GOI_A_CO_GIOI_VA_DAU_DIEZEL/260920_TDTC_CaXe_CaMay_Cong_Hop_Tuyen_A5.xlsx"),
    os.path.join(ROOT, "examples/HO_SO_CONG_HOP_TUYEN_A5/HUB_VINA_ALPHA_05-09_12-11/HUB/03_HO_SO_THUC_CHIEN_HUB_AND_SPOKE_5_GOI_VE_TINH/GOI_A_CO_GIOI_VA_DAU_DIEZEL/TDTC_CaXe_CaMay_DauDiezel_Cong_Hop_A5.xlsx")
]

D_START = datetime.date(2026, 9, 5)
D_FINISH = datetime.date(2026, 11, 12)
TOTAL_DAYS = (D_FINISH - D_START).days + 1  # 69 ngày
DATES = [D_START + datetime.timedelta(days=i) for i in range(TOTAL_DAYS)]

# Colors & Fills
FILL_NAVY = PatternFill("solid", fgColor="001B365D")
FILL_BLUE = PatternFill("solid", fgColor="002E75B6")
FILL_GREEN = PatternFill("solid", fgColor="00385723")
FILL_YELLOW = PatternFill("solid", fgColor="00FFF2CC")
FILL_RED_HDR = PatternFill("solid", fgColor="00C00000")
FILL_CORAL_HDR = PatternFill("solid", fgColor="00FCE4D6")
FILL_GANTT_BLUE = PatternFill("solid", fgColor="00BDD7EE")
FILL_GANTT_ORANGE = PatternFill("solid", fgColor="00FCE4D6")
FILL_MACH_GREEN = PatternFill("solid", fgColor="00E2EFDA")
FILL_GREY_LIGHT = PatternFill("solid", fgColor="00F2F2F2")

FONT_WHITE_13 = Font(name="Arial", size=13, bold=True, color="00FFFFFF")
FONT_WHITE_11 = Font(name="Arial", size=11, bold=True, color="00FFFFFF")
FONT_WHITE_85 = Font(name="Arial", size=8.5, bold=True, color="00FFFFFF")
FONT_WHITE_8 = Font(name="Arial", size=8, bold=True, color="00FFFFFF")
FONT_RED_TITLE = Font(name="Arial", size=10, bold=True, color="00C00000")
FONT_RED_85 = Font(name="Arial", size=8.5, bold=True, color="00C00000")
FONT_RED_8 = Font(name="Arial", size=8, bold=True, color="00C00000")
FONT_NAVY_BOLD = Font(name="Arial", size=8.5, bold=True, color="00002060")
FONT_MACH_VAL = Font(name="Arial", size=8, bold=True, color="001B365D")
FONT_DIESEL_VAL = Font(name="Arial", size=7.5, bold=False, color="00333333")
FONT_GANTT_BLUE = Font(name="Arial", size=8, bold=False, color="001B365D")
FONT_REGULAR_85 = Font(name="Arial", size=8.5, bold=False)

THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

WEEKDAYS_VN = ["T2", "T3", "T4", "T5", "T6", "T7", "CN"]

# 17 Công tác WBS
TASKS = [
    {
        "row": 8, "wbs": "ĐM-16", "name": "Đào đất hố móng cống hộp 3 mũi (H_tb = 3.5m)",
        "unit": "m3", "qty": 73800.0, "norm": 401.07, "shifts": 2,
        "start": datetime.date(2026, 9, 5), "finish": datetime.date(2026, 10, 4),
        "mach": "Máy xúc PC200 - PC300", "crew": 24, "val": 3.07, "critical": False
    },
    {
        "row": 9, "wbs": "ĐM-15", "name": "Đào hố móng 63 hố ga thu nước & đấu nối",
        "unit": "m3", "qty": 12650.0, "norm": 313.55, "shifts": 2,
        "start": datetime.date(2026, 9, 8), "finish": datetime.date(2026, 9, 30),
        "mach": "Máy xúc PC200", "crew": 12, "val": 0.88, "critical": False
    },
    {
        "row": 10, "wbs": "ĐM-29", "name": "Vận chuyển đất đào cự ly <5km (bãi thải/đắp)",
        "unit": "m3", "qty": 86450.0, "norm": 571.43, "shifts": 2,
        "start": datetime.date(2026, 9, 5), "finish": datetime.date(2026, 10, 4),
        "mach": "Ô tô tự đổ Howo 18 m3", "crew": 10, "val": 2.52, "critical": False
    },
    {
        "row": 11, "wbs": "ĐM-BOM", "name": "Bơm hạ mực nước ngầm & chống ngập móng",
        "unit": "ca", "qty": 138.0, "norm": 1.0, "shifts": 2,
        "start": datetime.date(2026, 9, 5), "finish": datetime.date(2026, 11, 12),
        "mach": "Máy bơm nước d80-d100", "crew": 4, "val": 1.0, "critical": False
    },
    {
        "row": 12, "wbs": "ĐM-53", "name": "Đệm cát / đá dăm 4x6 lót móng cống",
        "unit": "m3", "qty": 1250.0, "norm": 150.0, "shifts": 2,
        "start": datetime.date(2026, 9, 10), "finish": datetime.date(2026, 10, 8),
        "mach": "Đầm cóc + đầm bàn", "crew": 12, "val": 0.14, "critical": False
    },
    {
        "row": 13, "wbs": "ĐM-37", "name": "Bê tông lót móng M100# dày 50mm",
        "unit": "m3", "qty": 592.4, "norm": 45.0, "shifts": 2,
        "start": datetime.date(2026, 9, 12), "finish": datetime.date(2026, 10, 12),
        "mach": "Xe bồn + đầm bàn", "crew": 16, "val": 0.21, "critical": False
    },
    {
        "row": 14, "wbs": "ĐM-CT", "name": "Gia công & lắp dựng cốt thép cống hộp B20",
        "unit": "tấn", "qty": 753.36, "norm": 12.0, "shifts": 2,
        "start": datetime.date(2026, 9, 16), "finish": datetime.date(2026, 10, 26),
        "mach": "Cần cẩu 25T & Máy uốn cắt", "crew": 60, "val": 0.77, "critical": True
    },
    {
        "row": 15, "wbs": "ĐM-VK", "name": "Lắp dựng ván khuôn thép/phủ phim thân cống",
        "unit": "m2", "qty": 44940.7, "norm": 450.0, "shifts": 2,
        "start": datetime.date(2026, 9, 18), "finish": datetime.date(2026, 10, 28),
        "mach": "Cần cẩu 15T & Dàn giáo", "crew": 70, "val": 1.22, "critical": True
    },
    {
        "row": 16, "wbs": "ĐM-BT", "name": "Đổ bê tông thân cống B20 (đáy, vách, nắp)",
        "unit": "m3", "qty": 11870.0, "norm": 150.0, "shifts": 2,
        "start": datetime.date(2026, 9, 20), "finish": datetime.date(2026, 10, 31),
        "mach": "Máy bơm cần 37-43m + đầm", "crew": 40, "val": 0.94, "critical": True
    },
    {
        "row": 17, "wbs": "ĐM-VCBT", "name": "Vận chuyển bê tông thương phẩm B20 trạm trộn",
        "unit": "m3", "qty": 11870.0, "norm": 45.0, "shifts": 2,
        "start": datetime.date(2026, 9, 20), "finish": datetime.date(2026, 10, 31),
        "mach": "Xe bồn 8-10m3 (6 xe)", "crew": 12, "val": 3.14, "critical": True
    },
    {
        "row": 18, "wbs": "ĐM-46", "name": "Bê tông & cốt thép 63 hố ga BTCT",
        "unit": "m3", "qty": 409.5, "norm": 25.0, "shifts": 2,
        "start": datetime.date(2026, 9, 25), "finish": datetime.date(2026, 10, 25),
        "mach": "Xe bồn + đầm bàn", "crew": 22, "val": 0.26, "critical": False
    },
    {
        "row": 19, "wbs": "ĐM-25", "name": "Máy xúc lốp PC140 cẩu lắp & đầm cóc mang cống",
        "unit": "md", "qty": 2170.0, "norm": 100.0, "shifts": 2,
        "start": datetime.date(2026, 9, 20), "finish": datetime.date(2026, 10, 31),
        "mach": "Máy xúc lốp PC140", "crew": 8, "val": 0.26, "critical": False
    },
    {
        "row": 20, "wbs": "ĐM-17", "name": "Đắp cát/đất K95 hoàn trả mang cống",
        "unit": "m3", "qty": 54200.0, "norm": 255.0, "shifts": 2,
        "start": datetime.date(2026, 10, 20), "finish": datetime.date(2026, 11, 10),
        "mach": "Đầm cóc + xúc lật mang cống", "crew": 28, "val": 4.83, "critical": False
    },
    {
        "row": 21, "wbs": "ĐM-ỦI", "name": "Máy ủi D3-D5 san gạt hoàn trả đỉnh móng",
        "unit": "m3", "qty": 54200.0, "norm": 1425.0, "shifts": 2,
        "start": datetime.date(2026, 10, 20), "finish": datetime.date(2026, 11, 10),
        "mach": "Máy ủi bánh xích D3-D5", "crew": 4, "val": 0.86, "critical": False
    },
    {
        "row": 22, "wbs": "ĐM-LU", "name": "Máy lu rung 12-16T đầm nén K95 hoàn trả",
        "unit": "m3", "qty": 54200.0, "norm": 255.0, "shifts": 2,
        "start": datetime.date(2026, 10, 21), "finish": datetime.date(2026, 11, 11),
        "mach": "Máy lu rung 12T - 16T", "crew": 12, "val": 4.83, "critical": False
    },
    {
        "row": 23, "wbs": "ĐM-DAU", "name": "Xe téc cấp dầu lưu động 9m3 phục vụ máy móc",
        "unit": "ca", "qty": 138.0, "norm": 1.0, "shifts": 2,
        "start": datetime.date(2026, 9, 5), "finish": datetime.date(2026, 11, 12),
        "mach": "Xe téc cấp dầu 9 m3", "crew": 2, "val": 1.0, "critical": False
    },
    {
        "row": 24, "wbs": "ĐM-ĐIỆN", "name": "Máy phát điện 3 pha công nghiệp 25-45kVA",
        "unit": "ca", "qty": 138.0, "norm": 1.0, "shifts": 2,
        "start": datetime.date(2026, 9, 5), "finish": datetime.date(2026, 11, 12),
        "mach": "Máy phát điện 25-45kVA", "crew": 2, "val": 1.0, "critical": False
    }
]

# 11 Chủng loại MMTB
MACHINES = [
    {"row_m": 30, "row_f": 44, "code": "M1", "name": "Máy xúc bánh xích PC200 - PC300", "unit": "cái", "fuel_l": 89.6, "max_m": 8, "desc": "Đào móng, đào ga, xúc mang cống"},
    {"row_m": 31, "row_f": 45, "code": "M7", "name": "Ô tô tự đổ 18 m3", "unit": "xe", "fuel_l": 26.0, "max_m": 5, "desc": "Vận chuyển đất đào cự ly <5km"},
    {"row_m": 32, "row_f": 46, "code": "M5", "name": "Máy lu rung 12T - 16T", "unit": "cái", "fuel_l": 44.0, "max_m": 6, "desc": "Đầm nén K95 mang cống và đỉnh cống"},
    {"row_m": 33, "row_f": 47, "code": "M3", "name": "Máy ủi bánh xích D3 - D5", "unit": "cái", "fuel_l": 124.0, "max_m": 2, "desc": "San gạt hoàn trả hố móng cống"},
    {"row_m": 34, "row_f": 48, "code": "MC", "name": "Cần cẩu 25T & Cẩu tự hành 15T", "unit": "cái", "fuel_l": 28.0, "max_m": 3, "desc": "Hạ ván khuôn, cẩu lắp cốt thép"},
    {"row_m": 35, "row_f": 49, "code": "MB", "name": "Máy bơm bê tông cần 37 - 43m", "unit": "cái", "fuel_l": 65.0, "max_m": 2, "desc": "Đổ bê tông thân cống B20 & hố ga"},
    {"row_m": 36, "row_f": 50, "code": "XB", "name": "Xe bồn vận chuyển bê tông 8-10m3", "unit": "xe", "fuel_l": 42.0, "max_m": 6, "desc": "Chở bê tông từ trạm trộn về công trường"},
    {"row_m": 37, "row_f": 51, "code": "M8", "name": "Máy xúc bánh lốp PC140", "unit": "cái", "fuel_l": 64.0, "max_m": 1, "desc": "Cẩu lắp, đầm cóc mang cống"},
    {"row_m": 38, "row_f": 52, "code": "BP", "name": "Máy bơm nước hố móng d80-d100", "unit": "cái", "fuel_l": 14.0, "max_m": 2, "desc": "Bơm hạ mực nước ngầm 24/7"},
    {"row_m": 39, "row_f": 53, "code": "MP1", "name": "Xe téc cấp dầu lưu động 9m3", "unit": "xe", "fuel_l": 44.0, "max_m": 1, "desc": "Cấp dầu lưu động 2 ca/ngày"},
    {"row_m": 40, "row_f": 54, "code": "MP2", "name": "Máy phát điện 3 pha 25-45kVA", "unit": "cái", "fuel_l": 28.0, "max_m": 2, "desc": "Chiếu sáng ban đêm 2 ca"}
]


def update_workbook(wb_path: str):
    print(f"[*] Đang cập nhật tiến độ 05/09/2026 - 12/11/2026 cho: {wb_path}")
    wb = openpyxl.load_workbook(wb_path, data_only=False)
    ws1 = wb["01_TienDo_CaMay_Master"]

    # 1. Cập nhật Tiêu đề Row 3
    ws1["A3"] = f"MỐC TIẾN ĐỘ THI CÔNG TOÀN TUYẾN: TỪ 05/09/2026 ĐẾN 12/11/2026 (69 NGÀY) - 2 CA/NGÀY (20H/NGÀY) - 3 MŨI THI CÔNG ĐỒNG THỜI"
    ws1["A3"].font = FONT_RED_TITLE
    ws1["A3"].fill = FILL_YELLOW

    # 2. Xóa các cột timeline cũ (từ P đến cột cuối hiện tại)
    max_c = ws1.max_column
    for r in range(6, 56):
        for c in range(16, max_c + 1):
            cell = ws1.cell(r, c)
            cell.value = None
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border()

    # 3. Tạo Headers Timeline 69 ngày (Cột P=16 đến CF=84)
    for i, dt in enumerate(DATES):
        col = 16 + i
        col_letter = get_column_letter(col)
        ws1.column_dimensions[col_letter].width = 6.8

        is_sun = (dt.weekday() == 6)
        # Row 6: Ngày
        c6 = ws1.cell(6, col, dt.strftime("%d/%m"))
        c6.font = FONT_WHITE_8
        c6.fill = FILL_RED_HDR if is_sun else FILL_NAVY
        c6.alignment = Alignment(horizontal="center", vertical="center")
        c6.border = THIN_BORDER

        # Row 7: Thứ
        c7 = ws1.cell(7, col, WEEKDAYS_VN[dt.weekday()])
        c7.font = FONT_RED_8 if is_sun else Font(name="Arial", size=8, bold=True)
        c7.fill = FILL_CORAL_HDR if is_sun else FILL_GREY_LIGHT
        c7.alignment = Alignment(horizontal="center", vertical="center")
        c7.border = THIN_BORDER

        # Row 29: Header cho Bảng 2
        c29 = ws1.cell(29, col, dt.strftime("%d/%m"))
        c29.font = FONT_WHITE_8
        c29.fill = FILL_BLUE
        c29.alignment = Alignment(horizontal="center", vertical="center")
        c29.border = THIN_BORDER

    # 4. Ghi thông số 17 Task & Gantt
    for t in TASKS:
        r = t["row"]
        days = (t["finish"] - t["start"]).days + 1
        ws1.cell(r, 1, r - 7).alignment = Alignment(horizontal="center")
        ws1.cell(r, 2, t["wbs"]).alignment = Alignment(horizontal="center")
        ws1.cell(r, 3, t["name"])
        ws1.cell(r, 4, t["unit"]).alignment = Alignment(horizontal="center")
        ws1.cell(r, 5, t["qty"]).number_format = "#,##0.00" if isinstance(t["qty"], float) else "#,##0"
        ws1.cell(r, 6, t["norm"]).number_format = "#,##0.00"

        # Cột 7: Tổng số ca máy
        if t["unit"] == "ca":
            ws1.cell(r, 7, t["qty"]).number_format = "#,##0.0"
        else:
            ws1.cell(r, 7, round(t["qty"] / t["norm"], 1)).number_format = "#,##0.0"

        # Cột 8: Năng xuất ngày
        ws1.cell(r, 8, round(t["qty"] / days, 1)).number_format = "#,##0.0"
        # Cột 9: Thời gian (ngày)
        ws1.cell(r, 9, days).number_format = "#,##0"
        ws1.cell(r, 9).alignment = Alignment(horizontal="center")
        # Cột 10, 11: Ngày BĐ, Ngày KT
        c10 = ws1.cell(r, 10, t["start"])
        c10.number_format = "DD/MM/YYYY"
        c10.alignment = Alignment(horizontal="center")
        c11 = ws1.cell(r, 11, t["finish"])
        c11.number_format = "DD/MM/YYYY"
        c11.alignment = Alignment(horizontal="center")

        ws1.cell(r, 12, t["shifts"]).alignment = Alignment(horizontal="center")
        ws1.cell(r, 13, t["val"]).number_format = "0.00" if isinstance(t["val"], float) else "0"
        ws1.cell(r, 13).alignment = Alignment(horizontal="center")
        ws1.cell(r, 14, t["mach"])
        ws1.cell(r, 15, t["crew"]).alignment = Alignment(horizontal="center")

        # Tô Gantt cho task này
        for i, dt in enumerate(DATES):
            col = 16 + i
            cell = ws1.cell(r, col)
            cell.border = THIN_BORDER
            if t["start"] <= dt <= t["finish"]:
                cell.value = t["val"]
                cell.number_format = "0.00" if isinstance(t["val"], float) else "0"
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = FONT_GANTT_BLUE
                cell.fill = FILL_GANTT_ORANGE if t["critical"] else FILL_GANTT_BLUE
            else:
                cell.value = 0
                cell.number_format = "0"
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = Font(name="Arial", size=7, color="D9D9D9")

    # 5. Row 26: Tổng nhân công trên công trường toàn tuyến (Người/ngày)
    ws1.cell(26, 3, "TỔNG NHÂN CÔNG TRÊN CÔNG TRƯỜNG TUYẾN A5 (Người/ngày)")
    ws1.cell(26, 3).font = Font(name="Arial", size=10, bold=True, color="00FFFFFF")
    ws1.cell(26, 3).fill = FILL_NAVY
    for col_idx in range(1, 16):
        ws1.cell(26, col_idx).fill = FILL_NAVY

    for i, dt in enumerate(DATES):
        col = 16 + i
        tot_labor = sum(t["crew"] for t in TASKS if t["start"] <= dt <= t["finish"])
        cell = ws1.cell(26, col, tot_labor)
        cell.font = FONT_NAVY_BOLD
        cell.fill = FILL_YELLOW
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    # 6. Bảng 2: Tổng hợp ca máy & Phương tiện huy động theo ngày (R28..R40)
    for m in MACHINES:
        r_m = m["row_m"]
        ws1.cell(r_m, 1, m["code"]).alignment = Alignment(horizontal="center")
        ws1.cell(r_m, 3, m["name"])
        ws1.cell(r_m, 4, m["unit"]).alignment = Alignment(horizontal="center")
        ws1.cell(r_m, 5, m["fuel_l"]).number_format = "#,##0.0"
        ws1.cell(r_m, 6, m["max_m"]).alignment = Alignment(horizontal="center")
        ws1.cell(r_m, 7, m["desc"])

        for i, dt in enumerate(DATES):
            col = 16 + i
            cell = ws1.cell(r_m, col)
            cell.border = THIN_BORDER
            # Xác định số máy hoạt động
            active_m = 0
            code = m["code"]
            if code == "M1":  # Máy xúc
                if datetime.date(2026, 9, 5) <= dt <= datetime.date(2026, 10, 4):
                    active_m = 8
                elif datetime.date(2026, 10, 20) <= dt <= datetime.date(2026, 11, 10):
                    active_m = 4
            elif code == "M7":  # Howo
                if datetime.date(2026, 9, 5) <= dt <= datetime.date(2026, 10, 4):
                    active_m = 5
                elif datetime.date(2026, 10, 20) <= dt <= datetime.date(2026, 11, 10):
                    active_m = 4
            elif code == "M5":  # Lu
                if datetime.date(2026, 10, 21) <= dt <= datetime.date(2026, 11, 11):
                    active_m = 6
            elif code == "M3":  # Ủi
                if datetime.date(2026, 10, 20) <= dt <= datetime.date(2026, 11, 10):
                    active_m = 2
            elif code == "MC":  # Cẩu
                if datetime.date(2026, 9, 16) <= dt <= datetime.date(2026, 10, 28):
                    active_m = 3
            elif code == "MB":  # Bơm bê tông
                if datetime.date(2026, 9, 20) <= dt <= datetime.date(2026, 10, 31):
                    active_m = 2
            elif code == "XB":  # Xe bồn
                if datetime.date(2026, 9, 20) <= dt <= datetime.date(2026, 10, 31):
                    active_m = 6
            elif code == "M8":  # PC140
                if datetime.date(2026, 9, 20) <= dt <= datetime.date(2026, 10, 31):
                    active_m = 1
            elif code == "BP":  # Bơm ngầm 24/7
                active_m = 2
            elif code == "MP1":  # Téc dầu
                active_m = 1
            elif code == "MP2":  # Phát điện
                active_m = 2

            cell.value = active_m
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if active_m > 0:
                cell.font = FONT_MACH_VAL
                cell.fill = FILL_MACH_GREEN
            else:
                cell.font = Font(name="Arial", size=7, color="D9D9D9")
                cell.fill = PatternFill(fill_type=None)

    # 7. Bảng 3: Dầu Diezel tiêu thụ theo ngày (R42..R54)
    ws1.cell(42, 3, f"BẢNG TÍNH DẦU DIEZEL TIÊU THỤ THEO TIẾN ĐỘ THI CÔNG TUYẾN A5 (Lít/ngày)")
    ws1.cell(43, 3, "TỔNG SỐ LÍT DẦU DIEZEL TIÊU THỤ / NGÀY")
    ws1.cell(43, 3).font = FONT_RED_TITLE
    ws1.cell(43, 3).fill = FILL_YELLOW
    for col_idx in range(1, 16):
        ws1.cell(43, col_idx).fill = FILL_YELLOW

    for m in MACHINES:
        r_f = m["row_f"]
        ws1.cell(r_f, 1, m["code"]).alignment = Alignment(horizontal="center")
        ws1.cell(r_f, 3, f"Nhiên liệu dầu Diezel cho máy {m['code']}")
        ws1.cell(r_f, 4, "Lít").alignment = Alignment(horizontal="center")
        ws1.cell(r_f, 5, m["fuel_l"]).number_format = "#,##0.0"

    for i, dt in enumerate(DATES):
        col = 16 + i
        col_letter = get_column_letter(col)

        # Tổng lít dầu / ngày = SUM(P44:P54)
        c43 = ws1.cell(43, col, f"=SUM({col_letter}44:{col_letter}54)")
        c43.font = FONT_RED_85
        c43.fill = FILL_YELLOW
        c43.alignment = Alignment(horizontal="center", vertical="center")
        c43.number_format = "#,##0"
        c43.border = THIN_BORDER

        # Từng dòng máy: ={col}row_m*$Erow_f*2
        for m in MACHINES:
            r_m = m["row_m"]
            r_f = m["row_f"]
            cf = ws1.cell(r_f, col, f"={col_letter}{r_m}*$E{r_f}*2")
            cf.font = FONT_DIESEL_VAL
            cf.alignment = Alignment(horizontal="center", vertical="center")
            cf.number_format = "#,##0"
            cf.border = THIN_BORDER

    # Cập nhật Sheet 02 (Tổng ca máy toàn tuyến)
    if "02_TongHop_CaXe_CaMay_MMTB" in wb.sheetnames:
        ws2 = wb["02_TongHop_CaXe_CaMay_MMTB"]
        # Cập nhật các dòng bơm ngầm, téc dầu, phát điện theo 69 ngày * 2 ca = 138 ca
        ws2["F12"] = 138.0  # BP
        ws2["F13"] = 138.0  # MP1
        ws2["F14"] = 138.0  # MP2

    # Cập nhật Sheet 03 (Kế hoạch dầu diezel theo kỳ)
    if "03_KeHoach_Dau_Diezel" in wb.sheetnames:
        ws3 = wb["03_KeHoach_Dau_Diezel"]
        ws3["A1"] = "KẾ HOẠCH CẤP DẦU DIEZEL CHO MÁY MÓC THI CÔNG TUYẾN A5 (TỪ 05/09/2026 ĐẾN 12/11/2026)"
        ws3["F3"] = "Kỳ 1 (05/9 - 20/9)"
        ws3["G3"] = "Kỳ 2 (21/9 - 10/10)"
        ws3["H3"] = "Kỳ 3 (11/10 - 31/10)"
        ws3["I3"] = "Kỳ 4 (01/11 - 12/11)"
        ws3["D12"] = 138.0
        ws3["D13"] = 138.0
        ws3["D14"] = 138.0

    wb.save(wb_path)
    wb.close()
    print(f"[OK] Đã cập nhật thành công: {wb_path} (69 ngày: 05/09/2026 -> 12/11/2026)")


def main():
    for f in SRC_FILES:
        if os.path.exists(f):
            update_workbook(f)
        else:
            print(f"[!] Không tìm thấy tệp: {f}")


if __name__ == "__main__":
    main()

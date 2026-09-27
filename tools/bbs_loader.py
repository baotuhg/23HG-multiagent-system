# -*- coding: utf-8 -*-
"""
BBS LOADER — Đọc Bảng thống kê cốt thép (Bar Bending Schedule) THẬT từ file
Hỗ trợ: Excel (.xlsx/.xlsm), CSV (.csv, dấu phân cách , ; hoặc tab), JSON (.json)

Pure Python, Zero LLM. Không tự bịa dữ liệu: dòng nào sai được liệt kê trong
`errors` để người dùng sửa BBS, không âm thầm bỏ qua.

Nhận diện cột theo tiêu đề (không phân biệt hoa/thường, có dấu hay không dấu):
  Ký hiệu thanh / Số hiệu / mark                 → ký hiệu thanh (tùy chọn)
  Đường kính Ø (mm) / diameter_mm                → đường kính (mm)
  Mác thép / grade                                → mác thép (tùy chọn)
  Chiều dài 1 thanh (m) | (mm) / length_mm       → chiều dài (bắt buộc ghi đơn vị)
  Tổng số thanh / quantity                        → số lượng
  hoặc Số thanh / cấu kiện × Số cấu kiện          → số lượng
  Hạng mục kết cấu / component (tùy chọn)         → ghép vào ký hiệu để ghi nhãn

Ví dụ CSV tối thiểu:
  mark,diameter_mm,grade,length_mm,quantity
  F1,28,CB400-V,6578,83
"""

from __future__ import annotations
import csv
import json
import os
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Sequence, Tuple

from tools.cutting_stock_solver import CutDemand


STANDARD_REBAR_DIAMETERS = {6, 8, 10, 12, 14, 16, 18, 20, 22, 25, 28, 32, 36, 40}
MIN_PIECE_LENGTH_MM = 50          # Đoạn ngắn hơn 5cm gần như chắc chắn là lỗi dữ liệu/đơn vị
HEADER_SCAN_ROWS = 40
NON_REBAR_KEYWORDS = ("a416", "capdul", "cable", "strand", "taocap")

_FORMULA = object()  # Ô công thức Excel chưa có giá trị tính sẵn


class BBSLoadError(ValueError):
    """Không đọc được file BBS (sai định dạng, không tìm thấy cột bắt buộc...)."""


@dataclass
class BBSLoadResult:
    source: str
    demands: List[CutDemand] = field(default_factory=list)
    errors: List[str] = field(default_factory=list)    # Dòng sai dữ liệu — cần sửa BBS
    skipped: List[str] = field(default_factory=list)   # Dòng không phải thép thanh (cáp DƯL...)
    rows_read: int = 0

    @property
    def total_pieces(self) -> int:
        return sum(d.quantity for d in self.demands)


# ─────────────────────────────────────────────────────────────────────────────
# PUBLIC API
# ─────────────────────────────────────────────────────────────────────────────

def load_bbs(path: str, sheet: Optional[str] = None) -> BBSLoadResult:
    """Đọc file BBS. `sheet`: tên sheet Excel (mặc định tự tìm sheet có bảng BBS)."""
    if not os.path.exists(path):
        raise BBSLoadError(f"Không tìm thấy file BBS: {path}")

    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return _load_excel(path, sheet)
    if ext in (".csv", ".txt"):
        return _parse_rows(_read_csv(path), source=path)
    if ext == ".json":
        return _parse_rows(_read_json(path), source=path)
    raise BBSLoadError(f"Định dạng BBS chưa hỗ trợ: {ext} (dùng .xlsx, .csv hoặc .json)")


# ─────────────────────────────────────────────────────────────────────────────
# READERS
# ─────────────────────────────────────────────────────────────────────────────

def _load_excel(path: str, sheet: Optional[str]) -> BBSLoadResult:
    import openpyxl

    values_wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    formulas_wb = openpyxl.load_workbook(path, data_only=False, read_only=True)
    try:
        if sheet:
            if sheet not in values_wb.sheetnames:
                raise BBSLoadError(f"Không có sheet '{sheet}' trong {path}. Các sheet: {values_wb.sheetnames}")
            candidates = [sheet]
        else:
            candidates = list(values_wb.sheetnames)

        for name in candidates:
            rows = _merge_formula_markers(
                values_wb[name].iter_rows(values_only=True),
                formulas_wb[name].iter_rows(values_only=True),
            )
            if _find_header(rows) is not None:
                return _parse_rows(rows, source=f"{path}#{name}")
    finally:
        values_wb.close()
        formulas_wb.close()

    where = f"sheet '{sheet}'" if sheet else "bất kỳ sheet nào"
    raise BBSLoadError(
        f"Không tìm thấy bảng BBS trong {where} của {path} — cần các cột đường kính, "
        f"chiều dài (ghi đơn vị m hoặc mm) và số lượng"
    )


def _merge_formula_markers(values_rows, formula_rows) -> List[List[Any]]:
    """Ô có công thức nhưng chưa có giá trị tính sẵn → đánh dấu _FORMULA thay vì None."""
    merged = []
    for vals, forms in zip(values_rows, formula_rows):
        row = list(vals)
        for i, f in enumerate(forms):
            if i < len(row) and row[i] is None and isinstance(f, str) and f.startswith("="):
                row[i] = _FORMULA
        merged.append(row)
    return merged


def _read_csv(path: str) -> List[List[Any]]:
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        text = f.read()
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(text.splitlines(), dialect)]


def _read_json(path: str) -> List[List[Any]]:
    with open(path, "r", encoding="utf-8") as f:
        data = json.load(f)
    items = data.get("items", data.get("bbs", [])) if isinstance(data, dict) else data
    if not isinstance(items, list) or not all(isinstance(i, dict) for i in items):
        raise BBSLoadError("File JSON BBS phải là danh sách đối tượng hoặc {\"items\": [...]}")
    keys: List[str] = []
    for item in items:
        for k in item:
            if k not in keys:
                keys.append(k)
    return [keys] + [[item.get(k) for k in keys] for item in items]


# ─────────────────────────────────────────────────────────────────────────────
# PARSING
# ─────────────────────────────────────────────────────────────────────────────

def _norm(text: Any) -> str:
    s = unicodedata.normalize("NFD", str(text))
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    s = s.lower().replace("đ", "d")
    return re.sub(r"[^a-z0-9]", "", s)


def _length_unit(header: Any) -> Optional[str]:
    raw = str(header).lower().strip()
    if re.search(r"(^|[\s(\[_/-])mm([\s)\]]|$)", raw):
        return "mm"
    if re.search(r"(^|[\s(\[_/-])m([\s)\]]|$)", raw):
        return "m"
    return None


def _detect_columns(headers: Sequence[Any]) -> Dict[str, Any]:
    cols: Dict[str, Any] = {}
    for idx, raw in enumerate(headers):
        if raw is None or raw is _FORMULA:
            continue
        h = _norm(raw)
        if not h:
            continue
        if "tongchieudai" in h or "trongluong" in h or "khoiluong" in h or "weight" in h \
                or "totallength" in h:
            continue
        if "kyhieu" in h or "sohieu" in h or h in ("mark", "barmark"):
            cols.setdefault("mark", idx)
        elif "duongkinh" in h or h in ("d", "dmm", "dia", "diameter", "diametermm", "phi"):
            cols.setdefault("diameter", idx)
        elif "macthep" in h or "capthep" in h or h == "grade":
            cols.setdefault("grade", idx)
        elif "chieudai" in h or h.startswith("length"):
            cols.setdefault("length", (idx, _length_unit(raw), str(raw)))
        elif "tongsothanh" in h or h in ("quantity", "qty", "count", "soluong", "totalbars"):
            cols.setdefault("quantity", idx)
        elif "sothanh" in h:
            cols.setdefault("per_member", idx)
        elif "socaukien" in h or h in ("members", "membercount"):
            cols.setdefault("members", idx)
        elif "hangmuc" in h or h in ("component", "structure"):
            cols.setdefault("component", idx)
        elif "bophan" in h or h in ("part", "member"):
            cols.setdefault("part", idx)
    return cols


def _find_header(rows: Sequence[Sequence[Any]]) -> Optional[Tuple[int, Dict[str, Any]]]:
    for r, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        cols = _detect_columns(row)
        if "diameter" in cols and "length" in cols and ("quantity" in cols or "per_member" in cols):
            return r, cols
    return None


def _cell(row: Sequence[Any], idx: Optional[int]) -> Any:
    if idx is None or idx >= len(row):
        return None
    v = row[idx]
    if isinstance(v, str) and not v.strip():
        return None
    return v


def _to_number(v: Any) -> Optional[float]:
    if v is None or v is _FORMULA:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "")
    if "," in s and ("." not in s or s.rfind(",") > s.rfind(".")):
        s = s.replace(".", "").replace(",", ".")   # 6,578 hoặc 1.234,5 (dấu phẩy thập phân)
    else:
        s = s.replace(",", "")                     # 1,234.5 (dấu phẩy hàng nghìn)
    try:
        return float(s)
    except ValueError:
        return None


def _parse_rows(rows: List[List[Any]], source: str) -> BBSLoadResult:
    found = _find_header(rows)
    if found is None:
        raise BBSLoadError(
            f"{source}: không tìm thấy dòng tiêu đề BBS — cần các cột đường kính, "
            f"chiều dài (ghi đơn vị m hoặc mm) và số lượng"
        )
    header_idx, cols = found
    length_idx, unit, length_header = cols["length"]
    if unit is None:
        raise BBSLoadError(
            f"{source}: cột chiều dài '{length_header}' không ghi đơn vị — "
            f"đặt tên cột là 'length_mm' hoặc 'Chiều dài (m)'"
        )
    factor = 1000.0 if unit == "m" else 1.0
    result = BBSLoadResult(source=source)

    for offset, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        raw_len = _cell(row, length_idx)
        raw_qty = _cell(row, cols.get("quantity"))
        raw_per = _cell(row, cols.get("per_member"))
        if raw_len is None and raw_qty is None and raw_per is None:
            continue  # dòng trống, dòng tiêu đề nhóm, dòng tổng hợp

        result.rows_read += 1
        mark = str(_cell(row, cols.get("mark")) or f"R{offset}").strip()
        component = _cell(row, cols.get("component"))
        label = f"{str(component).strip()}:{mark}" if component else mark
        where = f"Dòng {offset} ({label})"
        grade = str(_cell(row, cols.get("grade")) or "").strip()
        part = str(_cell(row, cols.get("part")) or "")

        if any(k in _norm(grade + part + mark) for k in NON_REBAR_KEYWORDS):
            result.skipped.append(f"{where}: cáp DƯL / không phải thép thanh — không cắt từ cây 11.7m")
            continue

        diameter = _to_number(_cell(row, cols["diameter"]))
        if diameter is None:
            result.errors.append(f"{where}: thiếu hoặc sai đường kính '{_cell(row, cols['diameter'])}'")
            continue
        if diameter != int(diameter) or int(diameter) not in STANDARD_REBAR_DIAMETERS:
            result.errors.append(
                f"{where}: đường kính Ø{diameter:g} không phải thép thanh tiêu chuẩn "
                f"{sorted(STANDARD_REBAR_DIAMETERS)}"
            )
            continue

        if raw_len is _FORMULA:
            result.errors.append(f"{where}: chiều dài là công thức chưa được tính — mở và lưu lại file bằng Excel")
            continue
        length = _to_number(raw_len)
        if length is None:
            result.errors.append(f"{where}: sai chiều dài '{raw_len}'")
            continue
        length_mm = int(round(length * factor))
        if length_mm < MIN_PIECE_LENGTH_MM:
            result.errors.append(
                f"{where}: chiều dài {length_mm}mm quá ngắn — kiểm tra đơn vị/lệch cột"
            )
            continue

        quantity = _to_number(raw_qty)
        if quantity is None:
            per_member = _to_number(raw_per)
            members_raw = _cell(row, cols.get("members"))
            members = _to_number(members_raw) if members_raw is not None else 1.0
            if per_member is None or members is None:
                reason = "công thức chưa được tính" if _FORMULA in (raw_qty, raw_per, members_raw) else "trống/sai"
                result.errors.append(f"{where}: không xác định được số lượng thanh ({reason})")
                continue
            quantity = per_member * members
        if quantity < 0 or abs(quantity - round(quantity)) > 1e-6:
            result.errors.append(f"{where}: số lượng thanh {quantity:g} không hợp lệ")
            continue
        if quantity == 0:
            continue

        result.demands.append(CutDemand(
            length_mm=length_mm,
            quantity=int(round(quantity)),
            diameter_mm=int(diameter),
            mark=label,
            grade=grade,
        ))

    if not result.demands and not result.errors:
        raise BBSLoadError(f"{source}: bảng BBS không có dòng dữ liệu nào")
    return result

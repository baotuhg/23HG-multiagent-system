# -*- coding: utf-8 -*-
"""
TAKEOFF RULES — thư viện diễn giải khối lượng (bê tông, ván khuôn, đào đắp) từ kích thước hình học.

Mục tiêu: mỗi khối lượng đều có DIỄN GIẢI đọc được (công thức bằng số thật), tính bằng Decimal làm tròn
half-up, và không có "số chết" ẩn: kích thước và quy ước là đầu vào có nhãn.

Cấu kiện hỗ trợ (đủ cho nhà dân dụng, cầu, thoát nước, đường):
  - Khối chữ nhật: móng, cột, dầm, sàn, tường  (rect_concrete / rect_formwork)
  - Tròn: cọc khoan nhồi (+ đoạn đập đầu cọc), cột tròn  (pile, circular_column)
  - Đất: đào hào/hố mái taluy, đào đắp theo mặt cắt (diện tích trung bình đầu mút)
  - Cống hộp 1-n khoang có vút góc (BoxCulvert) và bảng diễn giải Excel của cống A5

QUY TẮC ĐO BÓC LÀ ĐẦU VÀO, KHÔNG PHẢI HẰNG SỐ TRONG CODE
  Văn bản đo bóc thay đổi theo thời gian (TT 17/2019 → TT 13/2021, sửa đổi bởi TT 01/2025 và TT 60/2025;
  từ 01/07/2026 theo nguồn thứ cấp có thể chuyển sang văn bản mới) nên các ngưỡng phụ thuộc điều khoản
  (trừ lỗ rỗng bao nhiêu, mặt nào tính ván khuôn...) nằm trong MeasurementProfile do kỹ sư QS chọn/nạp từ
  JSON (load_profile). Hồ sơ `verified=False` (mặc định, và cả PROFILE_TT13_2021_PL_VI lập từ bản OCR Phụ lục VI
  do người dùng cung cấp): mọi kết quả kèm cảnh báo cho tới khi đối chiếu bản gốc và đặt verified=True.
  Thông tư 12/2021/TT-BXD là định mức (hao phí), KHÔNG phải đo bóc.

Quy tắc đã mã hóa từ Phụ lục VI (kèm TT 13/2021/TT-BXD): lấy 3 số thập phân; bê tông không trừ cốt thép < 2% và lỗ
  rỗng < 0,1 m3; ván khuôn không trừ lỗ < 1 m2; xây/gỗ/hoàn thiện không trừ lỗ < 0,25 m2; mặt đường không trừ hố ga
  < 1 m2; đào đắp không cộng độ nở rời/co ngót và trừ công trình ngầm chiếm chỗ; ống thoát nước không tính đoạn hố
  ga; dàn giáo trong: cao > 3,6 m, +1 lớp mỗi 1,2 m, phần dư < 0,6 m bỏ; xuất Bảng 6.1/6.2 đúng cột mẫu.

Nguyên tắc thông dụng được dùng (không phụ thuộc điều khoản cụ thể):
  - Đo theo kích thước trong bản vẽ thiết kế; bê tông và ván khuôn tách riêng theo cấu kiện.
  - Ván khuôn = diện tích bề mặt bê tông tiếp xúc ván khuôn; mặt nào có ván khuôn là quy ước biện pháp thi công.

Mặt cắt cống hộp (đơn vị mét), tính trên 1 m dài:
  diện tích BT   = B_ngoài × H_ngoài − n × b_lòng × h_lòng + n × 4 × ½ × c²       (c: cạnh vút)
  chu vi lòng    = 2(b + h) − 4c(2 − √2)       (mỗi góc vút thay 2c bằng cạnh huyền c√2)
  ván khuôn/m    = k_trong × n × chu_vi_lòng + k_hông × 2H + k_trên × B + k_đáy × B
                   + k_đầu × 2 × diện_tích_BT / L_đốt

Pure Python, Zero LLM.
"""

from __future__ import annotations
import math
from dataclasses import dataclass, field, replace
from typing import Dict, Iterable, List, Optional, Sequence, Tuple
import json

from tools.money import round_half_up

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SQRT2 = math.sqrt(2.0)


@dataclass(frozen=True)
class BoxCulvert:
    outer_w: float      # bề rộng ngoài (m)
    outer_h: float      # chiều cao ngoài (m)
    inner_w: float      # bề rộng lòng một khoang (m)
    inner_h: float      # chiều cao lòng một khoang (m)
    cells: int = 1      # số khoang
    haunch: float = 0.0  # cạnh vút trong lòng (m), tam giác vuông cân

    def concrete_area(self) -> float:
        """Diện tích mặt cắt bê tông (m2) = thể tích trên 1 m dài (m3/m)."""
        return (self.outer_w * self.outer_h - self.cells * self.inner_w * self.inner_h
                + self.cells * 4 * 0.5 * self.haunch ** 2)

    def inner_perimeter(self) -> float:
        """Chu vi lòng một khoang (m) có trừ phần góc vút."""
        return 2 * (self.inner_w + self.inner_h) - 4 * self.haunch * (2 - SQRT2)

    def formwork_per_m(self, segment_length: float, inner: bool = True, sides: bool = True,
                       top: bool = False, bottom: bool = False, ends: bool = True) -> Dict[str, float]:
        """Ván khuôn tiếp xúc (m2) trên 1 m dài theo từng mặt và tổng."""
        parts = {
            "trong": self.cells * self.inner_perimeter() if inner else 0.0,
            "hong_ngoai": 2 * self.outer_h if sides else 0.0,
            "mat_tren": self.outer_w if top else 0.0,
            "mat_day": self.outer_w if bottom else 0.0,
            "dau_dot": 2 * self.concrete_area() / segment_length if ends else 0.0,
        }
        parts["tong"] = sum(parts.values())
        return parts


# ─────────────────────────────────────────────────────────────────────────────
# HỒ SƠ QUY TẮC ĐO BÓC + KẾT QUẢ CÓ DIỄN GIẢI
# ─────────────────────────────────────────────────────────────────────────────

# Khóa ngưỡng "không phải trừ lỗ rỗng nhỏ hơn X" — theo từng loại công tác
K_BETONG_M3 = "be_tong_m3"          # lỗ rỗng trong khối bê tông (m3)
K_VANKHUON_M2 = "van_khuon_m2"      # lỗ rỗng trên bề mặt bê tông, tính ván khuôn (m2)
K_XAY_M2 = "xay_m2"                 # khoảng trống trong khối xây (m2)
K_GO_M2 = "go_m2"                   # lỗ rỗng sàn, vách, trần gỗ (m2)
K_HOANTHIEN_M2 = "hoan_thien_m2"    # lỗ rỗng không phải hoàn thiện (m2)
K_MATDUONG_M2 = "mat_duong_m2"      # lỗ trống trên mặt đường: hố ga, hố thăm (m2)
K_DAT_KHAC_CAP_M3 = "dat_khac_cap_m3"   # đất/đá khác cấp trong hố đào không tách riêng (m3)


@dataclass(frozen=True)
class MeasurementProfile:
    """Quy tắc đo bóc phụ thuộc điều khoản văn bản. Giá trị do người dùng đặt từ văn bản, kèm nguồn."""
    name: str = "CHƯA XÁC ĐỊNH"
    source: str = ""                                   # văn bản + mục/trang đã đối chiếu
    verified: bool = False                             # True chỉ khi đã đối chiếu bản gốc
    decimals: int = 3                                  # khối lượng lấy đến 3 số sau dấu phẩy khi là số thập phân
    # "không phải trừ" lỗ rỗng NHỎ HƠN ngưỡng (lỗ >= ngưỡng thì trừ). Thiếu khóa = trừ mọi lỗ ghi trong bản vẽ.
    no_deduct_below: Dict[str, float] = field(default_factory=dict)
    rebar_no_deduct_below_ratio: Optional[float] = None   # không trừ thể tích cốt thép nếu hàm lượng < tỷ lệ này
    deduct_formwork_overlap: bool = True               # trừ phần ván khuôn chỗ cấu kiện giao nhau (dầm–sàn...)
    ocr_pending: Tuple[str, ...] = ()                  # số liệu lấy từ OCR cần đối chiếu bản gốc

    def threshold(self, key: str) -> Optional[float]:
        return self.no_deduct_below.get(key)

    def warnings(self) -> List[str]:
        if self.verified:
            return []
        msg = (f"Hồ sơ quy tắc '{self.name}' CHƯA đối chiếu bản gốc văn bản đo bóc hiện hành"
               + (f" (nguồn: {self.source})" if self.source else ""))
        out = [msg]
        out += [f"Cần đối chiếu: {p}" for p in self.ocr_pending]
        return out


DEFAULT_PROFILE = MeasurementProfile()

# Hồ sơ lập từ bản OCR Phụ lục VI "Phương pháp đo bóc khối lượng công trình" (kèm TT 13/2021/TT-BXD) do người dùng
# cung cấp. Các số bên dưới được hai bộ OCR đọc giống nhau nhưng nằm trong danh sách "cần kiểm tra" nên verified=False
# cho tới khi đối chiếu bản gốc. Văn bản này có thể đã hết hiệu lực/được thay thế tuỳ thời điểm lập hồ sơ.
PROFILE_TT13_2021_PL_VI = MeasurementProfile(
    name="TT 13/2021/TT-BXD — Phụ lục VI",
    source="Phụ lục VI TT 13/2021/TT-BXD: II.1.4d (làm tròn 3 số, Trang 2), II.5.2 (Trang 4), II.5.3 (Trang 4-5), "
           "II.5.4 (Trang 5), II.5.5 (Trang 5), II.5.9 (Trang 6), II.5.12-5.13 (Trang 7)",
    verified=False,
    decimals=3,
    no_deduct_below={
        K_BETONG_M3: 0.1,          # II.5.4: lỗ rỗng trong bê tông < 0,1 m3
        K_VANKHUON_M2: 1.0,        # II.5.5: lỗ rỗng trên bề mặt bê tông < 1 m2
        K_XAY_M2: 0.25,            # II.5.3: khoảng trống trong khối xây < 0,25 m2
        K_GO_M2: 0.25,             # II.5.12: sàn, vách, trần gỗ < 0,25 m2
        K_HOANTHIEN_M2: 0.25,      # II.5.13a: lỗ rỗng không phải hoàn thiện < 0,25 m2
        K_MATDUONG_M2: 1.0,        # II.5.9: hố ga, hố thăm... trên mặt đường < 1 m2
        K_DAT_KHAC_CAP_M3: 1.0,    # II.5.2: đất/đá khác cấp < 1 m3 không tách riêng
    },
    rebar_no_deduct_below_ratio=0.02,   # II.5.4: cốt thép hàm lượng < 2% thể tích cấu kiện bê tông
    ocr_pending=(
        "Trang 5, II.5.4: ngưỡng lỗ rỗng bê tông 0,1 m3 (OCR độ tin cậy thấp)",
        "Trang 5, II.5.5: ngưỡng lỗ rỗng ván khuôn 1 m2 (OCR độ tin cậy thấp)",
        "Trang 4, II.5.2: ngưỡng đất/đá khác cấp 1 m3 (OCR độ tin cậy thấp)",
        "Trang 6, II.5.9: ngưỡng lỗ trống mặt đường 1 m2 (OCR độ tin cậy thấp)",
        "Trang 7, II.5.12-5.13: ngưỡng 0,25 m2 (OCR độ tin cậy thấp)",
        "Trang 5, II.5.5: câu về ván khuôn tấm định hình > 3 m2 bị OCR nhiễu — chưa áp dụng",
    ),
)


def load_profile(path: str) -> MeasurementProfile:
    """Nạp hồ sơ quy tắc từ JSON (khóa trùng tên trường của MeasurementProfile)."""
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    unknown = set(data) - set(MeasurementProfile.__dataclass_fields__)
    if unknown:
        raise ValueError(f"Hồ sơ quy tắc có khóa lạ: {', '.join(sorted(unknown))}")
    if "ocr_pending" in data:
        data["ocr_pending"] = tuple(data["ocr_pending"])
    return MeasurementProfile(**data)


@dataclass(frozen=True)
class Quantity:
    """Một khối lượng đo bóc kèm diễn giải đọc được (đủ cột cho Bảng chi tiết khối lượng mẫu 6.2)."""
    name: str
    unit: str
    value: float                 # khối lượng toàn bộ
    formula: str                 # công thức bằng số thật, vd "2 × (1.2 × 1 × 0.4)"
    notes: Tuple[str, ...] = ()
    count: float = 1             # số bộ phận giống nhau (cột 6)
    per_unit: Optional[float] = None   # khối lượng một bộ phận (cột 8)
    drawing: str = ""            # ký hiệu bản vẽ (cột 2)
    code: str = ""               # mã hiệu công tác (cột 3)
    kind: str = ""               # loại công tác (be_tong, van_khuon, dao_hao...) — để phân nhóm khi tổng hợp

    def explain(self) -> str:
        extra = ("  [" + "; ".join(self.notes) + "]") if self.notes else ""
        return f"{self.name}: {self.formula} = {self.value:,.3f} {self.unit}{extra}"

    def with_ref(self, drawing: str = "", code: str = "", kind: str = "") -> "Quantity":
        return replace(self, drawing=drawing or self.drawing, code=code or self.code, kind=kind or self.kind)


def _g(x: float) -> str:
    return f"{x:g}"


def _q(profile: MeasurementProfile, name: str, unit: str, raw: float, formula: str,
       notes: Sequence[str] = (), count: float = 1) -> Quantity:
    notes = tuple(notes) + tuple(profile.warnings())
    total = float(round_half_up(raw, profile.decimals))
    per_unit = float(round_half_up(raw / count, profile.decimals)) if count else None
    return Quantity(name, unit, total, formula, notes, count=count, per_unit=per_unit)


def _deduct(gross: float, openings: Sequence[float], threshold: Optional[float]) -> Tuple[float, List[float], List[float]]:
    """Trừ các lỗ rỗng; lỗ NHỎ HƠN ngưỡng thì không trừ (ngưỡng None = trừ hết). Trả về (net, đã trừ, bỏ qua)."""
    taken = [o for o in openings if threshold is None or o >= threshold]
    skipped = [o for o in openings if threshold is not None and o < threshold]
    return gross - sum(taken), taken, skipped


# ── Khối chữ nhật ────────────────────────────────────────────────────────────

def rect_concrete(name: str, n: float, length: float, width: float, height: float,
                  openings_m3: Sequence[float] = (), rebar_volume_m3: float = 0.0,
                  profile: MeasurementProfile = DEFAULT_PROFILE) -> Quantity:
    """Bê tông cấu kiện chữ nhật: n × (dài × rộng × cao), trừ lỗ rỗng và cốt thép theo hồ sơ quy tắc.

    - Lỗ rỗng nhỏ hơn ngưỡng của hồ sơ không trừ.
    - Cốt thép: chỉ trừ thể tích cốt thép khi hàm lượng (rebar / thể tích BT một cấu kiện) ≥ ngưỡng của hồ sơ;
      hồ sơ không đặt ngưỡng → không trừ cốt thép (thông lệ). rebar_volume_m3 là của TOÀN BỘ n cấu kiện.
    """
    unit_v = length * width * height
    gross = n * unit_v
    net, taken, skipped = _deduct(gross, openings_m3, profile.threshold(K_BETONG_M3))
    f = f"{_g(n)} × ({_g(length)} × {_g(width)} × {_g(height)})" + "".join(f" − {_g(o)}" for o in taken)
    notes = []
    if skipped:
        notes.append(f"không trừ {len(skipped)} lỗ < {_g(profile.threshold(K_BETONG_M3))} m3 theo hồ sơ quy tắc")
    if rebar_volume_m3:
        ratio = rebar_volume_m3 / gross if gross else 0.0
        limit = profile.rebar_no_deduct_below_ratio
        if limit is not None and ratio >= limit:
            net -= rebar_volume_m3
            f += f" − {_g(rebar_volume_m3)}"
            notes.append(f"trừ cốt thép: hàm lượng {ratio:.2%} ≥ {limit:.0%}")
        else:
            notes.append(f"không trừ cốt thép (hàm lượng {ratio:.2%})")
    return _q(profile, name, "m3", net, f, notes, count=n)


_FORMWORK_KINDS = ("mong", "cot", "dam", "san", "tuong")


def rect_formwork(name: str, kind: str, n: float, length: float, width: float, height: float,
                  slab_thickness: float = 0.0, edges: bool = False, ends: bool = False,
                  openings_m2: Sequence[float] = (), profile: MeasurementProfile = DEFAULT_PROFILE) -> Quantity:
    """Ván khuôn tiếp xúc cấu kiện chữ nhật (m2) — bề mặt bê tông cần chống đỡ tạm thời khi đúc.

    mong : 4 mặt hông           = 2(dài + rộng) × cao
    cot  : 4 mặt (dài × rộng là tiết diện, cao là chiều cao cột) = 2(dài + rộng) × cao
    dam  : 2 hông + đáy         = dài × (2(cao − dày sàn) + rộng)   [trừ phần chìm trong sàn nếu có]
    san  : đáy (+ cạnh sàn)     = dài × rộng (+ 2(dài + rộng) × cao khi edges)   [cao = chiều dày sàn]
    tuong: 2 mặt (+ 2 đầu)      = 2 × dài × cao (+ 2 × rộng × cao khi ends)   [rộng = chiều dày tường]
    """
    if kind not in _FORMWORK_KINDS:
        raise ValueError(f"kind phải thuộc {_FORMWORK_KINDS}")
    L, W, H = length, width, height
    if kind in ("mong", "cot"):
        one, f1 = 2 * (L + W) * H, f"2 × ({_g(L)} + {_g(W)}) × {_g(H)}"
    elif kind == "dam":
        depth = H - (slab_thickness if profile.deduct_formwork_overlap else 0.0)
        one = L * (2 * depth + W)
        cut = f" − 2 × {_g(slab_thickness)}" if slab_thickness and profile.deduct_formwork_overlap else ""
        f1 = f"{_g(L)} × (2 × {_g(H)}{cut} + {_g(W)})"
    elif kind == "san":
        one, f1 = L * W, f"{_g(L)} × {_g(W)}"
        if edges:
            one += 2 * (L + W) * H
            f1 += f" + 2 × ({_g(L)} + {_g(W)}) × {_g(H)}"
    else:  # tuong
        one, f1 = 2 * L * H, f"2 × {_g(L)} × {_g(H)}"
        if ends:
            one += 2 * W * H
            f1 += f" + 2 × {_g(W)} × {_g(H)}"
    gross = n * one
    net, taken, skipped = _deduct(gross, openings_m2, profile.threshold(K_VANKHUON_M2))
    f = f"{_g(n)} × ({f1})" + "".join(f" − {_g(o)}" for o in taken)
    notes = [f"không trừ {len(skipped)} lỗ < {_g(profile.threshold(K_VANKHUON_M2))} m2 theo hồ sơ quy tắc"] if skipped else []
    return _q(profile, name, "m2", net, f, notes, count=n)


# ── Cấu kiện tròn, cọc, khoan ────────────────────────────────────────────────

def pile(name: str, n: float, diameter: float, length: float, cutoff: float = 0.0,
         profile: MeasurementProfile = DEFAULT_PROFILE) -> Tuple[Quantity, Optional[Quantity]]:
    """Cọc khoan nhồi: bê tông π/4·D²·L·n; kèm khối lượng đập đầu cọc (π/4·D²·cutoff·n) nếu có."""
    area = math.pi * diameter ** 2 / 4
    conc = _q(profile, name, "m3", n * area * length,
              f"{_g(n)} × (π/4 × {_g(diameter)}² × {_g(length)})", count=n)
    demo = None
    if cutoff:
        demo = _q(profile, name + " — đập đầu cọc", "m3", n * area * cutoff,
                  f"{_g(n)} × (π/4 × {_g(diameter)}² × {_g(cutoff)})", count=n)
    return conc, demo


def bored_length(name: str, n: float, drill_depth: float, profile: MeasurementProfile = DEFAULT_PROFILE) -> Quantity:
    """Công tác khoan (m): chiều sâu khoan đo dọc lỗ khoan, từ điểm bắt đầu tiếp xúc mặt đất đến đáy hố khoan.
    Ghi rõ đường kính, cấp đất đá, khoan trên cạn/dưới nước... ở phần mô tả công tác (mục II.5.8)."""
    return _q(profile, name, "m", n * drill_depth, f"{_g(n)} × {_g(drill_depth)}", count=n)


def circular_column(name: str, n: float, diameter: float, height: float,
                    profile: MeasurementProfile = DEFAULT_PROFILE) -> Tuple[Quantity, Quantity]:
    """Cột tròn: bê tông π/4·D²·H·n và ván khuôn mặt bên π·D·H·n."""
    conc = _q(profile, name, "m3", n * math.pi * diameter ** 2 / 4 * height,
              f"{_g(n)} × (π/4 × {_g(diameter)}² × {_g(height)})", count=n)
    form = _q(profile, name + " — ván khuôn", "m2", n * math.pi * diameter * height,
              f"{_g(n)} × (π × {_g(diameter)} × {_g(height)})", count=n)
    return conc, form


# ── Đào đắp ──────────────────────────────────────────────────────────────────

def trench_excavation(name: str, length: float, bottom_width: float, depth: float, slope_m: float = 0.0,
                      profile: MeasurementProfile = DEFAULT_PROFILE) -> Quantity:
    """Đào hào mái taluy 1:m (m = ngang/đứng): V = L × (b + m·H) × H. Không cộng thêm độ nở rời/co ngót."""
    area = (bottom_width + slope_m * depth) * depth
    return _q(profile, name, "m3", length * area,
              f"{_g(length)} × ({_g(bottom_width)} + {_g(slope_m)} × {_g(depth)}) × {_g(depth)}")


def pit_excavation(name: str, bottom_a: float, bottom_b: float, depth: float, slope_m: float = 0.0,
                   profile: MeasurementProfile = DEFAULT_PROFILE) -> Quantity:
    """Đào hố móng đáy a×b mái taluy 1:m: hình chóp cụt V = H/3 × (A1 + A2 + √(A1·A2))."""
    a1 = bottom_a * bottom_b
    a2 = (bottom_a + 2 * slope_m * depth) * (bottom_b + 2 * slope_m * depth)
    return _q(profile, name, "m3", depth / 3 * (a1 + a2 + math.sqrt(a1 * a2)),
              f"{_g(depth)}/3 × ({_g(a1)} + {_g(a2)} + √({_g(a1)} × {_g(a2)}))")


def average_end_volume(name: str, stations: Sequence[float], areas: Sequence[float],
                       profile: MeasurementProfile = DEFAULT_PROFILE) -> Quantity:
    """Khối lượng theo mặt cắt (diện tích trung bình đầu mút): V = Σ (A_i + A_{i+1})/2 × (x_{i+1} − x_i)."""
    if len(stations) != len(areas) or len(stations) < 2:
        raise ValueError("Cần ≥ 2 mặt cắt, số lý trình bằng số diện tích")
    if any(b <= a for a, b in zip(stations, stations[1:])):
        raise ValueError("Lý trình phải tăng dần")
    terms, total = [], 0.0
    for (x1, a1), (x2, a2) in zip(zip(stations, areas), zip(stations[1:], areas[1:])):
        total += (a1 + a2) / 2 * (x2 - x1)
        terms.append(f"({_g(a1)} + {_g(a2)})/2 × {_g(x2 - x1)}")
    return _q(profile, name, "m3", total, " + ".join(terms))


def net_of_buried_works(name: str, gross_volume: float, buried: Sequence[Tuple[str, float]],
                        profile: MeasurementProfile = DEFAULT_PROFILE) -> Quantity:
    """Khối lượng đào/đắp trừ các công trình ngầm chiếm chỗ (đường ống kỹ thuật, cống thoát nước...) — II.5.2.
    buried = [(tên, thể tích chiếm chỗ m3)]. Thể tích chiếm chỗ là phần bao ngoài của công trình ngầm."""
    f = _g(gross_volume) + "".join(f" − {_g(v)} ({n})" for n, v in buried)
    return _q(profile, name, "m3", gross_volume - sum(v for _, v in buried), f)


# ── Đường ống, dàn giáo ──────────────────────────────────────────────────────

def pipe_length(name: str, centerline: float, occupied_by_chambers: float = 0.0, drainage: bool = True,
                profile: MeasurementProfile = DEFAULT_PROFILE) -> Quantity:
    """Đường ống: chiều dài đo dọc tim ống. Ống thoát nước không tính đoạn ở hố ga, hố thu, hố thăm chiếm chỗ (II.5.10)."""
    cut = occupied_by_chambers if drainage else 0.0
    f = _g(centerline) + (f" − {_g(cut)}" if cut else "")
    return _q(profile, name, "m", centerline - cut, f)


def scaffold_extra_layers(height: float, base: float = 3.6, step: float = 1.2, min_part: float = 0.6) -> int:
    """Số lớp dàn giáo trong tính thêm: chỉ khi cao > 3,6 m; mỗi 1,2 m tăng thêm = 1 lớp; phần dư < 0,6 m không tính (II.5.16)."""
    if height <= base:
        return 0
    extra = height - base
    layers = int(extra // step)
    if extra - layers * step >= min_part - 1e-12:
        layers += 1
    return layers


def scaffold_inner(name: str, plan_area: float, height: float, profile: MeasurementProfile = DEFAULT_PROFILE) -> Quantity:
    """Dàn giáo trong (m2 hình chiếu bằng × số lớp): chỉ tính khi cao > 3,6 m; lớp gốc 3,6 m + lớp cộng dồn."""
    if height <= 3.6:
        return _q(profile, name, "m2", 0.0, f"cao {_g(height)} ≤ 3,6 m: không tính dàn giáo trong")
    layers = 1 + scaffold_extra_layers(height)
    return _q(profile, name, "m2", plan_area * layers, f"{_g(plan_area)} × (1 + {layers - 1} lớp thêm)")


def scaffold_column(name: str, perimeter: float, height: float, profile: MeasurementProfile = DEFAULT_PROFILE) -> Quantity:
    """Dàn giáo hoàn thiện trụ, cột độc lập: (chu vi + 3,6 m) × chiều cao (II.5.16)."""
    return _q(profile, name, "m2", (perimeter + 3.6) * height, f"({_g(perimeter)} + 3.6) × {_g(height)}")


# ── Xuất theo biểu mẫu 6.1 / 6.2 ─────────────────────────────────────────────

BANG_6_2_HEADER = ["STT", "KÝ HIỆU BẢN VẼ", "MÃ HIỆU CÔNG TÁC", "DANH MỤC CÔNG TÁC", "ĐƠN VỊ TÍNH",
                   "SỐ BỘ PHẬN GIỐNG NHAU", "DIỄN GIẢI TÍNH TOÁN", "KHỐI LƯỢNG MỘT BỘ PHẬN",
                   "KHỐI LƯỢNG TOÀN BỘ", "GHI CHÚ"]
BANG_6_1_HEADER = ["STT", "MÃ HIỆU CÔNG TÁC", "DANH MỤC CÔNG TÁC XÂY DỰNG", "ĐƠN VỊ TÍNH",
                   "CÁCH THỨC XÁC ĐỊNH", "KHỐI LƯỢNG", "GHI CHÚ"]


def bang_6_2(quantities: Sequence[Quantity]) -> List[List]:
    """Các dòng của Bảng chi tiết khối lượng công tác xây dựng (mẫu 6.2, Phụ lục VI TT 13/2021): 10 cột."""
    rows = [list(BANG_6_2_HEADER)]
    for i, q in enumerate(quantities, 1):
        rows.append([i, q.drawing, q.code, q.name, q.unit, q.count, q.formula, q.per_unit, q.value,
                     "; ".join(q.notes)])
    return rows


def bang_6_1(quantities: Sequence[Quantity], how: str = "Theo Bảng chi tiết khối lượng công tác xây dựng") -> List[List]:
    """Các dòng của Bảng tổng hợp khối lượng xây dựng (mẫu 6.1): gộp cùng (mã hiệu, tên, đơn vị)."""
    merged: Dict[Tuple[str, str, str], float] = {}
    for q in quantities:
        key = (q.code, q.name, q.unit)
        merged[key] = merged.get(key, 0.0) + q.value
    rows = [list(BANG_6_1_HEADER)]
    for i, ((code, name, unit), value) in enumerate(merged.items(), 1):
        rows.append([i, code, name, unit, how, round(value, 3), ""])
    return rows


# ─────────────────────────────────────────────────────────────────────────────
# ÁP VÀO SHEET DIỄN GIẢI (QS_DIEN_GIAI_CHI_TIET / QS_TAKEOFF)
# ─────────────────────────────────────────────────────────────────────────────

# Dòng bê tông thân → dòng ván khuôn tương ứng, và đầu vào của từng loại cống (K,L,M,N)
CULVERT_ROWS = {
    8: {"formwork_row": 11, "inputs": (2.0, 2.0, 1, 0.20)},
    9: {"formwork_row": 12, "inputs": (3.0, 3.0, 1, 0.25)},
    10: {"formwork_row": 13, "inputs": (3.0, 3.0, 2, 0.25)},
}
FLAG_FIRST_ROW = 17      # C17:C21 = công tắc mặt ván khuôn
FLAGS = [
    ("Ván khuôn mặt trong (lòng cống)", 1),
    ("Ván khuôn mặt hông ngoài (2 bên)", 1),
    ("Ván khuôn mặt trên (bản nắp)", 0),
    ("Ván khuôn mặt đáy (bản đáy)", 0),
    ("Ván khuôn đầu đốt (2 đầu mỗi đốt)", 1),
]
NOTE = ("Quy ước: ván khuôn tính theo diện tích bề mặt bê tông tiếp xúc ván khuôn; mặt nào có ván khuôn do biện pháp "
        "thi công quyết định (công tắc 1/0 bên dưới). Căn cứ đo bóc: văn bản hiện hành tại thời điểm lập hồ sơ do kỹ sư QS "
        "xác nhận (theo nguồn thứ cấp: TT 13/2021/TT-BXD và sửa đổi áp dụng đến 30/06/2026; từ 01/07/2026 cần kiểm tra "
        "văn bản thay thế như TT 37/2026/TT-BXD, QĐ 1041/QĐ-BXD - chưa đối chiếu toàn văn). TT 12/2021/TT-BXD là định "
        "mức, không phải đo bóc.")


def apply_culvert_derivation(ws) -> None:
    """Chuyển bảng diễn giải cống hộp A5 sang đầu vào có tên (sửa tại chỗ, giữ nguyên vị trí cột I).

    - Dòng bê tông thân (8-10): thêm K (rộng lòng), L (cao lòng), M (số khoang), N (cạnh vút);
      cột I trở thành công thức của các ô đó (giá trị không đổi so với công thức cũ).
    - Dòng ván khuôn (11-13): GIỮ NGUYÊN số đang dùng ở cột F (không tự đổi khối lượng hợp đồng);
      thêm P..S tính lại theo hình học, T = % chênh lệch, U = kết luận để kỹ sư QS quyết định.
    - Khối công tắc quy ước ván khuôn ở A15:C21 và ghi chú căn cứ.
    """
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    inp_fill = PatternFill("solid", fgColor="FFF2CC")   # ô đầu vào
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = {
        "K": "RỘNG LÒNG b (m)", "L": "CAO LÒNG h (m)", "M": "SỐ KHOANG n", "N": "CẠNH VÚT c (m)",
        "P": "VK TRONG (m2/m)", "Q": "VK HÔNG NGOÀI (m2/m)", "R": "VK ĐẦU ĐỐT (m2/m)",
        "S": "VK TÍNH LẠI (m2/m)", "T": "CHÊNH LỆCH SO VỚI SỐ ĐANG DÙNG (cột F)", "U": "KẾT LUẬN",
    }
    for col, text in headers.items():
        c = ws[f"{col}4"]
        c.value, c.font, c.fill, c.alignment, c.border = text, hdr_font, hdr_fill, center, border
        ws.column_dimensions[col].width = 18

    f_in, f_side, f_top, f_bot, f_end = (f"$C${FLAG_FIRST_ROW + k}" for k in range(5))
    for crow, spec in CULVERT_ROWS.items():
        for col, val in zip("KLMN", spec["inputs"]):
            c = ws[f"{col}{crow}"]
            c.value, c.fill, c.border = val, inp_fill, border
            c.number_format = "0" if col == "M" else "0.00"
        area = f"(F{crow}*G{crow}-M{crow}*K{crow}*L{crow}+M{crow}*4*0.5*N{crow}^2)"
        ws[f"I{crow}"].value = f"=D{crow}*E{crow}*{area}"

        r = spec["formwork_row"]
        ws[f"P{r}"].value = f"={f_in}*M{crow}*(2*(K{crow}+L{crow})-4*N{crow}*(2-SQRT(2)))"
        ws[f"Q{r}"].value = f"={f_side}*2*G{crow}+{f_top}*F{crow}+{f_bot}*F{crow}"
        ws[f"R{r}"].value = f"={f_end}*2*{area}/E{crow}"
        ws[f"S{r}"].value = f"=P{r}+Q{r}+R{r}"
        ws[f"T{r}"].value = f"=S{r}/F{r}-1"
        ws[f"U{r}"].value = f'=IF(ABS(T{r})>0.05,"LỆCH >5% - QS xác nhận","Khớp (<=5%)")'
        for col in "PQRSTU":
            c = ws[f"{col}{r}"]
            c.border = border
            c.number_format = "0.0%" if col == "T" else "0.000"
        ws[f"J{r}"].value = "Cột F = VK m2/m ĐANG DÙNG (không phải bề rộng); P..U là tính lại theo hình học"

    ws["A15"].value = "QUY ƯỚC VÁN KHUÔN (1 = có ván khuôn, 0 = không) - đổi ô vàng, cột P..U của dòng 11-13 tự tính lại"
    ws["A15"].font = Font(bold=True)
    for k, (label, default) in enumerate(FLAGS):
        r = FLAG_FIRST_ROW + k
        ws[f"B{r}"].value = label
        c = ws[f"C{r}"]
        c.value, c.fill, c.border, c.alignment = default, inp_fill, border, center
    note_row = FLAG_FIRST_ROW + len(FLAGS) + 1
    ws[f"A{note_row}"].value = NOTE
    ws[f"A{note_row}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=10)
    ws.row_dimensions[note_row].height = 62

# -*- coding: utf-8 -*-
"""
EXCEL EVAL — Tính giá trị công thức Excel chưa có kết quả lưu sẵn (Pure Python)

Khi nào cần: file .xlsx do phần mềm khác tạo (openpyxl, xuất từ web...) chỉ có công thức,
chưa từng được Excel tính → đọc bằng openpyxl(data_only=True) ra None. File đã mở và lưu
bằng Excel thì dùng luôn giá trị Excel đã tính (không tính lại).

Hỗ trợ (đủ cho bảng QS / BBS / tiến độ thông dụng):
  số, chuỗi, + - * / ^ &, ngoặc, dấu âm, tham chiếu ô / vùng (có $, sang sheet khác),
  SUM, SUMIF, SUMIFS, COUNTIF, COUNTIFS, ROUND, ROUNDUP, ROUNDDOWN, MIN, MAX,
  AVERAGE, ABS, SQRT, PI, PRODUCT.
Hàm khác → FormulaError (báo rõ, không đoán).
"""

from __future__ import annotations
import fnmatch
import math
import re
from typing import Any, Dict, List, Optional, Tuple

from openpyxl.utils import column_index_from_string


class FormulaError(ValueError):
    """Công thức không tính được (hàm chưa hỗ trợ, tham chiếu vòng, lỗi cú pháp...)."""


_TOKEN = re.compile(r"""
    (?P<ws>\s+)
  | (?P<str>"(?:[^"]|"")*")
  | (?P<ref>(?:(?:'(?:[^']|'')+'|[A-Za-z_][\w\.]*)!)?\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?)
  | (?P<num>\d+(?:\.\d*)?(?:[eE][+-]?\d+)?|\.\d+)
  | (?P<func>[A-Za-z_][A-Za-z0-9_\.]*)\s*\(
  | (?P<op><>|<=|>=|[-+*/^&(),:<>=%])
""", re.VERBOSE)


class WorkbookEvaluator:
    """Đọc workbook 2 lần (công thức + giá trị lưu sẵn) và tính ô thiếu giá trị theo yêu cầu."""

    def __init__(self, path: str):
        import openpyxl
        self._formulas: Dict[str, Dict[Tuple[int, int], Any]] = {}
        self._cached: Dict[str, Dict[Tuple[int, int], Any]] = {}
        fwb = openpyxl.load_workbook(path, data_only=False, read_only=True)
        vwb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            for ws in fwb.worksheets:
                self._formulas[ws.title] = _cells(ws)
                self._cached[ws.title] = _cells(vwb[ws.title])
        finally:
            fwb.close()
            vwb.close()
        self._computed: Dict[Tuple[str, int, int], Any] = {}
        self._stack: set = set()
        self.evaluated_cells = 0   # số ô đã phải tự tính (không có giá trị Excel lưu sẵn)

    @property
    def sheetnames(self) -> List[str]:
        return list(self._formulas)

    def max_row(self, sheet: str) -> int:
        return max((r for r, _ in self._formulas[sheet]), default=0)

    def max_col(self, sheet: str) -> int:
        return max((c for _, c in self._formulas[sheet]), default=0)

    def formula(self, sheet: str, row: int, col: int) -> Any:
        return self._formulas[sheet].get((row, col))

    def value(self, sheet: str, row: int, col: int) -> Any:
        """Giá trị ô: giá trị Excel lưu sẵn, hoặc tự tính công thức. Lỗi → FormulaError."""
        if sheet not in self._formulas:
            raise FormulaError(f"Không có sheet '{sheet}'")
        cached = self._cached[sheet].get((row, col))
        raw = self._formulas[sheet].get((row, col))
        if cached is not None or not (isinstance(raw, str) and raw.startswith("=")):
            return cached if cached is not None else raw
        key = (sheet, row, col)
        if key in self._computed:
            return self._computed[key]
        if key in self._stack:
            raise FormulaError(f"Tham chiếu vòng tại {sheet}!{_addr(row, col)}")
        self._stack.add(key)
        try:
            result = _Parser(self, sheet, raw[1:]).parse()
        except FormulaError as e:
            raise FormulaError(f"{sheet}!{_addr(row, col)} '{raw}': {e}") from None
        finally:
            self._stack.discard(key)
        if isinstance(result, list):
            result = result[0][0] if result and result[0] else None
        self._computed[key] = result
        self.evaluated_cells += 1
        return result

    def rows(self, sheet: str, missing=None) -> List[List[Any]]:
        """Toàn bộ giá trị sheet theo dòng; ô công thức không tính được → `missing`."""
        out = []
        ncol = self.max_col(sheet)
        for r in range(1, self.max_row(sheet) + 1):
            row = []
            for c in range(1, ncol + 1):
                try:
                    row.append(self.value(sheet, r, c))
                except FormulaError:
                    row.append(missing)
            out.append(row)
        return out

    def range_values(self, sheet: str, r1: int, c1: int, r2: int, c2: int) -> List[List[Any]]:
        return [[self.value(sheet, r, c) for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)]


def _cells(ws) -> Dict[Tuple[int, int], Any]:
    out = {}
    for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
        for c, v in enumerate(row, start=1):
            if v is not None:
                out[(r, c)] = v
    return out


def _addr(row: int, col: int) -> str:
    from openpyxl.utils import get_column_letter
    return f"{get_column_letter(col)}{row}"


def _num(v: Any) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        raise FormulaError(f"giá trị '{v}' không phải số")


def _flatten(v: Any) -> List[Any]:
    if isinstance(v, list):
        return [x for row in v for x in row]
    return [v]


def _match(value: Any, criterion: Any) -> bool:
    """Điều kiện kiểu Excel: số, chuỗi có * ?, hoặc '>5', '<>0', '=abc'."""
    if isinstance(criterion, (int, float)) and not isinstance(criterion, bool):
        try:
            return value is not None and value != "" and _num(value) == float(criterion)
        except FormulaError:
            return False
    text = "" if criterion is None else str(criterion)
    m = re.match(r"^(<>|<=|>=|<|>|=)?(.*)$", text, re.S)
    op, operand = m.group(1) or "=", m.group(2)
    try:
        number = float(operand.replace(",", "."))
        is_number = operand.strip() != ""
    except ValueError:
        number, is_number = None, False
    if is_number and op in ("<", ">", "<=", ">=", "=", "<>"):
        try:
            v = _num(value) if value not in (None, "") else None
        except FormulaError:
            v = None
        if v is None:
            return op == "<>"
        return {"<": v < number, ">": v > number, "<=": v <= number, ">=": v >= number,
                "=": v == number, "<>": v != number}[op]
    s = "" if value is None else str(value)
    hit = fnmatch.fnmatchcase(s.lower(), operand.lower()) if any(ch in operand for ch in "*?") \
        else s.lower() == operand.lower()
    return hit if op == "=" else (not hit if op == "<>" else False)


class _Parser:
    def __init__(self, ev: WorkbookEvaluator, sheet: str, text: str):
        self.ev, self.sheet = ev, sheet
        self.tokens = []
        pos = 0
        while pos < len(text):
            m = _TOKEN.match(text, pos)
            if not m:
                raise FormulaError(f"không đọc được '{text[pos:pos + 15]}'")
            pos = m.end()
            kind = m.lastgroup
            if kind != "ws":
                self.tokens.append((kind, m.group(kind)))
        self.i = 0

    def peek(self):
        return self.tokens[self.i] if self.i < len(self.tokens) else (None, None)

    def take(self, value=None):
        tok = self.peek()
        if value is not None and tok[1] != value:
            raise FormulaError(f"cần '{value}', gặp '{tok[1]}'")
        self.i += 1
        return tok

    def parse(self):
        v = self.compare()
        if self.i != len(self.tokens):
            raise FormulaError(f"thừa '{self.peek()[1]}'")
        return v

    def compare(self):
        left = self.concat()
        while self.peek()[1] in ("=", "<>", "<", ">", "<=", ">="):
            op = self.take()[1]
            right = self.concat()
            a, b = _flatten(left)[0], _flatten(right)[0]
            try:
                a, b = _num(a), _num(b)
            except FormulaError:
                a, b = str(a).lower(), str(b).lower()
            left = {"=": a == b, "<>": a != b, "<": a < b, ">": a > b, "<=": a <= b, ">=": a >= b}[op]
        return left

    def concat(self):
        left = self.additive()
        while self.peek()[1] == "&":
            self.take()
            right = self.additive()
            left = f"{_text(left)}{_text(right)}"
        return left

    def additive(self):
        left = self.term()
        while self.peek()[1] in ("+", "-"):
            op = self.take()[1]
            right = self.term()
            left = _num(_scalar(left)) + (_num(_scalar(right)) if op == "+" else -_num(_scalar(right)))
        return left

    def term(self):
        left = self.power()
        while self.peek()[1] in ("*", "/"):
            op = self.take()[1]
            right = _num(_scalar(self.power()))
            if op == "/" and right == 0:
                raise FormulaError("chia cho 0")
            left = _num(_scalar(left)) * right if op == "*" else _num(_scalar(left)) / right
        return left

    def power(self):
        left = self.unary()
        while self.peek()[1] == "^":
            self.take()
            left = _num(_scalar(left)) ** _num(_scalar(self.unary()))
        return left

    def unary(self):
        if self.peek()[1] == "-":
            self.take()
            return -_num(_scalar(self.unary()))
        if self.peek()[1] == "+":
            self.take()
            return self.unary()
        v = self.primary()
        if self.peek()[1] == "%":
            self.take()
            v = _num(_scalar(v)) / 100
        return v

    def primary(self):
        kind, text = self.take()
        if kind == "num":
            return float(text)
        if kind == "str":
            return text[1:-1].replace('""', '"')
        if kind == "ref":
            return self.reference(text)
        if kind == "func":
            return self.function(text.upper())
        if text == "(":
            v = self.compare()
            self.take(")")
            return v
        raise FormulaError(f"không hiểu '{text}'")

    def reference(self, text: str):
        sheet = self.sheet
        if "!" in text:
            sheet, text = text.rsplit("!", 1)
            if sheet.startswith("'"):
                sheet = sheet[1:-1].replace("''", "'")
        parts = text.replace("$", "").split(":")
        coords = []
        for p in parts:
            m = re.match(r"([A-Za-z]+)(\d+)", p)
            coords.append((int(m.group(2)), column_index_from_string(m.group(1).upper())))
        if len(coords) == 1:
            return self.ev.value(sheet, *coords[0])
        (r1, c1), (r2, c2) = coords
        return self.ev.range_values(sheet, min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2))

    def args(self) -> List[Any]:
        out = []
        if self.peek()[1] == ")":
            self.take()
            return out
        while True:
            out.append(self.compare())
            if self.peek()[1] == ",":
                self.take()
                continue
            self.take(")")
            return out

    def function(self, name: str):
        a = self.args()
        nums = lambda values: [_num(x) for v in values for x in _flatten(v)
                               if isinstance(x, (int, float)) and not isinstance(x, bool)]
        if name == "SUM":
            return sum(nums(a))
        if name == "PRODUCT":
            return math.prod(nums(a))
        if name in ("MIN", "MAX"):
            values = nums(a)
            return (min if name == "MIN" else max)(values) if values else 0.0
        if name == "AVERAGE":
            values = nums(a)
            if not values:
                raise FormulaError("AVERAGE không có số")
            return sum(values) / len(values)
        if name == "ABS":
            return abs(_num(_scalar(a[0])))
        if name == "PI":
            return math.pi
        if name == "SQRT":
            x = _num(_scalar(a[0]))
            if x < 0:
                raise FormulaError("SQRT của số âm (#NUM!)")
            return math.sqrt(x)
        if name in ("ROUND", "ROUNDUP", "ROUNDDOWN"):
            x, d = _num(_scalar(a[0])), int(_num(_scalar(a[1])) if len(a) > 1 else 0)
            f = 10 ** d
            if name == "ROUND":
                return math.floor(abs(x) * f + 0.5) / f * (1 if x >= 0 else -1)
            op = math.ceil if name == "ROUNDUP" else math.floor
            return op(abs(x) * f - 1e-9 if name == "ROUNDUP" else abs(x) * f + 1e-9) / f * (1 if x >= 0 else -1)
        if name in ("SUMIF", "COUNTIF"):
            rng, crit = _flatten(a[0]), _scalar(a[1])
            if name == "COUNTIF":
                return float(sum(1 for v in rng if _match(v, crit)))
            sums = _flatten(a[2]) if len(a) > 2 else rng
            return sum(_num(s) for v, s in zip(rng, sums) if _match(v, crit) and _is_num(s))
        if name in ("SUMIFS", "COUNTIFS"):
            if name == "SUMIFS":
                target, pairs = _flatten(a[0]), a[1:]
            else:
                target, pairs = None, a
            if len(pairs) % 2:
                raise FormulaError(f"{name} thiếu điều kiện")
            ranges = [(_flatten(pairs[k]), _scalar(pairs[k + 1])) for k in range(0, len(pairs), 2)]
            n = len(ranges[0][0])
            hits = [all(_match(rng[i], crit) for rng, crit in ranges) for i in range(n)]
            if target is None:
                return float(sum(hits))
            return sum(_num(target[i]) for i in range(n) if hits[i] and _is_num(target[i]))
        raise FormulaError(f"hàm {name} chưa hỗ trợ")


def _is_num(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _scalar(v: Any) -> Any:
    return _flatten(v)[0] if isinstance(v, list) else v


def _text(v: Any) -> str:
    v = _scalar(v)
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return "" if v is None else str(v)
